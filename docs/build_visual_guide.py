# ruff: noqa: E501 - Slide copy is kept beside its layout metadata for review.

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Hotspot:
    x: float
    y: float
    title: str
    text: str


@dataclass(frozen=True)
class VisualSlide:
    section: str
    title: str
    path: str
    asset: str
    hotspots: tuple[Hotspot, ...]
    mobile: bool = False
    crop: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0)
    note: str = "После действия обновите список и проверьте результат."


VISUAL_ASSETS = Path("visual-v6")


def _asset_path(guide: Any, filename: str) -> Path:
    path = guide.ASSETS / VISUAL_ASSETS / filename
    if not path.exists():
        path = guide.ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def _add_picture_frame(
    guide: Any,
    slide: Any,
    *,
    asset: str,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    crop: tuple[float, float, float, float],
) -> tuple[float, float, float, float]:
    path = _asset_path(guide, asset)
    guide.add_rect(slide, x, y, w, h, fill=guide.WHITE, line=guide.LINE)
    header_h = 0.32 if label else 0.0
    if label:
        guide.add_rect(slide, x, y, w, header_h, fill=guide.PURPLE_DARK, line=None)
        guide.add_text(
            slide,
            label,
            x + 0.14,
            y + 0.06,
            w - 0.28,
            0.20,
            size=9.5,
            fill=guide.WHITE,
            bold=True,
        )
    inner_x = x + 0.08
    inner_y = y + header_h + 0.08
    inner_w = w - 0.16
    inner_h = h - header_h - 0.16
    left, top, right, bottom = crop
    with Image.open(path) as source:
        source.load()
        crop_box = (
            round(source.width * left),
            round(source.height * top),
            round(source.width * (1 - right)),
            round(source.height * (1 - bottom)),
        )
        cropped = source.crop(crop_box)
        image_w, image_h = cropped.size
        scale = min(inner_w / image_w, inner_h / image_h)
        picture_w = image_w * scale
        picture_h = image_h * scale
        picture_x = inner_x + (inner_w - picture_w) / 2
        picture_y = inner_y + (inner_h - picture_h) / 2
        stream = BytesIO()
        cropped.save(stream, format="PNG", optimize=True)
        stream.seek(0)
        slide.shapes.add_picture(
            stream,
            Inches(picture_x),
            Inches(picture_y),
            width=Inches(picture_w),
            height=Inches(picture_h),
        )
    return picture_x, picture_y, picture_w, picture_h


def _add_hotspot(
    guide: Any,
    slide: Any,
    *,
    number: int,
    image_rect: tuple[float, float, float, float],
    point: Hotspot,
) -> None:
    px, py, pw, ph = image_rect
    marker_x = px + point.x * pw - 0.18
    marker_y = py + point.y * ph - 0.18
    ring = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(marker_x - 0.07),
        Inches(marker_y - 0.07),
        Inches(0.50),
        Inches(0.50),
    )
    ring.fill.background()
    ring.line.color.rgb = guide.color(guide.CORAL)
    ring.line.width = Pt(2.2)
    marker = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(marker_x),
        Inches(marker_y),
        Inches(0.36),
        Inches(0.36),
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = guide.color(guide.CORAL)
    marker.line.fill.background()
    frame = marker.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = str(number)
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.runs[0]
    run.font.name = guide.FONT
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = guide.color(guide.WHITE)


def _add_cursor(
    guide: Any,
    slide: Any,
    *,
    image_rect: tuple[float, float, float, float],
    point: Hotspot,
) -> None:
    px, py, pw, ph = image_rect
    cursor = slide.shapes.add_shape(
        MSO_SHAPE.UP_ARROW,
        Inches(px + point.x * pw + 0.10),
        Inches(py + point.y * ph + 0.12),
        Inches(0.28),
        Inches(0.40),
    )
    cursor.rotation = 315
    cursor.fill.solid()
    cursor.fill.fore_color.rgb = guide.color(guide.WHITE)
    cursor.line.color.rgb = guide.color(guide.PURPLE_DARK)
    cursor.line.width = Pt(1.4)


def _add_action_card(
    guide: Any,
    slide: Any,
    *,
    number: int,
    point: Hotspot,
    x: float,
    y: float,
    w: float,
    h: float,
) -> None:
    guide.add_rect(slide, x, y, w, h, fill=guide.WHITE, line=guide.LINE)
    guide.add_rect(slide, x, y, 0.06, h, fill=guide.CORAL, line=None)
    guide.add_rect(slide, x + 0.18, y + 0.15, 0.34, 0.34, fill=guide.CORAL, line=None, radius=True)
    guide.add_text(
        slide,
        str(number),
        x + 0.18,
        y + 0.16,
        0.34,
        0.28,
        size=11,
        fill=guide.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    guide.add_text(slide, point.title, x + 0.64, y + 0.10, w - 0.80, 0.34, size=11.8, bold=True)
    guide.add_text(slide, point.text, x + 0.64, y + 0.45, w - 0.80, h - 0.51, size=10.0, fill=guide.MUTED)


def _add_result_callout(guide: Any, slide: Any, *, text: str, x: float, y: float, w: float) -> None:
    guide.add_rect(slide, x, y, w, 0.96, fill=guide.TEAL_LIGHT, line=None)
    guide.add_rect(slide, x, y, 0.06, 0.96, fill=guide.TEAL, line=None)
    guide.add_text(slide, "Проверка результата", x + 0.22, y + 0.12, w - 0.36, 0.24, size=11.5, bold=True)
    guide.add_text(slide, text, x + 0.22, y + 0.40, w - 0.36, 0.44, size=10.0, fill=guide.MUTED)


def add_visual_slide(guide: Any, prs: Presentation, spec: VisualSlide) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide)
    guide.add_brand(slide, len(prs.slides), spec.title, spec.section)
    guide.add_rect(slide, 0.48, 1.80, 12.28, 0.36, fill=guide.PURPLE_LIGHT, line=None)
    guide.add_text(slide, "ПУТЬ", 0.66, 1.91, 0.46, 0.15, size=8.5, fill=guide.PURPLE, bold=True)
    guide.add_text(slide, spec.path, 1.20, 1.86, 11.25, 0.22, size=11.5, bold=True)
    if spec.mobile:
        frame_x, frame_y, frame_w, frame_h = 0.58, 2.26, 4.00, 4.60
        panel_x, panel_w = 4.88, 7.80
        label = "Экран на телефоне"
    else:
        frame_x, frame_y, frame_w, frame_h = 0.48, 2.26, 8.52, 4.60
        panel_x, panel_w = 9.30, 3.38
        label = "Экран приложения"
    image_rect = _add_picture_frame(
        guide,
        slide,
        asset=spec.asset,
        x=frame_x,
        y=frame_y,
        w=frame_w,
        h=frame_h,
        label=label,
        crop=spec.crop,
    )
    for index, point in enumerate(spec.hotspots, start=1):
        _add_hotspot(guide, slide, number=index, image_rect=image_rect, point=point)
    count = max(len(spec.hotspots), 1)
    available_h = 3.48
    gap = 0.08
    card_h = min(1.05, (available_h - gap * (count - 1)) / count)
    for index, point in enumerate(spec.hotspots, start=1):
        _add_action_card(
            guide,
            slide,
            number=index,
            point=point,
            x=panel_x,
            y=2.26 + (index - 1) * (card_h + gap),
            w=panel_w,
            h=card_h,
        )
    _add_result_callout(guide, slide, text=spec.note, x=panel_x, y=5.90, w=panel_w)
    guide.add_footer(slide, "Практическое руководство · интерфейс 0.86.2 · Algo MAX")


def add_visual_step_slide(
    guide: Any,
    prs: Presentation,
    spec: VisualSlide,
    *,
    active_index: int,
    show_result: bool,
) -> None:
    if active_index < 1 or active_index > len(spec.hotspots):
        raise ValueError(f"Invalid hotspot index {active_index} for {spec.asset}")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide)
    point = spec.hotspots[active_index - 1]
    guide.add_brand(slide, len(prs.slides), spec.title, spec.section)
    guide.add_rect(slide, 0.48, 1.80, 12.28, 0.36, fill=guide.PURPLE_LIGHT, line=None)
    guide.add_text(slide, "ПУТЬ", 0.66, 1.91, 0.46, 0.15, size=8.5, fill=guide.PURPLE, bold=True)
    guide.add_text(slide, spec.path, 1.20, 1.86, 11.25, 0.22, size=11.5, bold=True)

    if spec.mobile:
        frame_x, frame_y, frame_w, frame_h = 0.58, 2.26, 4.18, 4.60
        panel_x, panel_w = 5.08, 7.60
        label = "Экран на телефоне"
    else:
        frame_x, frame_y, frame_w, frame_h = 0.48, 2.26, 9.10, 4.60
        panel_x, panel_w = 9.86, 2.82
        label = "Экран приложения"

    image_rect = _add_picture_frame(
        guide,
        slide,
        asset=spec.asset,
        x=frame_x,
        y=frame_y,
        w=frame_w,
        h=frame_h,
        label=label,
        crop=spec.crop,
    )
    _add_hotspot(guide, slide, number=active_index, image_rect=image_rect, point=point)
    _add_cursor(guide, slide, image_rect=image_rect, point=point)
    if "qr" in spec.asset.casefold():
        guide.add_rect(slide, frame_x + 0.18, frame_y + 0.48, 1.78, 0.34, fill=guide.YELLOW, line=None)
        guide.add_text(
            slide,
            "ПРИМЕР, НЕ СКАНИРОВАТЬ",
            frame_x + 0.24,
            frame_y + 0.57,
            1.66,
            0.16,
            size=8.5,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    guide.add_rect(slide, panel_x, 2.26, panel_w, 0.46, fill=guide.PURPLE_DARK, line=None)
    guide.add_text(
        slide,
        f"ШАГ {active_index} ИЗ {len(spec.hotspots)}",
        panel_x + 0.18,
        2.39,
        panel_w - 0.36,
        0.18,
        size=9.5,
        fill=guide.WHITE,
        bold=True,
    )
    _add_action_card(
        guide,
        slide,
        number=active_index,
        point=point,
        x=panel_x,
        y=2.84,
        w=panel_w,
        h=1.64,
    )
    result_text = (
        spec.note
        if show_result
        else "После нажатия дождитесь изменения экрана и переходите к следующему шагу."
    )
    _add_result_callout(guide, slide, text=result_text, x=panel_x, y=4.72, w=panel_w)
    guide.add_rect(slide, panel_x, 5.92, panel_w, 0.66, fill=guide.YELLOW_LIGHT, line=None)
    guide.add_text(
        slide,
        "Курсор и номер показывают единственное действие этого кадра.",
        panel_x + 0.18,
        6.08,
        panel_w - 0.36,
        0.32,
        size=9.2,
        fill=guide.MUTED,
    )
    guide.add_footer(slide, "Пошаговое видео · Algo MAX")


def add_visual_error_slide(
    guide: Any,
    prs: Presentation,
    spec: VisualSlide,
    *,
    error_title: str,
    error_text: str,
    recovery: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide)
    guide.add_brand(slide, len(prs.slides), error_title, "Типичная ситуация")
    guide.add_rect(slide, 0.48, 1.80, 12.28, 0.36, fill=guide.CORAL_LIGHT, line=None)
    guide.add_text(slide, "ВОСПРОИЗВЕДЕННАЯ ОШИБКА", 0.66, 1.90, 3.20, 0.17, size=8.5, fill=guide.CORAL, bold=True)

    if spec.mobile:
        frame_x, frame_y, frame_w, frame_h = 0.58, 2.26, 4.18, 4.60
        panel_x, panel_w = 5.08, 7.60
        label = "Экран на телефоне"
    else:
        frame_x, frame_y, frame_w, frame_h = 0.48, 2.26, 9.10, 4.60
        panel_x, panel_w = 9.86, 2.82
        label = "Экран приложения"
    _add_picture_frame(
        guide,
        slide,
        asset=spec.asset,
        x=frame_x,
        y=frame_y,
        w=frame_w,
        h=frame_h,
        label=label,
        crop=spec.crop,
    )

    guide.add_rect(slide, panel_x, 2.26, panel_w, 1.64, fill=guide.CORAL_LIGHT, line=guide.CORAL)
    guide.add_text(slide, "Что произошло", panel_x + 0.20, 2.48, panel_w - 0.40, 0.30, size=12, bold=True)
    guide.add_text(slide, error_text, panel_x + 0.20, 2.90, panel_w - 0.40, 0.78, size=10.2, fill=guide.MUTED)
    guide.add_rect(slide, panel_x, 4.18, panel_w, 1.78, fill=guide.TEAL_LIGHT, line=guide.TEAL)
    guide.add_text(slide, "Как исправить", panel_x + 0.20, 4.40, panel_w - 0.40, 0.30, size=12, bold=True)
    guide.add_text(slide, recovery, panel_x + 0.20, 4.82, panel_w - 0.40, 0.90, size=10.2, fill=guide.MUTED)
    guide.add_footer(slide, "Пошаговое видео · Algo MAX")


def add_cover(guide: Any, prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide, guide.WHITE)
    guide.add_rect(slide, 0, 0, 5.15, 7.5, fill=guide.PURPLE_DARK, line=None)
    guide.add_rect(slide, 0.64, 0.62, 0.58, 0.58, fill=guide.YELLOW, line=None, radius=True)
    guide.add_text(
        slide,
        "A",
        0.64,
        0.64,
        0.58,
        0.50,
        size=22,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    guide.add_text(slide, "ALGO MAX", 1.43, 0.78, 2.10, 0.30, size=18, fill=guide.WHITE, bold=True)
    guide.add_text(slide, "Полное руководство\nпо ролям", 0.64, 1.70, 4.00, 1.34, size=30, fill=guide.WHITE, bold=True)
    guide.add_text(
        slide,
        "Конкретные кнопки, права доступа, ожидаемый результат и исправление ошибок",
        0.66,
        3.28,
        3.88,
        1.26,
        size=15,
        fill=guide.WHITE,
    )
    guide.add_rect(slide, 0.66, 5.58, 3.72, 0.58, fill=guide.YELLOW, line=None)
    guide.add_text(slide, "70+ актуальных экранов", 0.86, 5.75, 3.26, 0.24, size=13, bold=True)
    guide.add_text(slide, "Редакция 7 · интерфейс 0.86.2 · сентябрь 2026", 0.66, 6.70, 4.28, 0.24, size=10.2, fill="CBB9E8")
    for asset, x, y, w, h in (
        ("student-dashboard-desktop.png", 5.56, 0.64, 7.14, 2.02),
        ("store-desktop.png", 5.56, 2.84, 7.14, 2.02),
        ("orders-collect-desktop.png", 5.56, 5.04, 7.14, 1.80),
    ):
        _add_picture_frame(guide, slide, asset=asset, x=x, y=y, w=w, h=h, label="", crop=(0, 0, 0, 0))


def add_contents(guide: Any, prs: Presentation) -> None:
    guide.add_cards_slide(
        prs,
        kicker="Навигация",
        title="Содержание по рабочим разделам",
        cards=[
            ("1. Вход и роли", "Профиль, город, ребенок и доступные разделы.", "purple"),
            ("2. Магазин", "Поиск, фильтры, карточка, корзина и код.", "yellow"),
            ("3. Заказы", "Склад, сборка, доставка, получение и отмена.", "coral"),
            ("4. Ученики и AC", "Карточка, статус, баланс, начисления и отчет.", "teal"),
            ("5. Рассылки", "Аудитория, площадки, эмодзи и проверка.", "purple"),
            ("6–7. Управление", "Сотрудники, товары, склады, импорт, QR и помощь.", "teal"),
        ],
        note_title="Принцип руководства",
        note_text="Одна операция — один экран. Номера на скриншоте совпадают с точными действиями справа.",
        note_kind="yellow",
    )


def add_how_to_read(guide: Any, prs: Presentation) -> None:
    add_visual_slide(
        guide,
        prs,
        VisualSlide(
            "Как читать руководство",
            "Номер показывает точное место нажатия",
            "Пример: Магазин -> карточка товара",
            "store-mobile.png",
            (
                Hotspot(0.50, 0.23, "Найдите номер", "Метка стоит рядом с нужной кнопкой или полем."),
                Hotspot(0.47, 0.46, "Выполните действие", "Пояснение справа использует тот же номер."),
                Hotspot(0.50, 0.84, "Сверьте результат", "Нижний блок сообщает, что должно измениться."),
            ),
            mobile=True,
            note="Если экран отличается, сначала проверьте роль, выбранного ребенка и рабочий город.",
        ),
    )


def add_role_reference(guide: Any, prs: Presentation) -> None:
    guide.add_cards_slide(
        prs,
        kicker="Роли",
        title="Кто что делает в системе",
        cards=[
            ("Ученик", "Покупает за свои AC, видит баланс, историю и собственные заказы.", "yellow"),
            ("Родитель", "Переключает детей, оформляет покупку и дает ребенку QR-доступ.", "teal"),
            ("Преподаватель", "Начисляет AC своим ученикам, получает и выдает их заказы.", "purple"),
            ("Куратор", "Работает с разрешенными группами, заказами и адресными рассылками.", "coral"),
            ("Администратор", "Ведет учеников, склады, товары, сотрудников и полный цикл заказа.", "teal"),
            ("Директор", "Контролирует несколько городов, отчеты и права сотрудников.", "purple"),
        ],
        note_title="Данные разделены по городам",
        note_text="Сотрудник с ролями в нескольких городах переключает только назначенные ему филиалы и проверяет город перед изменением.",
        note_kind="yellow",
    )
    guide.add_cards_slide(
        prs,
        kicker="Заказы",
        title="Пять рабочих состояний заказа",
        cards=[
            ("1. Зарезервирован", "Заказ создан; склад для физического товара еще не подтвержден.", "yellow"),
            ("2. Ожидает доставки", "Склад выбран, товар собран или готовится к отправке на площадку.", "coral"),
            ("3. На площадке", "Доставка подтверждена; заказ доступен преподавателю.", "teal"),
            ("4. У учителя", "Преподаватель фактически получил комплект.", "purple"),
            ("5. Получен", "Семья видит «Получен», сотрудники — «Передан ученику».", "teal"),
            ("Отменен", "Отменять на любом этапе могут только администратор и директор.", "coral"),
        ],
        note_title="Статус — это факт",
        note_text="Не переводите заказ дальше заранее: каждое состояние должно соответствовать реальному действию.",
        note_kind="yellow",
    )
    guide.add_steps_slide(
        prs,
        kicker="Доступ",
        title="Что проверить перед любой операцией",
        steps=[
            ("Откройте из MAX", "Запускайте mini-app из подтвержденного аккаунта MAX."),
            ("Проверьте профиль", "Сверьте имя и активную роль в верхнем блоке."),
            ("Проверьте город", "Сотрудник с доступом к нескольким городам выбирает рабочий город в верхней панели."),
            ("Проверьте ребенка", "Родитель выбирает ребенка до магазина, корзины или истории."),
            ("Обновите данные", "После изменения нажмите кнопку с круговыми стрелками."),
        ],
        side_title="Если нужного раздела нет",
        side_items=[
            "Роль не назначена или отозвана.",
            "Выбран другой город.",
            "Связь с учеником неактивна.",
            "Нужно повторно открыть приложение из MAX.",
        ],
        callout_title="Не используйте чужие ссылки",
        callout_text="Персональные и одноразовые приглашения нельзя пересылать другому пользователю.",
        callout_kind="coral",
    )


def _profile_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Ученик", "Главная ученика на компьютере", "Главная", "student-dashboard-desktop.png", (
            Hotspot(0.40, 0.14, "Проверьте имя", "Приветствие и профиль должны относиться к ученику."),
            Hotspot(0.76, 0.14, "Обновите баланс", "Круговая стрелка повторно запрашивает AC."),
            Hotspot(0.92, 0.14, "Откройте корзину", "Счетчик показывает добавленные позиции."),
        ), note="Ученик видит только свой баланс, свою историю и свои заказы."),
        VisualSlide("Ученик", "Главная ученика на телефоне", "Нижнее меню -> Главная", "student-dashboard-mobile.png", (
            Hotspot(0.50, 0.15, "Проверьте кабинет", "Имя в приветствии подтверждает текущий профиль."),
            Hotspot(0.50, 0.58, "Откройте быстрый раздел", "Карточки ведут к магазину и заказам."),
            Hotspot(0.52, 0.94, "Используйте меню", "Основные разделы закреплены снизу."),
        ), mobile=True, note="Дополнительные пункты находятся в меню «Еще» и зависят от роли."),
        VisualSlide("Родитель", "Переключение между детьми", "Главная -> выбор ребенка", "parent-dashboard-desktop.png", (
            Hotspot(0.36, 0.16, "Выберите ребенка", "Баланс и корзина переключатся вместе с именем."),
            Hotspot(0.78, 0.16, "Сверьте AC", "Сумма относится только к активному ребенку."),
            Hotspot(0.73, 0.65, "Откройте заказ", "Список показывает заказы связанных детей."),
        ), note="Перед покупкой повторно проверьте имя ребенка в корзине."),
        VisualSlide("Родитель", "Главная родителя на телефоне", "Главная -> ребенок", "parent-dashboard-mobile.png", (
            Hotspot(0.50, 0.14, "Профиль родителя", "Имя подтверждает, кто открыл приложение."),
            Hotspot(0.50, 0.34, "Активный ребенок", "Нажмите имя, чтобы сменить профиль."),
            Hotspot(0.50, 0.82, "Разделы семьи", "Корзина и заказы доступны снизу."),
        ), mobile=True, crop=(0, 0, 0, 0.48), note="У каждого ребенка отдельные баланс, корзина, заказы и история AC."),
        VisualSlide("Преподаватель", "Ученики и заказы преподавателя", "Главная", "teacher-dashboard-desktop.png", (
            Hotspot(0.43, 0.34, "Список учеников", "Показаны только назначенные группы и ученики."),
            Hotspot(0.80, 0.33, "Заказы к выдаче", "Здесь видны комплекты своих учеников."),
            Hotspot(0.09, 0.52, "Рабочие разделы", "Начисления, заказы и отчет доступны слева."),
        ), note="Если ученика нет, сбросьте фильтр группы, затем обратитесь к администратору."),
        VisualSlide("Преподаватель", "Рабочий экран преподавателя на телефоне", "Главная -> нижнее меню", "teacher-dashboard-mobile.png", (
            Hotspot(0.50, 0.28, "Выберите группу", "Фильтр сокращает список до нужной группы."),
            Hotspot(0.50, 0.58, "Откройте ученика", "Карточка дает доступ к разрешенной истории."),
            Hotspot(0.58, 0.92, "Начисления", "Кнопка открывает массовое начисление AC."),
        ), mobile=True, crop=(0, 0, 0, 0.46), note="Преподаватель не видит чужие группы и не может отменять заказы."),
        VisualSlide("Преподаватель", "Первичная привязка к группам LMS", "Первый вход -> Укажите фамилию и имя", "teacher-profile-first-login-desktop.png", (
            Hotspot(0.40, 0.52, "Введите фамилию", "Используйте точное написание из LMS."),
            Hotspot(0.60, 0.52, "Введите имя", "Имя профиля MAX может отличаться."),
            Hotspot(0.50, 0.61, "Сохраните и найдите", "Кнопка запускает автоматический поиск групп."),
        ), note="После сохранения главная страница показывает найденные группы и учеников."),
        VisualSlide("Преподаватель", "Проверка найденных групп", "Рабочий профиль -> Изменить ФИО", "teacher-profile-groups-desktop.png", (
            Hotspot(0.40, 0.49, "Сверьте ФИО", "Фамилия и имя должны совпадать с LMS."),
            Hotspot(0.50, 0.57, "Проверьте результат", "Блок показывает число учеников и найденные группы."),
            Hotspot(0.57, 0.65, "Повторите поиск", "Сохранение заново рассчитывает привязки."),
        ), note="Если групп нет, исправьте ФИО; вручную дублировать преподавателя не нужно."),
        VisualSlide("Куратор", "Рабочая область куратора", "Главная", "curator-dashboard-desktop.png", (
            Hotspot(0.42, 0.34, "Проверьте учеников", "Список ограничен разрешенными группами."),
            Hotspot(0.80, 0.33, "Контролируйте выдачу", "Куратор видит связанные заказы."),
            Hotspot(0.09, 0.58, "Откройте рассылки", "Адресные новости доступны в меню."),
        ), note="Куратор может подтверждать доставку на площадку, но не отменять заказ."),
        VisualSlide("Администратор", "Главная администратора", "Главная", "admin-dashboard-desktop.png", (
            Hotspot(0.42, 0.34, "Ученики и группы", "Быстрый список показывает активных учеников города."),
            Hotspot(0.80, 0.34, "Заказы к выдаче", "Карточки отражают фактические статусы."),
            Hotspot(0.09, 0.58, "Управление", "Товары, склады и сотрудники находятся здесь."),
        ), note="Администратор работает только в назначенном городе."),
        VisualSlide("Несколько ролей", "Единый профиль сотрудника", "Рабочий профиль -> основное меню", "combined-role-profile-desktop.png", (
            Hotspot(0.35, 0.17, "Проверьте роли", "Активные роли перечислены через запятую."),
            Hotspot(0.23, 0.23, "Измените ФИО", "Имя меняется один раз для всего профиля."),
            Hotspot(0.09, 0.54, "Используйте общие права", "Меню объединяет возможности всех активных ролей."),
        ), note="Отозванные роли не показываются и не расширяют доступ."),
        VisualSlide("Несколько городов", "Переключение рабочего города", "Верхняя панель -> Город / партнер", "director-city-switcher-desktop.png", (
            Hotspot(0.50, 0.41, "Найдите город", "Введите название или выберите строку списка."),
            Hotspot(0.50, 0.49, "Выберите филиал", "Подсветка показывает текущий рабочий контекст."),
            Hotspot(0.50, 0.65, "Смените город", "Доступны только города с активными ролями сотрудника."),
        ), note="После выбора проверьте город в профиле до изменения учеников, заказов или остатков."),
        VisualSlide("Директор", "Контроль филиала директором", "Главная -> город", "director-dashboard-desktop.png", (
            Hotspot(0.42, 0.34, "Проверьте учеников", "Список относится к выбранному городу."),
            Hotspot(0.80, 0.34, "Контролируйте заказы", "Проверьте зависшие этапы перед выдачей."),
            Hotspot(0.09, 0.58, "Откройте управление", "Сотрудники, склады и импорт доступны директору."),
        ), note="При нескольких городах сначала переключите город, затем выполняйте изменения."),
    )


def _store_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Магазин", "Поиск, категории и сортировка", "Магазин", "store-desktop.png", (
            Hotspot(0.70, 0.27, "Введите название", "Поиск работает по части названия товара."),
            Hotspot(0.31, 0.39, "Выберите категорию", "«Все» возвращает полный каталог."),
            Hotspot(0.86, 0.39, "Настройте сортировку", "Выберите популярность, цену или новизну."),
        ), note="После фильтра количество карточек и подпись найденных товаров должны измениться."),
        VisualSlide("Магазин", "Выбор порядка товаров", "Магазин -> Сортировка", "store-sort-desktop.png", (
            Hotspot(0.65, 0.29, "Откройте меню", "Нажмите поле «Сортировка»."),
            Hotspot(0.65, 0.38, "Выберите порядок", "Доступны популярность, цена и новизна."),
            Hotspot(0.30, 0.25, "Сверьте категории", "Одинаковые названия с лишними пробелами объединены."),
        ), note="Меню открывается поверх каталога и не перекрывается карточками товаров."),
        VisualSlide("Магазин", "Каталог на телефоне", "Магазин", "store-mobile.png", (
            Hotspot(0.50, 0.20, "Поиск", "Введите название и дождитесь обновления списка."),
            Hotspot(0.49, 0.36, "Категории", "Лента прокручивается горизонтально."),
            Hotspot(0.50, 0.78, "Подробнее", "Откройте карточку выбранной награды."),
        ), mobile=True, note="Цена и наличие видны до добавления товара в корзину."),
        VisualSlide("Магазин", "Дополнительные фильтры", "Магазин -> Фильтры", "store-filters-mobile.png", (
            Hotspot(0.50, 0.30, "Выберите порядок", "Сортировка меняет порядок карточек."),
            Hotspot(0.50, 0.50, "Оставьте доступные", "«В наличии» скрывает недоступные позиции."),
            Hotspot(0.50, 0.70, "Сбросьте", "Кнопка сброса возвращает исходный каталог."),
        ), mobile=True, note="Активные переключатели выделяются; закрытие панели не сбрасывает выбор."),
        VisualSlide("Магазин", "Карточка товара", "Магазин -> Подробнее", "product-dialog-mobile.png", (
            Hotspot(0.50, 0.28, "Сверьте товар", "Проверьте название, описание, цену и тип выдачи."),
            Hotspot(0.50, 0.64, "Выберите количество", "Не превышайте доступный остаток."),
            Hotspot(0.50, 0.84, "Добавьте", "Кнопка переносит позицию в корзину выбранного ученика."),
        ), mobile=True, note="Название не перекрывает изображение и полностью читается в карточке."),
        VisualSlide("Корзина", "Проверка количества и суммы", "Корзина", "cart-mobile.png", (
            Hotspot(0.50, 0.30, "Измените количество", "Минус и плюс сразу пересчитывают сумму."),
            Hotspot(0.50, 0.57, "Сверьте остаток", "Показан баланс после списания."),
            Hotspot(0.50, 0.76, "Оформите", "Нажмите один раз и дождитесь подтверждения."),
        ), mobile=True, note="Получатель, итог и остаток на счете должны быть верными до подтверждения."),
        VisualSlide("Корзина", "Финальное подтверждение заказа", "Корзина -> Оформить заказ", "checkout-dialog-mobile.png", (
            Hotspot(0.50, 0.34, "Проверьте получателя", "У родителя здесь указано имя выбранного ребенка."),
            Hotspot(0.50, 0.54, "Сверьте итог", "Сумма списания не должна превышать баланс."),
            Hotspot(0.50, 0.78, "Подтвердите", "Повторное нажатие не требуется."),
        ), mobile=True, note="После успеха корзина очищается, а заказ появляется в соответствующей вкладке."),
        VisualSlide("Цифровой товар", "Получение и копирование кода", "Заказы -> Получены -> Подробнее", "visual/order-digital-mobile.png", (
            Hotspot(0.50, 0.58, "Найдите код", "Код хранится внутри выполненного цифрового заказа."),
            Hotspot(0.50, 0.72, "Скопируйте", "Используйте кнопку рядом с кодом."),
            Hotspot(0.50, 0.88, "Проверьте выдачу", "История подтверждает автоматическую выдачу."),
        ), mobile=True, note="Один цифровой код выдается один раз и исключается из повторной продажи."),
    )


def _order_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Семья", "Фильтры заказов ученика и родителя", "Заказы", "family-orders-desktop.png", (
            Hotspot(0.39, 0.23, "Выберите состояние", "Для семьи финальная вкладка называется «Получены»."),
            Hotspot(0.78, 0.16, "Найдите заказ", "Поиск принимает номер, ученика или товар."),
            Hotspot(0.84, 0.52, "Откройте карточку", "«Подробнее» показывает состав и историю."),
        ), note="Выпадающий фильтр не дублируется: состояние выбирается только вкладками."),
        VisualSlide("Семья", "Заказы на телефоне", "Заказы -> состояние", "family-orders-mobile.png", (
            Hotspot(0.50, 0.18, "Прокрутите статусы", "Вкладки доступны горизонтальным свайпом."),
            Hotspot(0.50, 0.45, "Сверьте карточку", "Видны товар, сумма и текущий этап."),
            Hotspot(0.50, 0.72, "Подробнее", "Откройте историю конкретного заказа."),
        ), mobile=True, crop=(0, 0, 0, 0.42), note="Ученик и родитель не видят кнопку отмены ни на одном этапе."),
        VisualSlide("Семья", "Карточка и история заказа", "Заказы -> Подробнее", "family-order-dialog-mobile.png", (
            Hotspot(0.50, 0.28, "Сверьте состав", "Проверьте товары, количество и сумму."),
            Hotspot(0.50, 0.58, "Читайте историю", "Каждая запись содержит дату и фактическое действие."),
            Hotspot(0.90, 0.08, "Закройте", "Крестик возвращает в список без изменения заказа."),
        ), mobile=True, note="Служебные подписи из внутренних сводок пользователю не показываются."),
        VisualSlide("Заказы", "Рабочие вкладки заказов", "Заказы", "admin-orders-tabs-desktop.png", (
            Hotspot(0.38, 0.20, "Выберите состояние", "Верхний ряд фильтрует фактический статус."),
            Hotspot(0.28, 0.43, "Выберите этап работы", "Назначить склад, Собрать или Распределить."),
            Hotspot(0.88, 0.43, "Обновите", "Круговая стрелка применяет завершенные отметки."),
        ), note="Переключатели не дублируются и всегда видны до первого нажатия."),
        VisualSlide("Назначить склад", "Очередь зарезервированных заказов", "Заказы -> Назначить склад", "orders-assign-desktop.png", (
            Hotspot(0.22, 0.45, "Выберите заказ", "Чекбокс включает заказ в массовую операцию."),
            Hotspot(0.62, 0.44, "Сверьте склад", "Подсказка показывает доступный источник."),
            Hotspot(0.96, 0.44, "Откройте", "Стрелка открывает карточку и состав."),
        ), note="Выбор чекбокса не сворачивает площадку и не меняет статус сам по себе."),
        VisualSlide("Назначить склад", "Склад для каждой позиции", "Заказ -> склад для списания", "order-warehouse-dialog-desktop.png", (
            Hotspot(0.48, 0.54, "Выберите склад", "Список показывает доступный остаток."),
            Hotspot(0.50, 0.91, "Подтвердите", "Кнопка фиксирует склады всех позиций."),
            Hotspot(0.94, 0.08, "Закройте", "Крестик оставляет заказ зарезервированным."),
        ), note="После подтверждения заказ переходит в «Ожидает доставки»."),
        VisualSlide("Назначить склад", "Назначение склада на телефоне", "Заказ -> склад для списания", "order-warehouse-dialog-mobile.png", (
            Hotspot(0.50, 0.48, "Проверьте позицию", "Название товара не перекрывает изображение."),
            Hotspot(0.50, 0.67, "Выберите склад", "Доступный остаток указан в списке."),
            Hotspot(0.50, 0.90, "Подтвердите", "Сохраните выбор один раз."),
        ), mobile=True, note="Если остатка нет, закройте карточку и проверьте другой склад или товар."),
        VisualSlide("Собрать", "Общий чек-лист по складам", "Заказы -> Собрать", "orders-collect-desktop.png", (
            Hotspot(0.22, 0.44, "Откройте склад", "Список разделен по источнику товара."),
            Hotspot(0.28, 0.65, "Отметьте товар", "Одинаковые позиции объединены, количество суммируется."),
            Hotspot(0.88, 0.35, "Обновите", "Только обновление фиксирует готовый набор."),
        ), note="Чекбокс служит промежуточной отметкой; строка не исчезает сразу после нажатия."),
        VisualSlide("Распределить", "Площадки и преподаватели", "Заказы -> Распределить", "orders-route-desktop.png", (
            Hotspot(0.24, 0.48, "Откройте площадку", "Внутри заказы сгруппированы по преподавателю."),
            Hotspot(0.28, 0.68, "Отметьте доставку", "Ставьте галочку после фактической доставки."),
            Hotspot(0.88, 0.35, "Подтвердите", "Рабочая кнопка применяет выбранные заказы."),
        ), note="Заказы в статусах «На площадке» и «У учителя» больше не показываются в сборке."),
        VisualSlide("Отмена", "Причина отмены заказа", "Заказы -> Подробнее -> Другие действия -> Отменить", "order-cancel-form-desktop.png", (
            Hotspot(0.50, 0.52, "Выберите причину", "Используйте готовый вариант или свой текст."),
            Hotspot(0.50, 0.72, "Проверьте возврат", "AC вернутся ученику после подтверждения."),
            Hotspot(0.50, 0.88, "Подтвердите", "Отмена необратима обычной кнопкой."),
        ), note="Отмену на любом этапе выполняют только администратор или директор."),
        VisualSlide("Преподаватель", "Заказы только своих учеников", "Заказы", "teacher-orders-mobile.png", (
            Hotspot(0.50, 0.16, "Выберите состояние", "Преподаватель видит только свою очередь."),
            Hotspot(0.50, 0.42, "Сверьте ученика", "Имя и группа должны быть вашими."),
            Hotspot(0.50, 0.68, "Откройте заказ", "Проверьте состав до получения."),
        ), mobile=True, crop=(0, 0, 0, 0.36), note="Количество и состав доставленных заказов видны в общей выжимке и карточках."),
        VisualSlide("Преподаватель", "Получение заказа на площадке", "Заказы -> На площадке", "teacher-orders-venue-mobile.png", (
            Hotspot(0.50, 0.16, "Откройте «На площадке»", "Здесь только фактически доставленные заказы."),
            Hotspot(0.50, 0.52, "Сверьте состав", "Проверьте товар и количество."),
            Hotspot(0.50, 0.72, "Нажмите «Учитель получил»", "Статус меняется после передачи преподавателю."),
        ), mobile=True, crop=(0, 0, 0, 0.28), note="Следующее действие — «Передать ученику» после фактической выдачи ребенку."),
        VisualSlide("Преподаватель", "Выдача заказа ученику", "Заказы -> У учителя", "teacher-orders-teacher-mobile.png", (
            Hotspot(0.50, 0.16, "Откройте «У учителя»", "Здесь находятся принятые преподавателем комплекты."),
            Hotspot(0.50, 0.50, "Сверьте ученика", "Проверьте имя, товар и количество перед выдачей."),
            Hotspot(0.50, 0.72, "Передайте ученику", "Нажмите кнопку только после фактической выдачи."),
        ), mobile=True, crop=(0, 0, 0, 0.30), note="После действия заказ перемещается во вкладку «Получены»."),
        VisualSlide("Преподаватель", "Результат выдачи", "Заказы -> Получены -> Подробнее", "teacher-order-issued-history-mobile.png", (
            Hotspot(0.50, 0.24, "Проверьте состояние", "Карточка подтверждает завершенную выдачу."),
            Hotspot(0.50, 0.68, "Сверьте историю", "Последняя запись — «Заказ передан ученику»."),
        ), mobile=True, note="Повторное нажатие не требуется: история хранит дату и итоговое состояние."),
    )


def _student_ac_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("История учеников", "Реестр и фильтр по группам", "История учеников", "student-registry-desktop.png", (
            Hotspot(0.35, 0.32, "Введите запрос", "Ищите по ФИО, группе, ID или преподавателю."),
            Hotspot(0.76, 0.32, "Выберите группу", "Новый фильтр групп сохраняется."),
            Hotspot(0.94, 0.52, "Откройте карточку", "Стрелка показывает данные и историю ученика."),
        ), note="Отдельного старого фильтра «История ученика» больше нет."),
        VisualSlide("Преподаватель", "Реестр только своих учеников", "История учеников", "teacher-student-registry-desktop.png", (
            Hotspot(0.35, 0.30, "Найдите ученика", "Поиск работает внутри разрешенных групп."),
            Hotspot(0.77, 0.30, "Выберите группу", "Чужие группы не отображаются."),
            Hotspot(0.94, 0.52, "Откройте карточку", "Преподаватель видит разрешенную историю."),
        ), note="Куратор работает по тому же принципу в пределах назначенных групп."),
        VisualSlide("Карточка ученика", "Статус, день рождения, баланс и история", "История учеников -> открыть ученика", "student-card-desktop.png", (
            Hotspot(0.34, 0.35, "Проверьте данные", "Группа, преподаватель и день рождения хранятся вместе."),
            Hotspot(0.70, 0.48, "Измените статус", "Укажите новое состояние и причину."),
            Hotspot(0.70, 0.68, "Скорректируйте AC", "Введите итоговый баланс и обязательную причину."),
        ), crop=(0.16, 0.24, 0.04, 0.16), note="История сохраняет автора, дату, старое и новое значение."),
        VisualSlide("Новый ученик", "Создание ученика и связи с родителем", "История учеников -> Добавить", "student-create-desktop.png", (
            Hotspot(0.34, 0.38, "Заполните ученика", "Укажите ФИО, ID, группу, статус и день рождения."),
            Hotspot(0.70, 0.48, "Добавьте родителя", "Связь родитель–ученик создается сразу при наличии данных."),
            Hotspot(0.70, 0.70, "Сохраните", "MAX-связь ребенка можно добавить позже."),
        ), crop=(0.16, 0.05, 0.04, 0.34), note="Связь с MAX необязательна; обязательна только корректная карточка ученика."),
        VisualSlide("Начисления", "Групповое начисление AC", "Начисления", "accrual-mobile.png", (
            Hotspot(0.50, 0.31, "Выберите группу", "Ниже останутся ученики этой группы."),
            Hotspot(0.50, 0.55, "Выберите причину", "Сумма связана с правилом текущего города."),
            Hotspot(0.50, 0.82, "Отметьте учеников", "Кнопка показывает количество выбранных."),
        ), mobile=True, crop=(0, 0, 0, 0.30), note="После успеха проверьте новые операции в истории AC."),
        VisualSlide("Начисления", "Своя причина и сумма", "Начисления -> Своя причина и сумма", "accrual-custom-mobile.png", (
            Hotspot(0.50, 0.54, "Выберите свой вариант", "Появятся два ручных поля."),
            Hotspot(0.50, 0.66, "Введите причину", "Текст должен объяснять начисление."),
            Hotspot(0.50, 0.76, "Введите сумму", "Проверьте число до отправки."),
        ), mobile=True, crop=(0, 0, 0, 0.30), note="Кастомная операция сохраняется с автором и доступна в отчете."),
        VisualSlide("Отчет AC", "Период, преподаватель и группа", "Отчет AC", "ac-report-desktop.png", (
            Hotspot(0.31, 0.30, "Выберите период", "Используйте быстрый интервал или даты."),
            Hotspot(0.38, 0.47, "Уточните преподавателя", "Можно оставить всех сотрудников."),
            Hotspot(0.74, 0.47, "Уточните группу", "Фильтры применяются вместе."),
        ), note="Нажмите «Показать» и сверьте сумму, причину, ученика и автора каждой операции."),
        VisualSlide("Начисления", "Системная причина на день рождения", "Начисления -> Настроить причины", "birthday-accrual-rule-desktop.png", (
            Hotspot(0.49, 0.36, "Найдите системную строку", "Значок подарка заменяет кнопку удаления."),
            Hotspot(0.58, 0.36, "Измените сумму", "Название причины остается неизменным."),
            Hotspot(0.59, 0.80, "Сохраните", "Кнопка закреплена и остается видимой при прокрутке."),
        ), note="Настроенная сумма начисляется автоматически один раз в день рождения ученика."),
    )


def _broadcast_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Рассылки", "Категория, формат и группы", "Рассылки -> Получатели", "broadcast-targets-desktop.png", (
            Hotspot(0.27, 0.34, "Выберите получателей", "Родители, ученики или обе категории."),
            Hotspot(0.29, 0.53, "Выберите формат", "Очные, онлайн или индивидуальные группы."),
            Hotspot(0.31, 0.75, "Уточните площадки", "Выберите площадки и конкретные группы."),
        ), note="Перед отправкой система рассчитывает фактическое число получателей."),
        VisualSlide("Рассылки", "Редактор площадки", "Рассылки -> Площадки -> Добавить", "broadcast-venue-editor-mobile.png", (
            Hotspot(0.50, 0.22, "Введите название", "Используйте адрес или понятное имя."),
            Hotspot(0.50, 0.40, "Добавьте слова", "Группы подбираются по вхождению в название."),
            Hotspot(0.50, 0.78, "Уточните вручную", "Ручные галочки имеют приоритет."),
        ), mobile=True, crop=(0, 0.22, 0, 0.36), note="После сохранения откройте площадку повторно и сверьте выбранные группы."),
        VisualSlide("Рассылки", "Текст новости и изображение", "Рассылки -> Новость", "broadcast-news-mobile.png", (
            Hotspot(0.50, 0.25, "Введите заголовок", "Он должен кратко объяснять тему."),
            Hotspot(0.50, 0.48, "Напишите сообщение", "Укажите дату, место и действие получателя."),
            Hotspot(0.50, 0.76, "Добавьте фото", "Поддерживаются JPEG, PNG и WebP."),
        ), mobile=True, crop=(0, 0.17, 0, 0.36), note="Черновик сохраняется автоматически до финальной отправки."),
        VisualSlide("Рассылки", "Добавление эмодзи", "Рассылки -> Новость -> кнопка эмодзи", "broadcast-news-emoji-mobile.png", (
            Hotspot(0.50, 0.46, "Откройте набор", "Кнопка показывает доступные эмодзи."),
            Hotspot(0.50, 0.62, "Выберите символ", "Он вставится в позицию курсора."),
            Hotspot(0.50, 0.78, "Закройте набор", "Текст сообщения сохранится."),
        ), mobile=True, crop=(0, 0.16, 0, 0.34), note="Не требуется системная клавиатура: эмодзи доступны внутри редактора."),
        VisualSlide("Рассылки", "Финальная проверка", "Рассылки -> Проверка", "broadcast-review-mobile.png", (
            Hotspot(0.50, 0.24, "Сверьте аудиторию", "Проверьте роли, площадки, группы и количество."),
            Hotspot(0.50, 0.52, "Сверьте сообщение", "Проверьте текст, ссылку и изображение."),
            Hotspot(0.50, 0.80, "Отправьте один раз", "Дождитесь итоговой статистики."),
        ), mobile=True, crop=(0, 0.18, 0, 0.39), note="После отправки текст не редактируется; исправление отправляется новой рассылкой."),
    )


def _management_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Управление", "Рабочие вкладки филиала", "Управление", "management-summary-desktop.png", (
            Hotspot(0.31, 0.42, "Выберите вкладку", "Сотрудники, товары, склады, импорт, связи и история."),
            Hotspot(0.30, 0.68, "Проверьте сводку", "Показатели относятся к текущему городу."),
            Hotspot(0.84, 0.42, "Обновите", "Кнопка заново загружает данные вкладки."),
        ), note="Лишних устаревших вкладок и дублирующего приветственного блока нет."),
        VisualSlide("Сотрудники", "Единые карточки активных сотрудников", "Управление -> Сотрудники", "staff-list-desktop.png", (
            Hotspot(0.45, 0.51, "Найдите сотрудника", "Поиск работает по имени, MAX ID и активной роли."),
            Hotspot(0.39, 0.62, "Сверьте роли", "Все активные роли человека указаны в одной карточке через запятую."),
            Hotspot(0.96, 0.62, "Откройте действия", "Меню позволяет изменить ФИО, добавить или отозвать роль."),
        ), note="Отозванные роли не отображаются; сотрудник исчезает после отзыва последней активной роли."),
        VisualSlide("Сотрудники", "Действия с единым профилем", "Сотрудники -> меню с тремя точками", "staff-actions-desktop.png", (
            Hotspot(0.96, 0.59, "Откройте меню", "Кнопка относится ко всему профилю сотрудника."),
            Hotspot(0.89, 0.67, "Выберите действие", "Измените ФИО или добавьте еще одну роль."),
            Hotspot(0.89, 0.79, "Отзовите конкретную роль", "Остальные активные роли сохраняются."),
        ), note="Директор может отозвать административную роль в назначенном ему городе."),
        VisualSlide("Сотрудники", "Добавление дополнительной роли", "Сотрудники -> Добавить роль", "staff-add-role-desktop.png", (
            Hotspot(0.77, 0.71, "Выберите роль", "Список ограничен полномочиями текущего пользователя."),
            Hotspot(0.88, 0.71, "Добавьте", "Новая роль прикрепляется к существующему MAX ID."),
            Hotspot(0.31, 0.62, "Проверьте карточку", "После обновления роли будут перечислены через запятую."),
        ), note="Повторная карточка сотрудника при добавлении роли не создается."),
        VisualSlide("Сотрудники", "Изменение ФИО сотрудника", "Сотрудники -> Изменить ФИО", "staff-edit-name-desktop.png", (
            Hotspot(0.58, 0.74, "Введите фамилию", "Новое значение применяется ко всему профилю."),
            Hotspot(0.73, 0.74, "Введите имя", "MAX ID остается прежним."),
            Hotspot(0.87, 0.74, "Сохраните", "Все активные роли продолжат работать."),
        ), note="После сохранения обновите список и найдите сотрудника по MAX ID."),
        VisualSlide("Сотрудники", "Создание одноразовой ссылки", "Управление -> Сотрудники -> Пригласить", "staff-invite-mobile.png", (
            Hotspot(0.50, 0.38, "Выберите роль", "Список ограничен правами приглашающего."),
            Hotspot(0.50, 0.52, "Создайте ссылку", "Она действует семь дней и один раз."),
            Hotspot(0.50, 0.72, "Закройте", "Без создания никаких изменений нет."),
        ), mobile=True, crop=(0, 0, 0, 0.30), note="Администратор не может приглашать роль выше собственных полномочий."),
        VisualSlide("Сотрудники", "Копирование приглашения", "Сотрудники -> созданная ссылка", "staff-invite-link-mobile.png", (
            Hotspot(0.50, 0.42, "Скопируйте", "Передайте ссылку конкретному сотруднику."),
            Hotspot(0.50, 0.58, "Сверьте срок", "После истечения нужна новая ссылка."),
            Hotspot(0.50, 0.76, "Проверьте роль", "После перехода сотрудник появится в активных."),
        ), mobile=True, crop=(0, 0, 0, 0.26), note="Повторный переход по использованной ссылке не создает новую связь."),
        VisualSlide("Товары", "Каталог и остатки", "Управление -> Товары и остатки", "products-desktop.png", (
            Hotspot(0.33, 0.34, "Фильтруйте", "Используйте статус, категорию и поиск."),
            Hotspot(0.78, 0.34, "Добавьте товар", "Откроется единая карточка товара."),
            Hotspot(0.90, 0.62, "Измените", "Редактируйте цену, тип выдачи и остатки."),
        ), note="SKU вручную не вводится: технический идентификатор формируется автоматически."),
        VisualSlide("Товары", "Фильтр каталога по складу", "Товары и остатки -> Все склады", "product-warehouse-filter-desktop.png", (
            Hotspot(0.87, 0.50, "Откройте склады", "Фильтр расположен рядом со статусом и категорией."),
            Hotspot(0.87, 0.57, "Выберите склад", "Список сузится без изменения остатков."),
            Hotspot(0.95, 0.51, "Сверьте счетчик", "Число показывает найденные товары."),
        ), note="В строках результата должно быть название выбранного склада."),
        VisualSlide("Товары", "Результат фильтра по складу", "Товары и остатки -> выбранный склад", "product-filtered-by-warehouse-desktop.png", (
            Hotspot(0.87, 0.50, "Сверьте выбранный склад", "Поле хранит активный фильтр."),
            Hotspot(0.42, 0.68, "Проверьте товар", "Метаданные подтверждают склад и количество."),
            Hotspot(0.89, 0.68, "Откройте карточку", "Склады и остатки проверяются через «Редактировать»."),
        ), note="Выбор «Все склады» возвращает полный каталог города."),
        VisualSlide("Товары", "Карточка физического товара", "Товары и остатки -> Добавить", "product-editor-desktop.png", (
            Hotspot(0.34, 0.30, "Заполните карточку", "Название, категория, цена, описание и фото."),
            Hotspot(0.70, 0.40, "Выберите выдачу", "Для обычного товара оставьте выдачу со склада."),
            Hotspot(0.70, 0.68, "Укажите остатки", "Количество задается отдельно для каждого склада."),
        ), note="После сохранения проверьте карточку товара в магазине."),
        VisualSlide("Товары", "Цифровой товар и пул кодов", "Товары -> тип «Код сразу после покупки»", "digital-product-editor-desktop.png", (
            Hotspot(0.34, 0.34, "Выберите цифровой тип", "Физический склад для него не нужен."),
            Hotspot(0.70, 0.48, "Добавьте коды", "Каждая строка — отдельный код выдачи."),
            Hotspot(0.70, 0.72, "Сверьте доступность", "Активных кодов должно хватать на продажи."),
        ), note="Код выдается сразу после покупки и исключается из доступного пула."),
        VisualSlide("Товары", "Массовый импорт товаров", "Товары и остатки -> Массовый импорт", "product-import-desktop.png", (
            Hotspot(0.29, 0.65, "Скачайте шаблон", "Не переименовывайте обязательные колонки."),
            Hotspot(0.61, 0.65, "Выберите файл", "Поддерживается подготовленный XLSX или CSV."),
            Hotspot(0.88, 0.65, "Импортируйте", "Дождитесь отчета по строкам."),
        ), note="Фото принимается по прямой ссылке, из Яндекс Диска или Google Drive при открытом доступе."),
        VisualSlide("Товары", "Подтверждение удаления товара", "Товары и остатки -> корзина", "product-delete-confirmation-desktop.png", (
            Hotspot(0.50, 0.46, "Сверьте название", "Удаляется именно указанный товар."),
            Hotspot(0.60, 0.58, "Прочитайте последствия", "Остатки удалятся во всех подключенных городах."),
            Hotspot(0.60, 0.66, "Подтвердите", "«Удалить товар» нельзя отменить."),
        ), note="Если у товара есть защищенная история движения, используйте скрытие вместо удаления."),
        VisualSlide("Склады", "Список складов и основной склад", "Управление -> Склады", "warehouses-desktop.png", (
            Hotspot(0.32, 0.34, "Выберите основной", "Он подставляется в новые складские операции."),
            Hotspot(0.78, 0.34, "Добавьте склад", "Укажите понятное название и адрес."),
            Hotspot(0.88, 0.62, "Измените склад", "Старые заказы при этом не переносятся."),
        ), note="Основной склад — персональная настройка сотрудника, а не города целиком."),
        VisualSlide("Склады", "Создание склада на телефоне", "Склады -> Добавить", "warehouse-editor-mobile.png", (
            Hotspot(0.50, 0.36, "Введите название", "Используйте название, понятное при сборке."),
            Hotspot(0.50, 0.54, "Укажите адрес", "Добавьте ориентир или примечание."),
            Hotspot(0.50, 0.78, "Сохраните", "Склад появится в товарах и заказах."),
        ), mobile=True, note="Код и технический тип склада пользователю заполнять не требуется."),
        VisualSlide("Склады", "Общий склад в нескольких городах", "Склады -> склад с совпадающим названием", "warehouse-connected-desktop.png", (
            Hotspot(0.37, 0.78, "Найдите общий склад", "Единая сущность отмечена статусом «Подключен»."),
            Hotspot(0.87, 0.78, "Проверьте действия", "Связь можно отключить от города без удаления общего склада."),
            Hotspot(0.89, 0.43, "Добавьте по названию", "Точное совпадение подключает уже существующий склад."),
        ), note="Товары и остатки общего склада доступны во всех подключенных городах."),
        VisualSlide("Импорт данных", "Проверка таблицы импорта", "Управление -> Импорт данных", "crm-import-desktop.png", (
            Hotspot(0.30, 0.38, "Скачайте шаблон", "Используйте актуальную структуру колонок."),
            Hotspot(0.52, 0.52, "Выберите лист", "Поддерживаются «Шаблон» и «Сделки»."),
            Hotspot(0.82, 0.65, "Проверьте файл", "Предпросмотр не изменяет базу."),
        ), note="Город берется из строки и сопоставляется только с доступными директору городами."),
        VisualSlide("Связи", "Родители, ученики и MAX-аккаунты", "Управление -> Связи", "contacts-desktop.png", (
            Hotspot(0.32, 0.34, "Фильтруйте связи", "Проверяйте роль и текущее состояние."),
            Hotspot(0.54, 0.48, "Найдите пользователя", "Поиск принимает имя и MAX ID."),
            Hotspot(0.91, 0.62, "Откройте действия", "Отозванную связь можно оформить заново с другой ролью."),
        ), note="Удаление связи не удаляет ученика, баланс, заказы или историю."),
        VisualSlide("История", "Аудит действий сотрудников", "Управление -> История", "admin-history-desktop.png", (
            Hotspot(0.30, 0.34, "Выберите период", "Ограничьте историю нужным интервалом."),
            Hotspot(0.58, 0.46, "Найдите событие", "Поиск работает по сотруднику и действию."),
            Hotspot(0.86, 0.62, "Сверьте детали", "В записи есть автор, время и объект изменения."),
        ), note="Устаревший внешний раздел удален; здесь только действия внутри текущей системы."),
    )


def _qr_help_slides() -> tuple[VisualSlide, ...]:
    return (
        VisualSlide("Родитель", "QR-код ребенка и ссылка", "Главная -> QR-коды детей", "parent-qr-mobile.png", (
            Hotspot(0.50, 0.22, "Выберите ребенка", "Каждый код относится к одному профилю."),
            Hotspot(0.50, 0.52, "Сохраните QR", "Изображение можно передать ребенку."),
            Hotspot(0.50, 0.72, "Скопируйте ссылку", "Она повторяет данные внутри QR-кода."),
        ), mobile=True, crop=(0, 0.16, 0, 0.46), note="Не открывайте детскую ссылку в родительском MAX-аккаунте."),
        VisualSlide("Родитель", "Полноэкранный QR-код", "QR-код ребенка -> Увеличить", "parent-qr-preview-mobile.png", (
            Hotspot(0.50, 0.48, "Покажите код", "Ребенок сканирует его со второго устройства."),
            Hotspot(0.50, 0.78, "Увеличьте яркость", "Полноэкранный режим упрощает сканирование."),
            Hotspot(0.90, 0.08, "Закройте", "Возвращает к списку детей."),
        ), mobile=True, note="После привязки повторно используйте обычный вход через MAX."),
        VisualSlide("Преподаватель", "QR-коды учеников группы", "Главная -> QR-коды детей", "teacher-qr-mobile.png", (
            Hotspot(0.50, 0.24, "Выберите группу", "Показаны только ученики преподавателя."),
            Hotspot(0.50, 0.48, "Откройте ученика", "Карточка показывает доступность QR."),
            Hotspot(0.50, 0.76, "Покажите код", "Код создается при активной родительской связи."),
        ), mobile=True, crop=(0, 0.14, 0, 0.36), note="Если родитель не подключен, экран объяснит перейти по ссылке из письма."),
        VisualSlide("Помощь", "Раздел помощи на компьютере", "Помощь", "help-desktop.png", (
            Hotspot(0.09, 0.58, "Откройте помощь", "Пункт находится в основном меню роли."),
            Hotspot(0.52, 0.40, "Скоро", "Встроенная справка пока недоступна."),
        ), note="До запуска встроенной справки используйте PDF и текстовый каталог действий."),
        VisualSlide("Помощь", "Раздел помощи на телефоне", "Еще -> Помощь", "help-mobile.png", (
            Hotspot(0.50, 0.48, "Проверьте страницу", "Она должна показывать только «Скоро...»."),
            Hotspot(0.50, 0.92, "Вернитесь в меню", "Откройте другой раздел через нижнюю панель."),
        ), mobile=True, note="Пустого отступа или зависания навигации после открытия помощи быть не должно."),
    )


def add_order_workflow_steps(guide: Any, prs: Presentation) -> None:
    guide.add_steps_slide(
        prs,
        kicker="Заказы",
        title="Как работает чек-лист сборки",
        steps=[
            ("Откройте склад", "Собирайте один источник товара за раз."),
            ("Возьмите количество", "Одинаковые позиции уже объединены в одну строку."),
            ("Поставьте галочку", "Отметка остается на месте и не сворачивает блок."),
            ("Проверьте весь набор", "Не подтверждайте заказ при фактической нехватке."),
            ("Нажмите обновить", "Только эта кнопка переводит полностью собранные заказы дальше."),
        ],
        side_title="Что не попадает в сборку",
        side_items=["Заказы без выбранного склада.", "Уже доставленные на площадку.", "Полученные преподавателем.", "Переданные ученику или отмененные."],
        callout_title="Галочка — не статус",
        callout_text="Это промежуточный чек-лист. Статус меняется только после подтверждения.",
        callout_kind="yellow",
    )
    guide.add_steps_slide(
        prs,
        kicker="Заказы",
        title="Финальная выдача и названия статусов",
        steps=[
            ("Доставка на площадку", "Администратор, директор или куратор подтверждает факт."),
            ("Получение учителем", "Преподаватель нажимает «Учитель получил»."),
            ("Проверка состава", "Перед выдачей сверьте ученика, товар и количество."),
            ("Передача ребенку", "Нажмите «Передать ученику» после фактической выдачи."),
            ("Проверка семьи", "У семьи заказ появится во вкладке «Получены»."),
        ],
        side_title="Кто может отменить",
        side_items=["Администратор города.", "Директор города.", "Не преподаватель и не куратор.", "Не ученик и не родитель."],
        callout_title="Два названия одного финала",
        callout_text="Семья видит «Получен», сотрудники — «Передан ученику».",
        callout_kind="teal",
    )


def add_daily_checklists(guide: Any, prs: Presentation) -> None:
    guide.add_cards_slide(
        prs,
        kicker="Ежедневная работа",
        title="Короткий маршрут каждой роли",
        cards=[
            ("Ученик / родитель", "Проверить профиль -> выбрать товар -> оформить -> следить за статусом.", "yellow"),
            ("Преподаватель", "Выбрать группу -> начислить AC -> получить заказ -> передать ребенку.", "teal"),
            ("Куратор", "Проверить свои группы -> доставку -> подготовить адресную рассылку.", "purple"),
            ("Администратор", "Назначить склад -> собрать -> распределить -> проверить остатки.", "coral"),
            ("Директор", "Выбрать город -> проверить отчет, сотрудников, заказы и импорт.", "teal"),
        ],
        note_title="Каждое действие заканчивается проверкой",
        note_text="Дождитесь сообщения, обновите список и откройте историю связанного заказа или ученика.",
        note_kind="yellow",
    )
    guide.add_steps_slide(
        prs,
        kicker="Диагностика",
        title="Что делать, если результат не появился",
        steps=[
            ("Не нажимайте повторно", "Особенно при покупке, начислении и изменении остатка."),
            ("Нажмите обновить", "Используйте круговую стрелку в текущем разделе."),
            ("Сбросьте фильтры", "Выберите «Все», очистите поиск и проверьте группу."),
            ("Проверьте роль и город", "Пустые списки часто связаны с областью доступа."),
            ("Зафиксируйте контекст", "Снимок, время, роль, город и номер объекта."),
        ],
        side_title="Что сообщить поддержке",
        side_items=["Номер заказа или ID ученика.", "Точное время по Москве.", "Роль и выбранный город.", "Текст сообщения об ошибке."],
        callout_title="История важнее повторного клика",
        callout_text="Сначала убедитесь, что операция не сохранилась, и только затем повторяйте действие.",
        callout_kind="coral",
    )


def add_final_slide(guide: Any, prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide, guide.PURPLE_DARK)
    guide.add_rect(slide, 0.68, 0.68, 0.58, 0.58, fill=guide.YELLOW, line=None, radius=True)
    guide.add_text(slide, "A", 0.68, 0.70, 0.58, 0.50, size=22, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    guide.add_text(slide, "ALGO MAX", 1.48, 0.82, 2.20, 0.32, size=18, fill=guide.WHITE, bold=True)
    guide.add_text(slide, "Перед подтверждением", 0.72, 1.86, 5.92, 0.62, size=31, fill=guide.WHITE, bold=True)
    guide.add_bullets(
        slide,
        ["Проверьте роль и город.", "Сверьте ученика, заказ или товар.", "Не повторяйте финансовую кнопку.", "Обновите данные после результата.", "При ошибке сохраните контекст."],
        0.78, 2.96, 5.70, 3.10, size=15, gap=0.62, accent=guide.YELLOW, fill=guide.WHITE,
    )
    guide.add_rect(slide, 7.02, 1.02, 5.62, 5.56, fill=guide.WHITE, line=None)
    guide.add_text(slide, "Главный принцип", 7.48, 1.54, 4.70, 0.40, size=21, bold=True)
    guide.add_text(
        slide,
        "Статус в приложении должен соответствовать реальному действию: товар собран, доставлен, получен преподавателем или передан ученику.",
        7.48, 2.32, 4.58, 1.62, size=19, bold=True, fill=guide.PURPLE_DARK,
    )
    guide.add_callout(slide, "Материалы", "PPTX, PDF, текстовое руководство и отчет из 130 исправлений находятся в docs и output/pdf.", 7.48, 4.58, 4.58, kind="teal")
    guide.add_text(slide, f"{len(prs.slides):02d}", 12.02, 6.88, 0.62, 0.28, size=11, fill="CBB9E8", align=PP_ALIGN.RIGHT)


def build_visual_presentation(guide: Any) -> Presentation:
    prs = Presentation()
    prs.slide_width = guide.SLIDE_W
    prs.slide_height = guide.SLIDE_H
    add_cover(guide, prs)
    add_contents(guide, prs)
    add_how_to_read(guide, prs)
    add_role_reference(guide, prs)

    guide.add_section_slide(prs, section="Раздел 1", title="Вход, профиль и роли", subtitle="Сначала убедитесь, от чьего имени и в каком городе открыта система.", items=["Ученик и родитель", "Преподаватель и куратор", "Администратор", "Директор"])
    for spec in _profile_slides():
        add_visual_slide(guide, prs, spec)

    guide.add_section_slide(prs, section="Раздел 2", title="Магазин и покупка", subtitle="Поиск награды, проверка суммы и получение цифрового кода.", items=["Поиск и категории", "Сортировка и фильтры", "Корзина и подтверждение", "Цифровая автовыдача"], accent=guide.TEAL, background=guide.TEAL_LIGHT)
    for spec in _store_slides():
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="Покупка", title="Проверки и частые ошибки корзины",
        steps=[("Проверьте ребенка", "У родителя корзина относится к выбранному профилю."), ("Проверьте наличие", "Недоступный физический товар нельзя оформить."), ("Проверьте баланс", "Итог не должен превышать доступные AC."), ("Подтвердите один раз", "Дождитесь сообщения и появления заказа."), ("Не ищите отмену", "Ученик и родитель не отменяют оформленный заказ.")],
        side_title="После покупки", side_items=["Корзина очищена.", "Баланс обновлен.", "Заказ появился в списке.", "Цифровой код выдан сразу."], callout_title="Физический товар", callout_text="Сначала имеет статус «Зарезервирован» и ждет назначения склада.", callout_kind="yellow",
    )

    guide.add_section_slide(prs, section="Раздел 3", title="Заказы и выдача", subtitle="Каждая кнопка подтверждает фактический этап движения товара.", items=["Назначить склад", "Собрать", "Распределить", "Получить и передать ученику"], accent=guide.CORAL, background=guide.CORAL_LIGHT)
    specs = _order_slides()
    for spec in specs[:8]:
        add_visual_slide(guide, prs, spec)
    add_order_workflow_steps(guide, prs)
    for spec in specs[8:]:
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="Права", title="Кто меняет состояние заказа",
        steps=[("Склад", "Администратор или директор назначает склад."), ("Сборка", "Администратор или директор ведет чек-лист."), ("Доставка", "Администратор, директор или куратор подтверждает площадку."), ("Получение", "Преподаватель подтверждает получение комплекта."), ("Выдача", "Преподаватель или уполномоченный сотрудник передает ученику.")],
        side_title="Отмена", side_items=["Только администратор.", "Только директор.", "С указанием причины.", "С возвратом AC ученику."], callout_title="Семья не отменяет заказ", callout_text="Кнопка отмены отсутствует у ученика и родителя на всех этапах.", callout_kind="coral",
    )

    guide.add_section_slide(prs, section="Раздел 4", title="Ученики и астрокоины", subtitle="Карточка ученика объединяет связи, статус, баланс и полную историю.", items=["Реестр и карточка", "Создание и перевод", "Начисление AC", "Отчет и правила"])
    for spec in _student_ac_slides()[:4]:
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="Ученики", title="Статус, выбытие и заморозка доступа",
        steps=[("Измените статус", "В карточке ученика выберите состояние и укажите причину."), ("Запустите срок", "После выбытия действует установленное число дней доступа."), ("Настройте паузу", "Директор задает период, который не расходует оставшиеся дни."), ("Сохраните историю", "Баланс, заказы и связи не удаляются."), ("Восстановите", "Активируйте статус — прежний баланс продолжит действовать.")],
        side_title="Единый расчет срока", side_items=["Вход по письму и QR.", "Кабинет и магазин.", "Корзина и новый заказ.", "Старые заказы доступны сотрудникам."], callout_title="Летняя пауза", callout_text="Дни внутри периода заморозки не уменьшают остаток доступа.", callout_kind="yellow",
    )
    guide.add_steps_slide(
        prs, kicker="Ученики", title="Перевод в группу и подарок на день рождения",
        steps=[("Откройте карточку", "Найдите ученика через поиск или фильтр группы."), ("Выберите новую группу", "Сохраните перевод без создания дубликата."), ("Проверьте связи", "Родитель, MAX-доступ, баланс и заказы сохраняются."), ("Проверьте историю", "Старая и новая группы записаны с автором и временем."), ("Укажите день рождения", "В эту дату один раз начисляется сумма системного правила города.")],
        side_title="Доступ к истории", side_items=["Преподаватель ученика.", "Куратор группы.", "Администратор.", "Директор."], callout_title="Перевод — не новый ученик", callout_text="Все накопленные данные продолжают относиться к прежней карточке.", callout_kind="teal",
    )
    for spec in _student_ac_slides()[4:]:
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="Начисления", title="Правила AC для конкретного города",
        steps=[("Откройте правила", "Работайте в выбранном городе."), ("Свяжите причину и сумму", "Стандартная причина подставляет фиксированное количество AC."), ("Отключите лишнее", "Неактивное правило исчезает из формы."), ("Оставьте ручной вариант", "Для исключения используйте свою причину и сумму."), ("Проверьте отчет", "Операция содержит автора и примененную причину.")],
        side_title="Кому доступна настройка", side_items=["Администратору.", "Директору.", "Преподаватель применяет правила.", "Куратор применяет правила."], callout_title="Автоматический подарок", callout_text="Сумма правила начисляется в день рождения при заполненной дате. Причину нельзя удалить.", callout_kind="yellow",
    )

    guide.add_section_slide(prs, section="Раздел 5", title="Рассылки", subtitle="Сначала аудитория, затем содержание и финальная проверка.", items=["Получатели и форматы", "Площадки и группы", "Текст, фото и эмодзи", "Проверка и история"], accent=guide.TEAL, background=guide.TEAL_LIGHT)
    for spec in _broadcast_slides():
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="Рассылки", title="Отправка, история и повтор",
        steps=[("Проверьте аудиторию", "Сверьте город, категории, форматы, площадки и группы."), ("Проверьте контент", "Откройте ссылку и прочитайте сообщение на телефоне."), ("Отправьте один раз", "Дождитесь итоговой статистики."), ("Откройте историю", "Сверьте успешные, пропущенные и ошибочные доставки."), ("Повторите корректно", "При ошибке создайте новую рассылку из сохраненного содержания.")],
        side_title="Перед отправкой", side_items=["Нет данных чужой группы.", "Указаны дата и место.", "Ссылка открывается.", "Фото читается на телефоне."], callout_title="Отправленное не редактируется", callout_text="Исправления рассылаются новым сообщением, чтобы история оставалась достоверной.", callout_kind="coral",
    )

    guide.add_section_slide(prs, section="Раздел 6", title="Управление филиалом", subtitle="Сотрудники, товары, склады и импорт изолированы по рабочему городу.", items=["Сотрудники и приглашения", "Товары и цифровые коды", "Склады", "Импорт, связи и история"])
    management_specs = _management_slides()
    for index, spec in enumerate(management_specs):
        add_visual_slide(guide, prs, spec)
        if spec.asset == "staff-invite-link-mobile.png":
            guide.add_steps_slide(
                prs, kicker="Сотрудники", title="Дополнительные роли, ФИО и отзыв",
                steps=[("Найдите сотрудника", "Используйте поиск по имени, MAX ID или активной роли."), ("Откройте меню", "Кнопка с тремя точками находится справа в карточке."), ("Измените ФИО", "Новое имя применяется ко всем ролям профиля."), ("Добавьте роль", "Новая роль появляется в той же карточке через запятую."), ("Отзовите роль", "Выберите конкретную роль; остальные полномочия сохранятся.")],
                side_title="Что сохраняется", side_items=["Одна карточка сотрудника.", "История действий.", "Остальные активные роли.", "Единый MAX ID."], callout_title="Отозванные роли скрыты", callout_text="Они не отображаются в карточке и не дают доступ к рабочим разделам.", callout_kind="yellow",
            )
        if spec.asset == "crm-import-desktop.png":
            guide.add_steps_slide(
                prs, kicker="Импорт данных", title="Автоматическое распределение по городам",
                steps=[("Заполните город", "Значение берется из каждой строки файла."), ("Загрузите XLSX", "Поддерживаются листы «Шаблон» и «Сделки»."), ("Проверьте предпросмотр", "Неизвестный город показывает ошибку."), ("Подтвердите", "Только теперь записи создаются по филиалам."), ("Сверьте итог", "Откройте реестр каждого доступного города.")],
                side_title="Ограничения", side_items=["Только доступные директору города.", "Город уже создан.", "Группа создается в городе строки.", "Повтор обновляет по ID."], callout_title="Заранее база не распределяется", callout_text="Распределение выполняется только при подтвержденном импорте директором.", callout_kind="teal",
            )
    guide.add_section_slide(prs, section="Раздел 7", title="QR-коды и помощь", subtitle="Детский вход создается через подтвержденную связь с родителем.", items=["QR родителя", "QR преподавателя", "Повторная привязка", "Встроенная помощь"], accent=guide.CORAL, background=guide.CORAL_LIGHT)
    for spec in _qr_help_slides()[:3]:
        add_visual_slide(guide, prs, spec)
    guide.add_steps_slide(
        prs, kicker="QR и связи", title="Если QR недоступен или связь ошибочна",
        steps=[("Проверьте родителя", "Для QR нужна активная связь родитель–ученик."), ("Откройте письмо", "Если связи нет, родитель переходит по ссылке школы."), ("Отзовите ошибочную связь", "Администратор делает это во вкладке «Связи»."), ("Выдайте новую роль", "После отзыва связь можно оформить заново."), ("Повторите QR", "Покажите код только нужного ученика.")],
        side_title="Безопасная привязка", side_items=["Один QR — один ученик.", "Ссылка передается лично.", "Родитель не открывает детскую ссылку.", "История сохраняется."], callout_title="QR не заменяет родителя", callout_text="Он дает ребенку вход только после подтвержденной родительской связи.", callout_kind="coral",
    )
    for spec in _qr_help_slides()[3:]:
        add_visual_slide(guide, prs, spec)
    add_daily_checklists(guide, prs)
    add_final_slide(guide, prs)
    return prs
