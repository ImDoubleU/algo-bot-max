from __future__ import annotations

import json
import re
import subprocess
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

from PIL import Image, ImageStat
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import build_video_series as video
from interaction_scenarios import SCENARIOS


ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = ROOT / "docs" / "video_series_manifest.json"
REPORT_PATH = ROOT / "docs" / "QA_REPORT_RU.md"
JSON_PATH = ROOT / "output" / "training-materials-qa.json"
FFMPEG = ROOT / "output" / ".audit-tools" / "imageio_ffmpeg" / "binaries" / "ffmpeg-win-x86_64-v7.1.exe"

EXPECTED_IDS = (
    "promo",
    "detailed",
    "student",
    "parent",
    "teacher",
    "curator",
    "administrator",
    "director",
)
DURATION_RANGES = {
    "promo": (45, 60),
    "detailed": (360, 480),
    "student": (210, 255),
    "parent": (210, 245),
    "teacher": (270, 335),
    "curator": (165, 215),
    "administrator": (720, 850),
    "director": (600, 730),
}
ROLE_REQUIRED_ASSETS = {
    "student": {"student-dashboard-mobile.png", "store-sort-desktop.png", "checkout-dialog-mobile.png"},
    "parent": {"parent-dashboard-mobile.png", "parent-qr-mobile.png", "family-orders-mobile.png"},
    "teacher": {"teacher-profile-first-login-desktop.png", "teacher-profile-groups-desktop.png", "teacher-orders-teacher-mobile.png"},
    "curator": {"curator-dashboard-desktop.png", "curator-ac-report-desktop.png", "broadcast-review-mobile.png"},
    "administrator": {"staff-add-role-desktop.png", "product-warehouse-filter-desktop.png", "birthday-accrual-rule-desktop.png"},
    "director": {"director-city-switcher-desktop.png", "warehouse-connected-desktop.png", "staff-edit-name-desktop.png"},
}
FORBIDDEN = (
    "супер" + "админ",
    "0.78.1",
    "импорт crm",
    "подарок 50 ac",
    "ru_ru-dmitri",
)


@dataclass
class Check:
    pass_name: str
    item: str
    status: str
    detail: str


def check(results: list[Check], pass_name: str, item: str, condition: bool, detail: str) -> None:
    results.append(Check(pass_name, item, "PASS" if condition else "FAIL", detail))


def project_path(value: str) -> Path:
    return ROOT / Path(value.replace("\\", "/"))


def pptx_text(path: Path) -> str:
    prs = Presentation(path)
    return "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def probe_video(path: Path) -> dict[str, object]:
    result = subprocess.run(
        [str(FFMPEG), "-hide_banner", "-i", str(path)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    value = result.stderr
    duration = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", value)
    dimensions = re.search(r"Video:.*?(\d{3,4})x(\d{3,4}).*?(\d+(?:\.\d+)?)\s+fps", value)
    if not duration or not dimensions:
        raise RuntimeError(f"Не удалось прочитать видео {path.name}: {value[-2000:]}")
    hours, minutes, seconds = duration.groups()
    return {
        "duration": int(hours) * 3600 + int(minutes) * 60 + float(seconds),
        "width": int(dimensions.group(1)),
        "height": int(dimensions.group(2)),
        "fps": float(dimensions.group(3)),
        "audio_streams": len(re.findall(r"Stream #.*Audio:", value)),
        "subtitle_streams": len(re.findall(r"Stream #.*Subtitle:", value)),
    }


def decode_first_frame(path: Path) -> bool:
    result = subprocess.run(
        [str(FFMPEG), "-hide_banner", "-v", "error", "-i", str(path), "-frames:v", "1", "-f", "framehash", "-"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    return result.returncode == 0 and bool(re.search(r"^0,", result.stdout, flags=re.MULTILINE))


def srt_metrics(path: Path) -> tuple[int, bool]:
    text = path.read_text(encoding="utf-8")
    ranges = re.findall(
        r"(\d\d):(\d\d):(\d\d),(\d{3}) --> (\d\d):(\d\d):(\d\d),(\d{3})",
        text,
    )

    def seconds(parts: tuple[str, ...]) -> float:
        h, m, s, ms = map(int, parts)
        return h * 3600 + m * 60 + s + ms / 1000

    previous_end = 0.0
    ordered = True
    for raw in ranges:
        start, end = seconds(raw[:4]), seconds(raw[4:])
        ordered = ordered and start >= previous_end - 0.01 and end > start
        previous_end = end
    return len(ranges), ordered


def scan_forbidden(paths: list[Path]) -> list[str]:
    hits: list[str] = []
    for path in paths:
        if path.suffix.casefold() == ".pptx":
            content = pptx_text(path)
        else:
            content = path.read_text(encoding="utf-8")
        folded = content.casefold()
        for phrase in FORBIDDEN:
            if phrase in folded:
                hits.append(f"{path.name}: {phrase}")
    return hits


def audit() -> tuple[list[Check], dict[str, object]]:
    results: list[Check] = []
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    entries = list(manifest["videos"])
    ids = tuple(str(entry["id"]) for entry in entries)

    check(results, "Фактическая", "Состав серии", ids == EXPECTED_IDS, f"Видео: {', '.join(ids)}")
    check(results, "Фактическая", "Режим озвучки", manifest.get("voice") == "manual", "В манифесте указан живой голос; синтетическая дорожка не создается.")
    check(results, "Фактическая", "Сценарии", len(SCENARIOS) >= 40 and len({item.id for item in SCENARIOS}) == len(SCENARIOS), f"Проверено сценариев: {len(SCENARIOS)}")

    app_sources = {
        "Привязка LMS": (ROOT / "app/web/static/miniapp/app.js", "Сохранить и найти группы"),
        "Дополнительная роль": (ROOT / "app/web/static/miniapp/app-admin.js", "Добавить роль"),
        "Фильтр склада": (ROOT / "app/web/static/miniapp/app-admin.js", "Склад товара"),
        "Фото из облака": (ROOT / "app/web/static/miniapp/app-admin.js", "Google Drive"),
        "Защита дня рождения": (ROOT / "app/web/static/miniapp/app.js", "Причина «С днем рождения» обязательна"),
    }
    for item, (path, term) in app_sources.items():
        content = path.read_text(encoding="utf-8")
        check(results, "Фактическая", item, term in content, f"Текст интерфейса подтвержден в {path.name}.")

    source_docs = [
        ROOT / "docs/INTERACTION_GUIDE_RU.md",
        ROOT / "docs/VIDEO_SERIES_RU.md",
        ROOT / "docs/SCREENSHOT_REGISTER_RU.md",
        ROOT / "docs/Algo_MAX_Руководство_по_ролям.pptx",
        ROOT / "docs/Algo_MAX_Презентация_продукта.pptx",
    ]
    source_docs.extend(project_path(str(entry["presentation"])) for entry in entries)
    source_docs.extend(project_path(str(entry["script"])) for entry in entries)
    forbidden_hits = scan_forbidden(source_docs)
    check(results, "Фактическая", "Публичные роли и версии", not forbidden_hits, "Запрещенные упоминания не найдены." if not forbidden_hits else "; ".join(forbidden_hits))

    broken_links: list[str] = []
    navigation_docs = [ROOT / "docs/VIDEO_SERIES_RU.md", ROOT / "output/video-series/README.md"]
    for path in navigation_docs:
        content = path.read_text(encoding="utf-8")
        for target in re.findall(r"\[[^]]+\]\(([^)]+)\)", content):
            if re.match(r"^(?:https?://|mailto:|#)", target, flags=re.IGNORECASE):
                continue
            resolved = (path.parent / target.replace("\\", "/")).resolve()
            if not resolved.exists():
                broken_links.append(f"{path.name}: {target}")
    check(
        results,
        "Практическая",
        "Навигация по комплекту",
        not broken_links,
        "Все ссылки на PPTX, PDF, MP4, сценарии и субтитры открываются." if not broken_links else "; ".join(broken_links),
    )

    complete_scenarios = all(
        item.role and item.before and item.path and len(item.steps) >= 3 and item.result and item.verification and item.correction and item.limitations
        for item in SCENARIOS
    )
    check(results, "Практическая", "Полнота процедур", complete_scenarios, "У каждой процедуры есть условия, путь, шаги, результат, проверка, исправление и ограничения.")
    quote_steps = sum("«" in description or "»" in description for item in SCENARIOS for _, description in item.steps)
    total_steps = sum(len(item.steps) for item in SCENARIOS)
    check(results, "Практическая", "Названия кнопок", quote_steps >= total_steps * 0.45, f"Точные экранные подписи используются в {quote_steps} из {total_steps} шагов.")

    specs = video.all_visual_specs()
    missing_assets: list[str] = []
    blank_assets: list[str] = []
    for asset, spec in specs.items():
        try:
            path = video.visual._asset_path(video.guide, asset)
        except FileNotFoundError:
            missing_assets.append(asset)
            continue
        with Image.open(path) as image:
            variance = sum(ImageStat.Stat(image.convert("RGB").resize((64, 64))).var)
            if image.width < 360 or image.height < 360 or variance < 20:
                blank_assets.append(asset)
    check(results, "Практическая", "Скриншоты", not missing_assets and not blank_assets, f"Доступно и непусто: {len(specs)}; пропуски: {missing_assets}; проблемные: {blank_assets}")

    role_by_id = {role.video_id: role for role in video.ROLE_VIDEOS}
    coverage_errors: list[str] = []
    for video_id, required in ROLE_REQUIRED_ASSETS.items():
        missing = required - set(role_by_id[video_id].assets)
        if missing:
            coverage_errors.append(f"{video_id}: {sorted(missing)}")
    check(results, "Практическая", "Обязательный охват ролей", not coverage_errors, "Все новые ключевые сценарии включены." if not coverage_errors else "; ".join(coverage_errors))

    script_errors: list[str] = []
    for entry in entries:
        narration = json.loads(project_path(str(entry["narration"])).read_text(encoding="utf-8"))
        script = project_path(str(entry["script"])).read_text(encoding="utf-8")
        if "| Таймкод | Экран | Действие курсора | Текст диктора | Текст на экране |" not in script:
            script_errors.append(f"{entry['id']}: нет покадровой таблицы")
        if len(narration) != int(entry["slides"]):
            script_errors.append(f"{entry['id']}: слайды и текст не совпадают")
        if any(not item.get("cursor_action") or not item.get("screen_text") or not item.get("duration_seconds") for item in narration):
            script_errors.append(f"{entry['id']}: неполные данные кадра")
    check(results, "Практическая", "Сценарии и покадровые планы", not script_errors, "Для восьми роликов готовы таймкоды, курсор, диктор и экранный текст." if not script_errors else "; ".join(script_errors))

    presentation_paths = [ROOT / "docs/Algo_MAX_Руководство_по_ролям.pptx", ROOT / "docs/Algo_MAX_Презентация_продукта.pptx"]
    presentation_paths.extend(project_path(str(entry["presentation"])) for entry in entries)
    visual_errors: list[str] = []
    presentation_metrics: dict[str, dict[str, int | float]] = {}
    for path in presentation_paths:
        prs = Presentation(path)
        pictures = 0
        min_font = 1000.0
        for slide_number, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.left < -20_000 or shape.top < -20_000 or shape.left + shape.width > prs.slide_width + 20_000 or shape.top + shape.height > prs.slide_height + 20_000:
                    visual_errors.append(f"{path.name}, слайд {slide_number}: объект за границей")
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures += 1
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip() and run.font.size:
                                min_font = min(min_font, run.font.size.pt)
        if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.001:
            visual_errors.append(f"{path.name}: формат не 16:9")
        if min_font < 7.5:
            visual_errors.append(f"{path.name}: минимальный шрифт {min_font:.1f} pt")
        presentation_metrics[path.name] = {"slides": len(prs.slides), "pictures": pictures, "minimum_font_pt": round(min_font, 1)}
    check(results, "Визуальная", "Геометрия PPTX", not visual_errors, "Формат 16:9, объекты внутри слайдов, минимальный шрифт не ниже 7,5 pt." if not visual_errors else "; ".join(visual_errors[:20]))

    video_metrics: dict[str, dict[str, object]] = {}
    media_errors: list[str] = []
    for entry in entries:
        video_id = str(entry["id"])
        video_path = project_path(str(entry["video"]))
        pdf_path = project_path(str(entry["pdf"]))
        if not video_path.exists() or not pdf_path.exists():
            media_errors.append(f"{video_id}: нет MP4 или PDF")
            continue
        probe = probe_video(video_path)
        lower, upper = DURATION_RANGES[video_id]
        cues, ordered = srt_metrics(project_path(str(entry["subtitles_srt"])))
        pdf_pages = len(PdfReader(pdf_path).pages)
        ok = (
            lower <= float(probe["duration"]) <= upper
            and probe["width"] == 1920
            and probe["height"] == 1080
            and abs(float(probe["fps"]) - 25) <= 0.1
            and probe["audio_streams"] == 0
            and probe["subtitle_streams"] >= 1
            and ordered
            and cues >= int(entry["slides"])
            and pdf_pages == int(entry["slides"])
            and decode_first_frame(video_path)
        )
        if not ok:
            media_errors.append(f"{video_id}: {probe}, cues={cues}, pdf={pdf_pages}")
        video_metrics[video_id] = {**probe, "slides": int(entry["slides"]), "srt_cues": cues, "pdf_pages": pdf_pages}
    check(results, "Визуальная", "PDF, MP4 и субтитры", not media_errors, "Все PDF совпадают со слайдами; MP4 1080p/25 fps, без аудио, с субтитрами и декодируемым кадром." if not media_errors else "; ".join(media_errors))


    results.append(Check("Практическая", "Рабочая база и внешняя доставка", "NOT VERIFIED", "По условию рабочая база не изменялась. Реальная отправка в MAX, запись операций и производственный вход не выполнялись; интерфейсные маршруты проверены в безопасном локальном режиме."))

    summary = {
        "date": date.today().isoformat(),
        "ui_version": "0.86.4",
        "scenarios": len(SCENARIOS),
        "screenshots": len(specs),
        "presentations": presentation_metrics,
        "videos": video_metrics,
    }
    return results, summary


def write_report(results: list[Check], summary: dict[str, object]) -> None:
    failures = [item for item in results if item.status == "FAIL"]
    lines = [
        "# QA-отчет по обучающим материалам Algo MAX",
        "",
        f"Дата проверки: {summary['date']}. Интерфейс: {summary['ui_version']}.",
        "",
        f"Итог: **{'FAIL' if failures else 'PASS'}**. Проверено сценариев: {summary['scenarios']}; скриншотов: {summary['screenshots']}.",
        "",
        "## Три прохода",
        "",
        "| Проход | Проверка | Статус | Результат |",
        "|---|---|---|---|",
    ]
    for item in results:
        lines.append(f"| {item.pass_name} | {item.item} | **{item.status}** | {item.detail} |")

    lines.extend(["", "## Итоговые файлы", "", "| Видео | Слайды | Длительность | PDF | Звук |", "|---|---:|---:|---:|---|"])
    for video_id in EXPECTED_IDS:
        metrics = summary["videos"].get(video_id, {})
        if not metrics:
            lines.append(f"| {video_id} | — | — | — | не проверено |")
            continue
        duration = float(metrics["duration"])
        lines.append(
            f"| {video_id} | {metrics['slides']} | {int(duration) // 60}:{int(duration) % 60:02d} | "
            f"{metrics['pdf_pages']} стр. | нет, мастер для живой озвучки |"
        )

    lines.extend(
        [
            "",
            "## Оставшиеся ограничения",
            "",
            "- Реальные финансовые, складские и статусные изменения не выполнялись, чтобы не менять рабочие данные.",
            "- Доставка сообщений через MAX и вход по настоящим приглашениям требуют отдельной приемки на тестовых аккаунтах.",
            "- Встроенный раздел «Помощь» пока показывает «Скоро...», поэтому основными материалами остаются PDF, PPTX и текстовое руководство.",
            "- VPS не изменялся и развертывание не выполнялось.",
            "",
        ]
    )
    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    results, summary = audit()
    write_report(results, summary)
    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    JSON_PATH.write_text(
        json.dumps({"checks": [asdict(item) for item in results], "summary": summary}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    failures = [item for item in results if item.status == "FAIL"]
    print(f"QA checks: {len(results)}, failures: {len(failures)}")
    print(REPORT_PATH)
    if failures:
        raise SystemExit("\n".join(f"{item.pass_name}: {item.item}: {item.detail}" for item in failures))


if __name__ == "__main__":
    main()
