# ruff: noqa: E501 - Narration copy stays beside each video definition for review.

from __future__ import annotations

import json
import shutil
from collections.abc import Iterable
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import date
from pathlib import Path

import build_product_presentation as product
import build_role_guide as guide
import build_visual_guide as visual
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT = Path(__file__).resolve().parents[1]
SERIES_DIR = ROOT / "docs" / "video-series"
MANIFEST_PATH = ROOT / "docs" / "video_series_manifest.json"
OVERVIEW_PATH = ROOT / "docs" / "VIDEO_SERIES_RU.md"
VIDEO_OUTPUT_DIR = ROOT / "output" / "video-series"


@contextmanager
def footer_style(text: str):
    previous_footer = guide.add_footer

    def add_series_footer(slide, _text: str = text) -> None:
        product._guide_footer(slide, text)

    guide.add_footer = add_series_footer
    try:
        yield
    finally:
        guide.add_footer = previous_footer


@dataclass(frozen=True)
class RoleVideo:
    order: int
    video_id: str
    role: str
    role_genitive: str
    title: str
    subtitle: str
    cover_asset: str
    capabilities: tuple[str, ...]
    assets: tuple[str, ...]
    intro: str
    outro: str
    target_duration: str


ROLE_VIDEOS: tuple[RoleVideo, ...] = (
    RoleVideo(
        3,
        "student",
        "Ученик",
        "ученика",
        "Как пользоваться кабинетом ученика",
        "Баланс, магазин, корзина, заказы и цифровые награды",
        "student-dashboard-mobile.png",
        (
            "проверять баланс и историю астрокоинов;",
            "находить награды и добавлять их в корзину;",
            "оформлять заказ и следить за его состоянием;",
            "получать и копировать код цифрового товара.",
        ),
        (
            "student-dashboard-mobile.png",
            "store-mobile.png",
            "store-sort-desktop.png",
            "product-dialog-mobile.png",
            "cart-mobile.png",
            "checkout-dialog-mobile.png",
            "family-orders-mobile.png",
            "family-order-dialog-mobile.png",
            "visual/order-digital-mobile.png",
        ),
        "Сейчас покажу кабинет ученика. Проверим баланс, найдем награду, оформим заказ и посмотрим, где получить цифровой код.",
        "Мы прошли весь путь ученика: от баланса до полученного заказа. Ученик не отменяет созданный заказ сам; это делает администратор или директор.",
        "3:30–4:15",
    ),
    RoleVideo(
        4,
        "parent",
        "Родитель",
        "родителя",
        "Как управлять профилями детей",
        "Переключение детей, QR-доступ, покупки и семейные заказы",
        "parent-dashboard-mobile.png",
        (
            "переключаться между связанными детьми;",
            "проверять отдельный баланс и корзину каждого ребенка;",
            "передавать ребенку QR-код или персональную ссылку;",
            "оформлять и отслеживать заказ выбранного ребенка.",
        ),
        (
            "parent-dashboard-mobile.png",
            "parent-qr-mobile.png",
            "parent-qr-preview-mobile.png",
            "store-mobile.png",
            "cart-mobile.png",
            "checkout-dialog-mobile.png",
            "family-orders-mobile.png",
            "family-order-dialog-mobile.png",
        ),
        "Сейчас покажу кабинет родителя. Сначала выберем ребенка, затем откроем QR-доступ, оформим покупку и проверим заказ.",
        "Главная проверка для родителя — имя выбранного ребенка. Вместе с ним переключаются баланс, корзина, история и заказы. Детскую ссылку не открываем в родительском аккаунте MAX.",
        "3:30–4:00",
    ),
    RoleVideo(
        5,
        "teacher",
        "Преподаватель",
        "преподавателя",
        "Ежедневная работа преподавателя",
        "Свои ученики, начисления, QR-коды, прием и выдача заказов",
        "teacher-dashboard-mobile.png",
        (
            "работать только со своими группами и учениками;",
            "начислять астрокоины по правилу или своей причине;",
            "показывать ученикам персональные QR-коды;",
            "принимать комплект и отмечать фактическую выдачу.",
        ),
        (
            "teacher-profile-first-login-desktop.png",
            "teacher-profile-groups-desktop.png",
            "teacher-dashboard-mobile.png",
            "teacher-student-registry-desktop.png",
            "accrual-mobile.png",
            "accrual-custom-mobile.png",
            "teacher-qr-mobile.png",
            "teacher-orders-mobile.png",
            "teacher-orders-venue-mobile.png",
            "teacher-orders-teacher-mobile.png",
            "teacher-order-issued-history-mobile.png",
        ),
        "Сейчас покажу первый вход преподавателя и ежедневную работу. Введем ФИО как в LMS, проверим найденные группы, начислим AC и пройдем выдачу заказа.",
        "Группы преподавателя находятся по ФИО из LMS. Заказ меняем только после реального приема или выдачи. Отменять заказ преподаватель не может.",
        "4:30–5:30",
    ),
    RoleVideo(
        6,
        "curator",
        "Куратор",
        "куратора",
        "Работа куратора с учениками и новостями",
        "Группы, начисления, контроль заказов и адресные рассылки",
        "curator-dashboard-desktop.png",
        (
            "работать с разрешенными группами и учениками;",
            "начислять астрокоины выбранным ученикам;",
            "контролировать связанные заказы без права отмены;",
            "собирать аудиторию и отправлять адресные новости.",
        ),
        (
            "curator-dashboard-desktop.png",
            "curator-accrual-desktop.png",
            "curator-ac-report-desktop.png",
            "curator-broadcast-targets-desktop.png",
            "broadcast-news-mobile.png",
            "broadcast-news-emoji-mobile.png",
            "broadcast-review-mobile.png",
        ),
        "Сейчас покажу работу куратора: доступные группы, начисления, контроль заказов и рассылку по точной аудитории.",
        "Перед отправкой куратор всегда проверяет группы и число получателей. Заказы он контролирует, но отмена доступна только администратору и директору.",
        "2:30–3:15",
    ),
    RoleVideo(
        7,
        "administrator",
        "Администратор",
        "администратора",
        "Операционная работа администратора",
        "Ученики, AC, заказы, товары, склады, сотрудники, импорт и рассылки",
        "admin-dashboard-desktop.png",
        (
            "вести карточки учеников, связи, начисления и отчет AC;",
            "назначать склад, собирать и распределять заказы;",
            "управлять физическими и цифровыми товарами;",
            "работать со складами, импортом, сотрудниками, рассылками и историей.",
        ),
        (
            "combined-role-profile-desktop.png",
            "student-registry-desktop.png",
            "student-card-desktop.png",
            "student-create-desktop.png",
            "orders-assign-desktop.png",
            "order-warehouse-dialog-desktop.png",
            "orders-collect-desktop.png",
            "orders-route-desktop.png",
            "order-cancel-form-desktop.png",
            "products-desktop.png",
            "product-warehouse-filter-desktop.png",
            "product-filtered-by-warehouse-desktop.png",
            "product-editor-desktop.png",
            "digital-product-editor-desktop.png",
            "product-import-desktop.png",
            "product-delete-confirmation-desktop.png",
            "warehouses-desktop.png",
            "warehouse-connected-desktop.png",
            "crm-import-desktop.png",
            "contacts-desktop.png",
            "admin-staff-invite-desktop.png",
            "staff-actions-desktop.png",
            "staff-add-role-desktop.png",
            "staff-edit-name-desktop.png",
            "admin-history-desktop.png",
            "admin-accrual-desktop.png",
            "birthday-accrual-rule-desktop.png",
            "admin-ac-report-desktop.png",
            "admin-broadcast-targets-desktop.png",
            "broadcast-news-mobile.png",
            "broadcast-review-mobile.png",
        ),
        "Сейчас покажу ежедневную работу администратора: ученики, AC, заказы, сотрудники, товары, склады, импорт, связи, рассылки и история.",
        "Администратор проводит заказ по фактическим этапам, ведет данные города и проверяет результат каждого изменения через обновление и историю.",
        "12–14 минут",
    ),
    RoleVideo(
        8,
        "director",
        "Директор",
        "директора",
        "Управление филиалами и командой",
        "Несколько городов, сотрудники, заказы, остатки, отчеты и рассылки",
        "director-dashboard-desktop.png",
        (
            "переключать назначенные города и проверять контекст;",
            "приглашать сотрудников, выдавать и отзывать роли;",
            "контролировать заказы, склады, товары и импорт;",
            "проверять отчет AC и адресные коммуникации.",
        ),
        (
            "director-city-switcher-desktop.png",
            "combined-role-profile-desktop.png",
            "director-management-summary-desktop.png",
            "director-staff-list-desktop.png",
            "staff-invite-mobile.png",
            "staff-actions-desktop.png",
            "staff-add-role-desktop.png",
            "staff-edit-name-desktop.png",
            "director-orders-tabs-desktop.png",
            "director-orders-assign-desktop.png",
            "director-orders-collect-desktop.png",
            "director-orders-route-desktop.png",
            "director-order-cancel-form-desktop.png",
            "director-products-desktop.png",
            "product-warehouse-filter-desktop.png",
            "product-import-desktop.png",
            "product-delete-confirmation-desktop.png",
            "director-warehouses-desktop.png",
            "warehouse-connected-desktop.png",
            "director-crm-import-desktop.png",
            "director-accrual-desktop.png",
            "director-ac-report-desktop.png",
            "birthday-accrual-rule-desktop.png",
            "director-admin-history-desktop.png",
            "director-broadcast-targets-desktop.png",
            "broadcast-news-mobile.png",
            "broadcast-review-mobile.png",
        ),
        "Сейчас покажу кабинет директора. Начнем с выбора города, затем проверим сотрудников, роли, заказы, товары, склады, импорт, отчеты и рассылки.",
        "Перед каждым изменением директор проверяет город. После работы он обновляет список и сверяет сотрудника, заказ, остаток или начисление в истории.",
        "10–12 минут",
    ),
)


@dataclass(frozen=True)
class RoleError:
    title: str
    problem: str
    recovery: str


ROLE_ERRORS: dict[str, dict[str, tuple[RoleError, ...]]] = {
    "student": {
        "product-dialog-mobile.png": (
            RoleError(
                "Покупка недоступна",
                "Баланса не хватает или товар уже закончился.",
                "Вернитесь в каталог, включите фильтр наличия и выберите доступную награду по балансу.",
            ),
        ),
    },
    "parent": {
        "parent-dashboard-mobile.png": (
            RoleError(
                "Выбран не тот ребенок",
                "Баланс, корзина и заказы относятся к другому детскому профилю.",
                "Вернитесь к переключателю, выберите нужное имя и заново проверьте получателя в корзине.",
            ),
        ),
        "parent-qr-mobile.png": (
            RoleError(
                "Связь ребенка неактивна",
                "Персональная ссылка недоступна после отзыва родительской связи.",
                "Попросите администратора проверить связь и сформировать новое подключение для родителя.",
            ),
        ),
    },
    "teacher": {
        "teacher-profile-groups-desktop.png": (
            RoleError(
                "Группы не найдены",
                "ФИО не совпало с именем преподавателя в данных LMS.",
                "Нажмите «Изменить ФИО», перепишите фамилию и имя как в LMS и снова выберите «Сохранить и найти группы».",
            ),
        ),
        "teacher-qr-mobile.png": (
            RoleError(
                "QR-код недоступен",
                "Родитель ученика еще не подключен к системе.",
                "Попросите родителя открыть письмо школы и перейти по своей персональной ссылке.",
            ),
        ),
        "teacher-orders-venue-mobile.png": (
            RoleError(
                "Заказ еще не доставлен",
                "Кнопка приема недоступна, пока заказ не появился на площадке.",
                "Не меняйте состояние заранее; дождитесь доставки и обновите список заказов.",
            ),
        ),
    },
    "curator": {
        "curator-broadcast-targets-desktop.png": (
            RoleError(
                "Аудитория пуста",
                "Выбранная группа недоступна или фильтры исключили всех получателей.",
                "Сбросьте лишние условия, выберите доступную группу и снова рассчитайте аудиторию.",
            ),
        ),
    },
    "administrator": {
        "order-warehouse-dialog-desktop.png": (
            RoleError(
                "На складе нет остатка",
                "Выбранный склад не может зарезервировать нужное количество товара.",
                "Проверьте остальные склады или сначала исправьте фактический остаток товара.",
            ),
        ),
        "order-cancel-form-desktop.png": (
            RoleError(
                "Не указана причина",
                "Система не подтверждает отмену без объяснения операции.",
                "Выберите готовую причину или заполните собственный текст, затем повторите подтверждение.",
            ),
        ),
        "crm-import-desktop.png": (
            RoleError(
                "Ошибка строки импорта",
                "Предпросмотр нашел неизвестный город или незаполненное обязательное поле.",
                "Исправьте отмеченную строку в таблице и загрузите файл повторно до записи данных.",
            ),
        ),
        "product-import-desktop.png": (
            RoleError(
                "Фотография не загрузилась",
                "Облачная ссылка закрыта, временная или ведет на страницу без публичного доступа.",
                "Откройте ссылку в приватном окне, включите просмотр по ссылке и повторно импортируйте товар.",
            ),
        ),
        "product-delete-confirmation-desktop.png": (
            RoleError(
                "Удаление недоступно",
                "У товара есть защищенная история движения остатков или заказов.",
                "Закройте подтверждение и скройте товар обычной кнопкой глаза, чтобы сохранить историю.",
            ),
        ),
    },
    "director": {
        "director-management-summary-desktop.png": (
            RoleError(
                "Выбран неверный город",
                "Сводка и рабочие списки относятся не к тому филиалу.",
                "Смените город в верхней панели и повторно проверьте контекст перед изменением.",
            ),
            RoleError(
                "Действие недоступно",
                "Город не назначен директору или роль не дает нужного права.",
                "Не обходите ограничение; проверьте назначенные города и полномочия сотрудника.",
            ),
        ),
        "product-import-desktop.png": (
            RoleError(
                "Фотография не загрузилась",
                "Ссылка на Яндекс Диск или Google Drive недоступна без авторизации.",
                "Откройте общий доступ, проверьте ссылку в приватном окне и повторите импорт строки.",
            ),
        ),
    },
}


ROLE_FINAL_ACTIONS: dict[str, tuple[tuple[tuple[str, str], ...], str]] = {
    "student": (
        (
            ("Проверить баланс", "Обновите главную и сверьте доступные AC."),
            ("Выбрать награду", "Откройте магазин и добавьте доступный товар."),
            ("Найти заказ", "После покупки проверьте его состояние в заказах."),
        ),
        "Откройте магазин, выберите награду по балансу и проверьте ее в корзине. После оформления найдите заказ в нужной вкладке.",
    ),
    "parent": (
        (
            ("Выбрать ребенка", "Сверьте имя, баланс и отдельную корзину."),
            ("Проверить доступ", "Откройте QR-код или скопируйте детскую ссылку."),
            ("Сверить получателя", "Перед оплатой повторно проверьте имя ребенка."),
        ),
        "Выберите ребенка, проверьте его баланс и откройте QR-доступ. Перед покупкой еще раз сверьте имя получателя в корзине.",
    ),
    "teacher": (
        (
            ("Связать группы", "Введите ФИО как в LMS и проверьте найденных учеников."),
            ("Начислить AC", "Проведите одну операцию и проверьте историю."),
            ("Проверить выдачу", "Откройте заказы, доставленные на площадку."),
        ),
        "Сначала проверьте ФИО и найденные группы. Затем начислите AC одному ученику, откройте историю и проверьте заказы, доставленные на площадку.",
    ),
    "curator": (
        (
            ("Проверить группы", "Сверьте доступную аудиторию филиала."),
            ("Подготовить новость", "Выберите формат, площадки и группы."),
            ("Сверить отправку", "Проверьте число получателей до публикации."),
        ),
        "Откройте рассылки, выберите доступную группу и рассчитайте аудиторию. Отправляйте новость только после проверки числа получателей.",
    ),
    "administrator": (
        (
            ("Открыть заказы", "Начните с очереди зарезервированных заказов."),
            ("Назначить склад", "Проверьте остаток для каждой позиции."),
            ("Обновить очередь", "Убедитесь, что заказ ожидает доставки."),
        ),
        "Откройте заказы и начните с вкладки «Назначить склад». Выберите один заказ, проверьте остатки и подтвердите склады его позиций.",
    ),
    "director": (
        (
            ("Выбрать город", "Сверьте филиал в верхней панели."),
            ("Проверить сотрудников", "Откройте единые карточки и активные роли."),
            ("Проверить операции", "Сверьте заказы и отчет AC города."),
        ),
        "Выберите рабочий город и откройте список сотрудников. Проверьте активные роли, затем сверьте заказы и отчет AC этого филиала.",
    ),
}


def add_narration(
    items: list[dict[str, str]],
    title: str,
    text: str,
    *,
    screen_type: str = "information",
    chapter: str | None = None,
    cursor_action: str = "Без нажатия: показать экран",
    screen_text: str = "",
) -> None:
    item = {
        "title": title,
        "text": " ".join(text.split()),
        "screen_type": screen_type,
        "cursor_action": " ".join(cursor_action.split()),
        "screen_text": " ".join((screen_text or title).split()),
    }
    if chapter:
        item["chapter"] = chapter
    items.append(item)


def add_cover(
    prs: Presentation,
    *,
    kicker: str,
    title: str,
    subtitle: str,
    asset: str,
    label: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide, guide.WHITE)
    guide.add_rect(slide, 0, 0, 5.12, 7.5, fill=guide.PURPLE_DARK, line=None)
    guide.add_rect(slide, 0.62, 0.62, 0.58, 0.58, fill=guide.YELLOW, line=None, radius=True)
    guide.add_text(
        slide,
        "A",
        0.62,
        0.64,
        0.58,
        0.50,
        size=22,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    guide.add_text(slide, "ALGO MAX", 1.42, 0.78, 2.25, 0.30, size=18, fill=guide.WHITE, bold=True)
    guide.add_text(slide, kicker.upper(), 0.62, 1.48, 3.90, 0.26, size=11, fill="D9C8F5", bold=True)
    title_size = 28 if len(title) > 30 else 33
    guide.add_text(slide, title, 0.62, 1.92, 3.95, 1.50, size=title_size, fill=guide.WHITE, bold=True)
    guide.add_text(slide, subtitle, 0.62, 3.70, 3.88, 1.15, size=16, fill="E8DFFC")
    guide.add_text(slide, "Серия видеоинструкций · 2026", 0.62, 6.62, 3.70, 0.24, size=10, fill="CBB9E8")
    visual._add_picture_frame(
        guide,
        slide,
        asset=asset,
        x=5.50,
        y=0.56,
        w=7.30,
        h=6.38,
        label=label,
        crop=(0.0, 0.0, 0.0, 0.0),
    )


def add_outro(
    prs: Presentation,
    *,
    title: str,
    items: tuple[tuple[str, str], ...],
    note: str,
    panel_title: str = "Что проверить после просмотра",
    status_text: str = "Руководство по роли завершено",
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide, guide.WHITE)
    guide.add_rect(slide, 0, 0, 4.62, 7.5, fill=guide.PURPLE_DARK, line=None)
    guide.add_text(slide, "ALGO MAX", 0.62, 0.72, 2.40, 0.34, size=19, fill=guide.WHITE, bold=True)
    title_size = 25 if len(title) > 32 else 30
    guide.add_text(slide, title, 0.62, 1.62, 3.72, 1.75, size=title_size, fill=guide.WHITE, bold=True)
    guide.add_text(slide, status_text, 0.62, 5.92, 3.36, 0.55, size=13, fill="D9C8F5")
    guide.add_text(slide, panel_title, 5.28, 1.08, 6.75, 0.45, size=23, bold=True)
    for index, (item_title, item_description) in enumerate(items, start=1):
        guide.add_step(
            slide,
            index,
            item_title,
            item_description,
            5.30,
            1.88 + (index - 1) * 1.12,
            6.70,
        )
    guide.add_callout(slide, "Следующий материал", note, 5.28, 5.72, 6.84, kind="yellow")


def add_role_capabilities(prs: Presentation, role: RoleVideo) -> None:
    guide.add_cards_slide(
        prs,
        kicker=f"Роль: {role.role}",
        title="Что доступно и за что отвечает роль",
        cards=[
            (f"Задача {index}", item.rstrip(";."), kind)
            for index, (item, kind) in enumerate(
                zip(role.capabilities, ("purple", "yellow", "teal", "coral"), strict=True),
                start=1,
            )
        ],
        note_title="Границы доступа сохраняются на каждом экране",
        note_text="Если нужного раздела, группы или действия нет, сначала проверьте роль и выбранный город, затем обратитесь к ответственному администратору.",
        note_kind="purple",
    )


def _lower_first(text: str) -> str:
    value = text.strip()
    return value[:1].lower() + value[1:] if value else value


def visual_step_narration(
    role: RoleVideo,
    spec: visual.VisualSlide,
    point: visual.Hotspot,
    *,
    step_index: int,
    task_index: int,
    final_step: bool,
) -> str:
    action = f"{point.title.rstrip('.!?')}. {point.text}"
    if step_index == 1:
        lead_templates = (
            f"Открываю «{spec.title}».",
            f"Переходим к задаче «{spec.title}».",
            f"Следующая задача — «{spec.title}».",
            f"Теперь разбираю «{spec.title}».",
            f"Для этой операции открываю «{spec.title}».",
            f"На экране «{spec.title}» начинаю с первого действия.",
        )
        lead = lead_templates[task_index % len(lead_templates)]
    elif final_step:
        lead_templates = (
            "Завершаю операцию.",
            "Остается последнее действие.",
            "Последний шаг.",
            "В конце.",
            "Теперь завершаю.",
            "Финальное действие.",
        )
        lead = lead_templates[task_index % len(lead_templates)]
    else:
        lead_templates = (
            "Дальше.",
            "Затем.",
            "После этого.",
            "Следующее действие.",
            "Продолжаю.",
            "Теперь.",
        )
        lead = lead_templates[task_index % len(lead_templates)]
    if step_index == 1:
        parts = [lead, action]
    else:
        parts = [f"{lead.rstrip('.')} — {_lower_first(action)}"]
    if final_step:
        result_templates = (
            "Проверяю результат: {}",
            "После этого убеждаюсь: {}",
            "Готово: {}",
            "Контроль результата: {}",
        )
        parts.append(result_templates[task_index % len(result_templates)].format(_lower_first(spec.note)))
    return " ".join(parts)


def error_narration(error: RoleError) -> str:
    return f"{error.problem} {error.recovery}"


def role_visual_spec(role: RoleVideo, spec: visual.VisualSlide) -> visual.VisualSlide:
    if role.video_id == "administrator" and spec.asset == "crm-import-desktop.png":
        return replace(
            spec,
            note="Город берется из строки; администратор подтверждает импорт только в текущем филиале.",
        )
    if role.video_id == "administrator" and spec.asset == "admin-history-desktop.png":
        return replace(
            spec,
            note="В истории остаются только действия внутри текущей системы: автор, время и измененный объект.",
        )
    if role.video_id == "director" and spec.asset == "staff-invite-mobile.png":
        return replace(
            spec,
            note="После перехода сотрудник появляется в активных ролях выбранного города.",
        )
    if role.video_id == "teacher" and spec.asset == "teacher-student-registry-desktop.png":
        return replace(
            spec,
            note="Преподаватель видит только своих учеников и назначенные ему группы.",
        )
    return spec


def all_visual_specs() -> dict[str, visual.VisualSlide]:
    groups: Iterable[tuple[visual.VisualSlide, ...]] = (
        visual._profile_slides(),
        visual._store_slides(),
        visual._order_slides(),
        visual._student_ac_slides(),
        visual._broadcast_slides(),
        visual._management_slides(),
        visual._qr_help_slides(),
    )
    specs: dict[str, visual.VisualSlide] = {}
    for group in groups:
        for spec in group:
            if spec.asset in specs:
                raise ValueError(f"Duplicate visual asset: {spec.asset}")
            specs[spec.asset] = spec

    role_variants: dict[str, tuple[str, str | None]] = {
        "curator-broadcast-targets-desktop.png": (
            "broadcast-targets-desktop.png",
            "Куратор видит только доступные ему группы и проверяет итоговое число получателей.",
        ),
        "curator-ac-report-desktop.png": (
            "ac-report-desktop.png",
            "Куратор проверяет операции только по доступным ему группам и ученикам.",
        ),
        "admin-ac-report-desktop.png": (
            "ac-report-desktop.png",
            "Администратор сверяет период, ученика, причину, сумму и автора операции.",
        ),
        "admin-broadcast-targets-desktop.png": (
            "broadcast-targets-desktop.png",
            "До перехода к тексту администратор проверяет город и рассчитанную аудиторию.",
        ),
        "director-management-summary-desktop.png": (
            "management-summary-desktop.png",
            "Все показатели сводки относятся к выбранному директором городу.",
        ),
        "director-staff-list-desktop.png": (
            "staff-list-desktop.png",
            "Директор может настроить или отозвать роль сотрудника назначенного города.",
        ),
        "director-orders-tabs-desktop.png": ("admin-orders-tabs-desktop.png", None),
        "director-orders-assign-desktop.png": ("orders-assign-desktop.png", None),
        "director-orders-collect-desktop.png": ("orders-collect-desktop.png", None),
        "director-orders-route-desktop.png": ("orders-route-desktop.png", None),
        "director-order-cancel-form-desktop.png": (
            "order-cancel-form-desktop.png",
            "После подтверждения заказ отменяется, причина сохраняется, а AC возвращаются ученику.",
        ),
        "director-products-desktop.png": ("products-desktop.png", None),
        "director-warehouses-desktop.png": ("warehouses-desktop.png", None),
        "director-crm-import-desktop.png": (
            "crm-import-desktop.png",
            "Директор подтверждает импорт только после проверки городов и строк предпросмотра.",
        ),
        "director-ac-report-desktop.png": ("ac-report-desktop.png", None),
        "director-admin-history-desktop.png": (
            "admin-history-desktop.png",
            "История показывает действия сотрудников выбранного директором города.",
        ),
        "director-broadcast-targets-desktop.png": ("broadcast-targets-desktop.png", None),
    }
    for asset, (base_asset, note) in role_variants.items():
        base_spec = specs[base_asset]
        specs[asset] = replace(base_spec, asset=asset, note=note or base_spec.note)

    for role_name, asset in (
        ("Куратор", "curator-accrual-desktop.png"),
        ("Администратор", "admin-accrual-desktop.png"),
        ("Директор", "director-accrual-desktop.png"),
    ):
        specs[asset] = visual.VisualSlide(
            "Начисления",
            f"Начисление AC: {role_name.lower()}",
            "Начисления",
            asset,
            (
                visual.Hotspot(0.72, 0.45, "Выберите группу", "После выбора появится список доступных учеников."),
                visual.Hotspot(0.70, 0.57, "Выберите причину", "Сумма подставляется из правила текущего города."),
                visual.Hotspot(0.86, 0.68, "Отметьте учеников", "Перед начислением проверьте список и итоговое количество."),
            ),
            note="После успеха откройте историю AC и убедитесь, что сумма и причина записаны один раз.",
        )

    specs["admin-staff-invite-desktop.png"] = visual.VisualSlide(
        "Сотрудники",
        "Одноразовая ссылка сотрудника",
        "Управление -> Сотрудники -> Пригласить сотрудника",
        "admin-staff-invite-desktop.png",
        (
            visual.Hotspot(0.30, 0.58, "Выберите роль", "Список ограничен полномочиями администратора."),
            visual.Hotspot(0.57, 0.59, "Создайте ссылку", "Ссылка действует семь дней и срабатывает один раз."),
            visual.Hotspot(0.82, 0.59, "Закройте форму", "Без создания ссылки данные не меняются."),
        ),
        note="Ссылку передают конкретному сотруднику; повторное использование не создает новую связь.",
    )
    return specs


def build_role_video(role: RoleVideo, specs: dict[str, visual.VisualSlide]) -> tuple[Presentation, list[dict[str, str]]]:
    prs = Presentation()
    prs.slide_width = guide.SLIDE_W
    prs.slide_height = guide.SLIDE_H
    narration: list[dict[str, str]] = []

    add_cover(
        prs,
        kicker=f"Видео по роли · {role.role}",
        title=role.title,
        subtitle=role.subtitle,
        asset=role.cover_asset,
        label=f"Интерфейс роли: {role.role.lower()}",
    )
    cover_narration = role.intro
    add_narration(
        narration,
        role.role,
        cover_narration,
        screen_type="cover",
        chapter="Начало",
        cursor_action="Показать название роли и первый экран кабинета",
        screen_text=role.subtitle,
    )

    previous_section = ""
    for task_index, asset in enumerate(role.assets):
        try:
            spec = specs[asset]
        except KeyError as exc:
            raise KeyError(f"Unknown visual asset for {role.video_id}: {asset}") from exc
        role_spec = role_visual_spec(role, spec)
        chapter = role_spec.section if role_spec.section != previous_section else None
        previous_section = role_spec.section
        for index, point in enumerate(role_spec.hotspots, start=1):
            final_step = index == len(role_spec.hotspots)
            visual.add_visual_step_slide(
                guide,
                prs,
                role_spec,
                active_index=index,
                show_result=final_step,
            )
            add_narration(
                narration,
                f"{role_spec.title}: шаг {index}",
                visual_step_narration(
                    role,
                    role_spec,
                    point,
                    step_index=index,
                    task_index=task_index,
                    final_step=final_step,
                ),
                screen_type="action",
                chapter=chapter if index == 1 else None,
                cursor_action=f"Подвести курсор к метке {index}: {point.title}",
                screen_text=f"Шаг {index}. {point.title}",
            )

        for error in ROLE_ERRORS.get(role.video_id, {}).get(asset, ()):
            visual.add_visual_error_slide(
                guide,
                prs,
                role_spec,
                error_title=error.title,
                error_text=error.problem,
                recovery=error.recovery,
            )
            add_narration(
                narration,
                error.title,
                error_narration(error),
                screen_type="error",
                cursor_action="Показать проблемное состояние и выделить способ исправления",
                screen_text=f"Если не получилось: {error.title}",
            )

    final_items, final_narration = ROLE_FINAL_ACTIONS[role.video_id]
    add_outro(
        prs,
        title=f"Роль\n«{role.role}»\nготова к работе",
        items=final_items,
        note="Для редких операций и спорных ситуаций используйте полное руководство по ролям.",
    )
    add_narration(
        narration,
        "Итог",
        f"{role.outro} Практика после просмотра: {final_narration}",
        screen_type="outro",
        chapter="Итог",
        cursor_action="Показать итоговый чек-лист роли",
        screen_text="Повторите одну операцию и проверьте результат",
    )
    return prs, narration


def add_promo_showcase(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide)
    guide.add_brand(slide, len(prs.slides), "Один путь от мотивации до награды", "Как это работает")
    columns = (
        ("student-dashboard-mobile.png", "Ученик", "видит баланс и результат"),
        ("store-mobile.png", "Магазин", "превращает AC в награду"),
        ("teacher-orders-mobile.png", "Преподаватель", "принимает и выдает заказ"),
    )
    for index, (asset, title, caption) in enumerate(columns):
        x = 0.58 + index * 4.18
        visual._add_picture_frame(
            guide,
            slide,
            asset=asset,
            x=x,
            y=1.82,
            w=3.78,
            h=4.24,
            label=title,
            crop=(0.0, 0.0, 0.0, 0.0),
        )
        guide.add_text(slide, caption, x, 6.25, 3.78, 0.34, size=13, bold=True, align=PP_ALIGN.CENTER)
    guide.add_footer(slide, "Короткое знакомство · Algo MAX")


def add_order_flow(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.set_background(slide)
    guide.add_brand(slide, len(prs.slides), "Каждый заказ виден на понятном этапе", "Прозрачная выдача")
    statuses = (
        ("1", "Зарезервирован", "ждет склад", "purple"),
        ("2", "Ожидает доставки", "товар собирают", "yellow"),
        ("3", "Доставлен на площадку", "доставка завершена", "teal"),
        ("4", "Учитель получил заказ", "комплект принят", "coral"),
        ("5", "Получен", "вкладка «Получены»", "teal"),
    )
    palettes = {
        "purple": (guide.PURPLE_LIGHT, guide.PURPLE),
        "yellow": (guide.YELLOW_LIGHT, guide.YELLOW),
        "teal": (guide.TEAL_LIGHT, guide.TEAL),
        "coral": (guide.CORAL_LIGHT, guide.CORAL),
    }
    for index, (number, title, caption, kind) in enumerate(statuses):
        x = 0.48 + index * 2.55
        background, accent = palettes[kind]
        guide.add_rect(slide, x, 2.06, 2.25, 2.74, fill=background, line=None)
        guide.add_rect(slide, x + 0.18, 2.28, 0.48, 0.48, fill=accent, line=None, radius=True)
        guide.add_text(slide, number, x + 0.18, 2.34, 0.48, 0.30, size=14, bold=True, align=PP_ALIGN.CENTER)
        title_size = 12.5 if len(title) > 20 else 14 if len(title) > 15 else 15
        guide.add_text(slide, title, x + 0.18, 3.00, 1.88, 0.70, size=title_size, bold=True)
        guide.add_text(slide, caption, x + 0.18, 3.82, 1.88, 0.48, size=11.5, fill=guide.MUTED)
    guide.add_callout(
        slide,
        "Одинаковая картина для семьи и сотрудников",
        "История фиксирует фактическое действие, дату и ответственного пользователя.",
        2.28,
        5.42,
        8.78,
        kind="yellow",
    )
    guide.add_footer(slide, "Короткое знакомство · Algo MAX")


def build_promo_video() -> tuple[Presentation, list[dict[str, str]]]:
    prs = Presentation()
    prs.slide_width = guide.SLIDE_W
    prs.slide_height = guide.SLIDE_H
    narration: list[dict[str, str]] = []

    add_cover(
        prs,
        kicker="Короткое знакомство",
        title="Algo MAX",
        subtitle="Мотивация учеников, магазин наград и прозрачная выдача в одном сервисе",
        asset="student-dashboard-desktop.png",
        label="Личный кабинет ученика",
    )
    add_narration(
        narration,
        "Algo MAX",
        "Algo MAX: мотивация превращается в награду.",
        screen_type="cover",
        chapter="Знакомство",
    )

    add_promo_showcase(prs)
    add_narration(
        narration,
        "Путь ученика",
        "Ученик получает астрокоины и сразу видит новый баланс. Затем выбирает награду, оформляет заказ и следит за движением.",
        screen_type="action",
        chapter="Путь награды",
    )

    guide.add_cards_slide(
        prs,
        kicker="Один рабочий контур",
        title="Данные связаны между ролями и филиалами",
        cards=[
            ("Начисления", "Каждая операция остается в истории ученика.", "purple"),
            ("Остатки", "Склад и доступное количество видны до подтверждения.", "yellow"),
            ("Заказы", "Состояние меняется только после фактического действия.", "teal"),
            ("Доступ", "Роль и город ограничивают рабочие данные.", "coral"),
        ],
        note_title="Одинаковая картина для всех участников",
        note_text="Семья и сотрудники видят один фактический путь награды без ручных уточнений.",
        note_kind="purple",
    )
    add_narration(
        narration,
        "Связанные данные",
        "Начисления, товары, остатки и заказы связаны с конкретным учеником и городом. Сотрудники видят свою очередь действий, а семья получает понятную историю без ручных уточнений.",
        screen_type="information",
        chapter="Связанные данные",
    )

    add_order_flow(prs)
    add_narration(
        narration,
        "Заказ",
        "Заказ проходит резерв, доставку, площадку и преподавателя. После выдачи семья видит его во вкладке «Получены». История сохраняет время каждого перехода между состояниями.",
        screen_type="action",
        chapter="Прозрачная выдача",
    )

    guide.add_cards_slide(
        prs,
        kicker="Шесть пользовательских ролей",
        title="Каждый видит только свою работу",
        cards=[
            ("Семья", "Ученик и родитель: баланс, магазин, корзина и личные заказы.", "purple"),
            ("Команда", "Преподаватель и куратор: ученики, начисления, выдача и новости.", "yellow"),
            ("Операции", "Администратор: ученики, товары, склады и полный цикл заказа.", "teal"),
            ("Филиалы", "Директор: сотрудники, отчеты и контроль нескольких городов.", "coral"),
        ],
        note_title="Интерфейс подстраивается под роль",
        note_text="Пользователь не отвлекается на разделы и данные, которые ему недоступны.",
        note_kind="teal",
    )
    add_narration(
        narration,
        "Роли",
        "Семья покупает, преподаватель выдает, администрация управляет остатками. Каждый пользователь видит только свои данные и рабочие действия. Так меньше ручных уточнений между семьей и филиалом.",
        screen_type="information",
        chapter="Роли",
    )

    add_outro(
        prs,
        title="Откройте Algo MAX в MAX",
        items=(
            ("Откройте свой кабинет", "Проверьте профиль и доступный рабочий раздел."),
            ("Выполните один сценарий", "Начислите AC или выберите доступную награду."),
            ("Сверьте результат", "Обновите экран и проверьте историю операции."),
        ),
        note="Подробный обзор показывает всю систему, а ролевое видео объясняет конкретные кнопки вашего кабинета.",
        panel_title="Следующее действие",
        status_text="Один сервис для мотивации и выдачи",
    )
    add_narration(
        narration,
        "Следующее действие",
        "Откройте Algo MAX и проверьте свой кабинет. Затем выберите обзор или видео своей роли. Выполните одну операцию и сверьте результат.",
        screen_type="outro",
        chapter="Начать работу",
        cursor_action="Показать варианты продолжения обучения",
        screen_text="Откройте кабинет и выберите видео своей роли",
    )
    return prs, narration


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def build_detailed_video() -> tuple[Presentation, list[dict[str, str]]]:
    prs = product.build_presentation()
    source_narration = [dict(item) for item in product.NARRATION]
    for index in range(len(prs.slides) - 1, 16, -1):
        delete_slide(prs, index)

    add_outro(
        prs,
        title="Обзор завершен",
        items=(
            ("Выбрать свою роль", "Откройте инструкцию только для доступного вам кабинета."),
            ("Повторить сценарий", "Следуйте маршруту экрана и нажимайте указанные кнопки по порядку."),
            ("Сверить со справочником", "Для редкой операции найдите соответствующий раздел полного руководства."),
        ),
        note="Отдельные ролевые видео показывают конкретные кнопки, ограничения и проверку результата каждой операции.",
        panel_title="Как продолжить обучение",
        status_text="Подробный обзор завершен",
    )
    narration: list[dict[str, str]] = []
    chapters = {
        0: "Продукт и доступ",
        4: "Пользовательские роли",
        10: "Астрокоины и магазин",
        12: "Физические и цифровые заказы",
        14: "Данные и коммуникации",
    }
    for index, item in enumerate(source_narration[:17]):
        add_narration(
            narration,
            str(item["title"]),
            str(item["text"]),
            screen_type="overview",
            chapter=chapters.get(index),
        )
    add_narration(
        narration,
        "Следующий шаг",
        "Теперь выберите видео своей роли и повторите одну рабочую операцию. В нем показаны точные кнопки, результат и способ исправить типичную ошибку.",
        screen_type="outro",
        chapter="Следующий шаг",
    )
    return prs, narration


def narration_duration(item: dict[str, str]) -> float:
    words = max(len(item["text"].split()), 1)
    spoken = words / 132 * 60
    minimums = {
        "cover": 5.5,
        "action": 5.0,
        "error": 7.0,
        "information": 6.0,
        "overview": 7.0,
        "outro": 7.0,
    }
    return round(max(spoken + 0.8, minimums.get(item.get("screen_type", "information"), 6.0)), 2)


def prepare_narration(narration: list[dict[str, str]]) -> list[dict[str, str | float]]:
    return [dict(item, duration_seconds=narration_duration(item)) for item in narration]


def estimate_seconds(narration: list[dict[str, object]]) -> int:
    return max(1, round(sum(float(item.get("duration_seconds") or narration_duration(item)) for item in narration)))


def format_duration(seconds: int) -> str:
    minutes, remainder = divmod(seconds, 60)
    return f"{minutes}:{remainder:02d}"


def format_timecode(seconds: float) -> str:
    minutes, remainder = divmod(round(seconds), 60)
    return f"{minutes:02d}:{remainder:02d}"


def table_cell(value: object) -> str:
    return " ".join(str(value).split()).replace("|", "\\|")


def link_from_docs(value: object) -> str:
    normalized = str(value).replace("\\", "/")
    if normalized.startswith("docs/"):
        return normalized.removeprefix("docs/")
    return f"../{normalized}"


def write_script(path: Path, title: str, purpose: str, narration: list[dict[str, object]]) -> None:
    lines = [
        f"# {title}",
        "",
        f"**Назначение:** {purpose}",
        f"**Оценочная длительность:** {format_duration(estimate_seconds(narration))}",
        "",
        "Темп чтения: 125–140 слов в минуту. Пауза между смысловыми блоками — около одной секунды.",
        "",
        "## Покадровый план",
        "",
        "| Таймкод | Экран | Действие курсора | Текст диктора | Текст на экране |",
        "|---|---|---|---|---|",
        "",
    ]
    elapsed = 0.0
    for item in narration:
        duration = float(item["duration_seconds"])
        timecode = f"{format_timecode(elapsed)}–{format_timecode(elapsed + duration)}"
        lines.append(
            "| "
            + " | ".join(
                (
                    timecode,
                    table_cell(item["title"]),
                    table_cell(item.get("cursor_action", "Без нажатия")),
                    table_cell(item["text"]),
                    table_cell(item.get("screen_text", item["title"])),
                )
            )
            + " |"
        )
        elapsed += duration
    lines.extend(
        (
            "",
            "## Текст для чтения",
            "",
        )
    )
    for index, item in enumerate(narration, start=1):
        lines.extend((f"### {index}. {item['title']}", "", str(item["text"]), ""))
    path.write_text("\n".join(lines), encoding="utf-8")


def save_video_source(
    *,
    order: int,
    video_id: str,
    filename: str,
    title: str,
    kind: str,
    purpose: str,
    target_duration: str,
    prs: Presentation,
    narration: list[dict[str, str]],
) -> dict[str, object]:
    prepared_narration = prepare_narration(narration)
    if len(prs.slides) != len(prepared_narration):
        raise ValueError(f"{video_id}: {len(prs.slides)} slides, {len(narration)} narration items")
    pptx_path = SERIES_DIR / f"{order:02d}_{filename}.pptx"
    narration_path = SERIES_DIR / f"{order:02d}_{filename}.narration.json"
    script_path = SERIES_DIR / f"{order:02d}_{filename}.script.md"
    video_path = VIDEO_OUTPUT_DIR / f"{order:02d}_{filename}.mp4"
    pdf_path = ROOT / "output" / "pdf" / "video-series" / f"{order:02d}_{filename}.pdf"
    subtitle_srt_path = video_path.with_suffix(".srt")
    subtitle_vtt_path = video_path.with_suffix(".vtt")
    chapters_path = video_path.with_suffix(".chapters.txt")
    prs.save(pptx_path)
    narration_path.write_text(json.dumps(prepared_narration, ensure_ascii=False, indent=2), encoding="utf-8")
    write_script(script_path, title, purpose, prepared_narration)
    return {
        "order": order,
        "id": video_id,
        "kind": kind,
        "title": title,
        "purpose": purpose,
        "target_duration": target_duration,
        "estimated_seconds": estimate_seconds(prepared_narration),
        "slides": len(prs.slides),
        "presentation": str(pptx_path.relative_to(ROOT)),
        "pdf": str(pdf_path.relative_to(ROOT)),
        "narration": str(narration_path.relative_to(ROOT)),
        "script": str(script_path.relative_to(ROOT)),
        "video": str(video_path.relative_to(ROOT)),
        "subtitles_srt": str(subtitle_srt_path.relative_to(ROOT)),
        "subtitles_vtt": str(subtitle_vtt_path.relative_to(ROOT)),
        "chapters": str(chapters_path.relative_to(ROOT)),
    }


def write_overview(entries: list[dict[str, object]]) -> None:
    lines = [
        "# Серия видео по Algo MAX",
        "",
        "Материалы разделены по задаче. Короткое видео знакомит с продуктом, подробное показывает сквозную систему, а ролевые видео обучают конкретной работе без повторения общей теории.",
        "",
        "## Состав",
        "",
        "| № | Видео | Для чего | Слайды | Оценка | Целевая длительность | Материалы |",
        "|---:|---|---|---:|---:|---|---|",
    ]
    for entry in entries:
        lines.append(
            f"| {entry['order']} | {entry['title']} | {entry['purpose']} | {entry['slides']} | {format_duration(int(entry['estimated_seconds']))} | {entry['target_duration']} | "
            f"[PPTX]({link_from_docs(entry['presentation'])}) · [PDF]({link_from_docs(entry['pdf'])}) · "
            f"[сценарий]({link_from_docs(entry['script'])}) · [MP4 без голоса]({link_from_docs(entry['video'])}) · "
            f"[SRT]({link_from_docs(entry['subtitles_srt'])}) |"
        )
    lines.extend(
        (
            "",
            "## Правила использования",
            "",
            "1. Для презентации продукта отправляйте только короткое видео.",
            "2. Новому партнеру или руководителю после короткого видео отправляйте подробный обзор.",
            "3. Сотруднику или члену семьи отправляйте только видео его роли и при необходимости ссылку на полное руководство.",
            "4. Для обновления интерфейса сначала замените скриншоты в `user-guide-assets/visual-v6`, затем пересоберите серию.",
            "",
            "## Пересборка",
            "",
            "```powershell",
            "max_bot_venv\\Scripts\\python.exe docs\\build_video_series.py",
            "powershell.exe -NoProfile -ExecutionPolicy Bypass -File docs\\export_video_series.ps1",
            "powershell.exe -NoProfile -ExecutionPolicy Bypass -File docs\\build_video_contact_sheets.ps1",
            "max_bot_venv\\Scripts\\python.exe docs\\audit_training_materials.py",
            "```",
            "",
            "MP4 создаются без звуковой дорожки и готовы для записи живого голоса. SRT, VTT и главы находятся в `output/video-series`; PPTX и тексты диктора — в `docs/video-series`; PDF — в `output/pdf/video-series`.",
            "",
        )
    )
    OVERVIEW_PATH.write_text("\n".join(lines), encoding="utf-8")


def write_output_readme(entries: list[dict[str, object]]) -> None:
    lines = [
        "# Готовые видео Algo MAX",
        "",
        "В каталоге находятся восемь мастер-видео без звуковой дорожки. Они рассчитаны на запись живого голоса по готовым таймкодам.",
        "",
        "| № | Видео | Длительность | Материалы |",
        "|---:|---|---:|---|",
    ]
    for entry in entries:
        video_name = Path(str(entry["video"])).name
        pdf_name = Path(str(entry["pdf"])).name
        script_name = Path(str(entry["script"])).name
        pptx_name = Path(str(entry["presentation"])).name
        lines.append(
            f"| {entry['order']} | [{entry['title']}]({video_name}) | {format_duration(int(entry['estimated_seconds']))} | "
            f"[PDF](../pdf/video-series/{pdf_name}) · [PPTX](../../docs/video-series/{pptx_name}) · "
            f"[текст диктора](../../docs/video-series/{script_name}) |"
        )
    lines.extend(
        (
            "",
            "Для каждого MP4 рядом находятся SRT, VTT и файл глав. Полный итог проверки: `../../docs/QA_REPORT_RU.md`.",
            "",
        )
    )
    (VIDEO_OUTPUT_DIR / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    SERIES_DIR.mkdir(parents=True, exist_ok=True)
    VIDEO_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    specs = all_visual_specs()
    entries: list[dict[str, object]] = []

    with footer_style("Короткое знакомство · Algo MAX"):
        promo_prs, promo_narration = build_promo_video()
    entries.append(
        save_video_source(
            order=1,
            video_id="promo",
            filename="Algo_MAX_Короткое_знакомство",
            title="Algo MAX — короткое знакомство",
            kind="promo",
            purpose="Быстро и понятно показать ценность продукта без обучения кнопкам.",
            target_duration="45–60 секунд",
            prs=promo_prs,
            narration=promo_narration,
        )
    )

    with footer_style("Подробный обзор · Algo MAX"):
        detailed_prs, detailed_narration = build_detailed_video()
    entries.append(
        save_video_source(
            order=2,
            video_id="detailed",
            filename="Algo_MAX_Подробный_обзор",
            title="Algo MAX — подробный обзор продукта",
            kind="overview",
            purpose="Показать роли, данные и полный цикл от начисления до выдачи награды.",
            target_duration="6–8 минут",
            prs=detailed_prs,
            narration=detailed_narration,
        )
    )

    for role in ROLE_VIDEOS:
        with footer_style(f"Роль: {role.role} · Algo MAX"):
            role_prs, role_narration = build_role_video(role, specs)
        entries.append(
            save_video_source(
                order=role.order,
                video_id=role.video_id,
                filename=f"Algo_MAX_Роль_{role.role}",
                title=f"Algo MAX — {role.role}",
                kind="role",
                purpose=f"Показать пользователю с ролью «{role.role}» доступные разделы, кнопки и ограничения.",
                target_duration=role.target_duration,
                prs=role_prs,
                narration=role_narration,
            )
        )

    MANIFEST_PATH.write_text(
        json.dumps(
            {
                "series": "Algo MAX",
                "version": date.today().isoformat(),
                "voice": "manual",
                "audio_mode": "silent_master_for_human_voiceover",
                "videos": entries,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    write_overview(entries)
    write_output_readme(entries)

    active_source_paths = {
        ROOT / str(entry[key])
        for entry in entries
        for key in (
            "presentation",
            "pdf",
            "narration",
            "script",
            "video",
            "subtitles_srt",
            "subtitles_vtt",
            "chapters",
        )
    }
    for stale_path in (*SERIES_DIR.iterdir(), *VIDEO_OUTPUT_DIR.glob("*.mp4")):
        if stale_path.is_file() and stale_path not in active_source_paths:
            stale_path.unlink(missing_ok=True)
    active_video_ids = {str(entry["id"]) for entry in entries}
    for generated_dir in (VIDEO_OUTPUT_DIR / "slides", VIDEO_OUTPUT_DIR / "audio"):
        if not generated_dir.exists():
            continue
        for child in generated_dir.iterdir():
            if child.is_dir() and child.name not in active_video_ids:
                shutil.rmtree(child)
    narrated_dir = VIDEO_OUTPUT_DIR / "narrated"
    if narrated_dir.exists():
        for child in narrated_dir.glob("*.pptx"):
            if child.stem not in active_video_ids:
                child.unlink()
    print(f"Created {len(entries)} video sources")
    for entry in entries:
        print(f"[{entry['order']:02d}] {entry['id']}: {entry['slides']} slides, ~{format_duration(int(entry['estimated_seconds']))}")
    print(MANIFEST_PATH)


if __name__ == "__main__":
    main()
