from __future__ import annotations

import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
PRESENTATION = ROOT / "docs" / "Algo_MAX_Руководство_по_ролям.pptx"
ASSETS = ROOT / "docs" / "user-guide-assets" / "visual-v6"
INDEX = ROOT / "app" / "web" / "static" / "miniapp" / "index.html"

EXPECTED_SLIDES = 100
MIN_PICTURES = 70
MIN_UNIQUE_IMAGES = 70
FORBIDDEN_TEXT = (
    "Выполненные",
    "К действию",
    "Оформить возврат",
    "Из сводки",
    "Сводка по филиалу",
    "супер" + "админ",
    "0.78.1",
)


def slide_text(slide) -> str:
    return "\n".join(
        shape.text.strip()
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def main() -> None:
    errors: list[str] = []
    prs = Presentation(PRESENTATION)
    if len(prs.slides) != EXPECTED_SLIDES:
        errors.append(f"slides: expected {EXPECTED_SLIDES}, got {len(prs.slides)}")

    picture_count = 0
    unique_images: set[str] = set()
    text_runs = 0
    smallest_font = 1000.0
    geometry_tolerance = 20_000

    for slide_number, slide in enumerate(prs.slides, start=1):
        text = slide_text(slide)
        if not text:
            errors.append(f"slide {slide_number}: no text")
        if "�" in text:
            errors.append(f"slide {slide_number}: replacement character")
        for forbidden in FORBIDDEN_TEXT:
            if forbidden.casefold() in text.casefold():
                errors.append(f"slide {slide_number}: forbidden text {forbidden!r}")

        for shape in slide.shapes:
            if shape.left < -geometry_tolerance or shape.top < -geometry_tolerance:
                errors.append(f"slide {slide_number}: shape starts outside slide")
            if shape.left + shape.width > prs.slide_width + geometry_tolerance:
                errors.append(f"slide {slide_number}: shape exceeds right edge")
            if shape.top + shape.height > prs.slide_height + geometry_tolerance:
                errors.append(f"slide {slide_number}: shape exceeds bottom edge")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                unique_images.add(hashlib.sha256(shape.image.blob).hexdigest())
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            text_runs += 1
                            smallest_font = min(smallest_font, run.font.size.pt)

    if picture_count < MIN_PICTURES:
        errors.append(f"pictures: expected at least {MIN_PICTURES}, got {picture_count}")
    if len(unique_images) < MIN_UNIQUE_IMAGES:
        errors.append(
            f"unique images: expected at least {MIN_UNIQUE_IMAGES}, got {len(unique_images)}"
        )
    if smallest_font < 8.0:
        errors.append(f"font size: smallest explicit font is {smallest_font:.1f} pt")
    if text_runs == 0:
        errors.append("no explicitly sized text runs found")

    if len(list(ASSETS.glob("*.png"))) < 75:
        errors.append("visual-v6: expected at least 75 PNG assets")
    if "v=0.86.4" not in INDEX.read_text(encoding="utf-8"):
        errors.append("index.html: static cache version is not 0.86.4")

    if errors:
        raise SystemExit("Guide validation failed:\n- " + "\n- ".join(errors))

    print(
        "Guide validation passed: "
        f"{len(prs.slides)} slides, {picture_count} pictures, "
        f"{len(unique_images)} unique images, "
        f"minimum explicit font {smallest_font:.1f} pt"
    )


if __name__ == "__main__":
    main()
