from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.crm_import import CRM_TEMPLATE_SHEET_NAME, CrmStudentRow, parse_crm_students

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "astrocoins-import-preview.xlsx"

PURPLE = "6F30D0"
DARK = "24152F"
YELLOW = "FFD43B"
GREEN = "DDF7EB"
GREEN_DARK = "087F5B"
RED = "FDE7EA"
RED_DARK = "B42345"
LAVENDER = "F5F0FF"
GRAY = "F4F4F7"
WHITE = "FFFFFF"
BORDER_COLOR = "DED8E7"


class LegacyDumpError(RuntimeError):
    pass


@dataclass(frozen=True)
class LegacyStudent:
    user_id: int
    name: str
    email: str | None
    birthday: str | None
    deleted_at: str | None
    balance: int | None
    groups: tuple[str, ...]
    parents: tuple[str, ...]

    @property
    def is_active(self) -> bool:
        return self.deleted_at is None


@dataclass(frozen=True)
class MatchProposal:
    current: CrmStudentRow
    candidates: tuple[LegacyStudent, ...]
    selected: LegacyStudent | None
    method: str


@dataclass(frozen=True)
class MatchResult:
    ready: tuple[MatchProposal, ...]
    review: tuple[MatchProposal, ...]
    missing: tuple[MatchProposal, ...]
    legacy_students: tuple[LegacyStudent, ...]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Сформировать XLSX-предпросмотр переноса астрокоинов из старого SQL-дампа"
    )
    parser.add_argument("--sql-path", type=Path, required=True, help="Путь к SQL-дампу MariaDB")
    parser.add_argument("--crm-path", type=Path, required=True, help="Текущая CRM-выгрузка XLSX")
    parser.add_argument(
        "--crm-sheet",
        default=CRM_TEMPLATE_SHEET_NAME,
        help="Название листа CRM; при отсутствии подходящий лист определяется автоматически",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Итоговый XLSX")
    return parser.parse_args()


def _sql_unescape(value: str) -> str:
    replacements = {
        "0": "\0",
        "n": "\n",
        "r": "\r",
        "t": "\t",
        "b": "\b",
        "Z": "\x1a",
        "\\": "\\",
        "'": "'",
        '"': '"',
    }
    result: list[str] = []
    index = 0
    while index < len(value):
        if value[index] == "\\" and index + 1 < len(value):
            index += 1
            result.append(replacements.get(value[index], value[index]))
        else:
            result.append(value[index])
        index += 1
    return "".join(result)


def _parse_sql_value(raw_value: str) -> object:
    value = raw_value.strip()
    if value.upper() == "NULL":
        return None
    if len(value) >= 2 and value[0] == "'" and value[-1] == "'":
        return _sql_unescape(value[1:-1])
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def _parse_sql_tuples(values: str) -> list[list[object]]:
    rows: list[list[object]] = []
    index = 0
    while index < len(values):
        while index < len(values) and (values[index].isspace() or values[index] == ","):
            index += 1
        if index >= len(values) or values[index] == ";":
            break
        if values[index] != "(":
            raise LegacyDumpError(f"Не удалось разобрать SQL около позиции {index}")

        index += 1
        fields: list[object] = []
        field_start = index
        quoted = False
        escaped = False
        while index < len(values):
            character = values[index]
            if quoted:
                if escaped:
                    escaped = False
                elif character == "\\":
                    escaped = True
                elif character == "'":
                    quoted = False
            else:
                if character == "'":
                    quoted = True
                elif character == ",":
                    fields.append(_parse_sql_value(values[field_start:index]))
                    field_start = index + 1
                elif character == ")":
                    fields.append(_parse_sql_value(values[field_start:index]))
                    rows.append(fields)
                    index += 1
                    break
            index += 1
        else:
            raise LegacyDumpError("SQL-дамп закончился внутри строки INSERT")
    return rows


def _extract_table_rows(sql: str, table_name: str) -> list[dict[str, object]]:
    pattern = re.compile(
        rf"INSERT INTO `{re.escape(table_name)}` \((.*?)\) VALUES\s*",
        re.DOTALL,
    )
    rows: list[dict[str, object]] = []
    position = 0
    expected_columns: list[str] | None = None

    while match := pattern.search(sql, position):
        columns = [column.strip().strip("`") for column in match.group(1).split(",")]
        if expected_columns is None:
            expected_columns = columns
        elif columns != expected_columns:
            raise LegacyDumpError(f"В таблице {table_name} меняется набор колонок INSERT")

        end = match.end()
        quoted = False
        escaped = False
        while end < len(sql):
            character = sql[end]
            if quoted:
                if escaped:
                    escaped = False
                elif character == "\\":
                    escaped = True
                elif character == "'":
                    quoted = False
            else:
                if character == "'":
                    quoted = True
                elif character == ";":
                    break
            end += 1
        if end >= len(sql):
            raise LegacyDumpError(f"Не найден конец INSERT для таблицы {table_name}")

        for values in _parse_sql_tuples(sql[match.end() : end + 1]):
            if len(values) != len(columns):
                raise LegacyDumpError(
                    f"В таблице {table_name} ожидалось {len(columns)} значений, "
                    f"получено {len(values)}"
                )
            rows.append(dict(zip(columns, values, strict=True)))
        position = end + 1

    if not rows:
        raise LegacyDumpError(f"В дампе не найдены данные таблицы {table_name}")
    return rows


def _normalized_text(value: object) -> str:
    source = str(value or "").lower().replace("ё", "е")
    normalized = "".join(
        character if character.isalnum() else " " for character in source
    )
    return " ".join(normalized.split())


def _person_key(last_name: object, first_name: object | None = None) -> tuple[str, ...]:
    tokens = _normalized_text(f"{last_name or ''} {first_name or ''}").split()
    return tuple(tokens[:2])


def _legacy_parent_matches(current_parent_names: str | None, legacy: LegacyStudent) -> bool:
    current_tokens = set(_normalized_text(current_parent_names).split())
    if len(current_tokens) < 2:
        return False
    return any(
        len(current_tokens.intersection(_normalized_text(parent_name).split())) >= 2
        for parent_name in legacy.parents
    )


def load_legacy_students(sql_path: Path) -> tuple[LegacyStudent, ...]:
    if not sql_path.is_file():
        raise LegacyDumpError(f"SQL-дамп не найден: {sql_path}")
    sql = sql_path.read_text(encoding="utf-8")
    users = _extract_table_rows(sql, "users")
    coin_rows = _extract_table_rows(sql, "store_user_coins")
    child_links = _extract_table_rows(sql, "children_user")
    group_links = _extract_table_rows(sql, "student_group_users")
    groups = _extract_table_rows(sql, "student_groups")

    users_by_id = {int(row["id"]): row for row in users}
    balances = {int(row["user_id"]): int(row["amount"]) for row in coin_rows}
    group_names = {int(row["id"]): str(row["name"]) for row in groups}

    groups_by_student: dict[int, set[str]] = defaultdict(set)
    for row in group_links:
        group_name = group_names.get(int(row["group_id"]))
        if group_name:
            groups_by_student[int(row["user_id"])].add(group_name)

    parents_by_student: dict[int, set[str]] = defaultdict(set)
    for row in child_links:
        parent = users_by_id.get(int(row["user_id"]))
        if parent and parent.get("name"):
            parents_by_student[int(row["children_id"])].add(str(parent["name"]))

    students = []
    for row in users:
        if int(row["role"]) != 4:
            continue
        user_id = int(row["id"])
        students.append(
            LegacyStudent(
                user_id=user_id,
                name=str(row["name"]),
                email=str(row["email"]) if row.get("email") else None,
                birthday=str(row["birthday"]) if row.get("birthday") else None,
                deleted_at=str(row["deleted_at"]) if row.get("deleted_at") else None,
                balance=balances.get(user_id),
                groups=tuple(sorted(groups_by_student[user_id])),
                parents=tuple(sorted(parents_by_student[user_id])),
            )
        )
    return tuple(
        sorted(
            students,
            key=lambda student: (not student.is_active, student.name, student.user_id),
        )
    )


def match_students(
    legacy_students: tuple[LegacyStudent, ...],
    current_students: list[CrmStudentRow],
) -> MatchResult:
    active_by_name: dict[tuple[str, ...], list[LegacyStudent]] = defaultdict(list)
    for legacy in legacy_students:
        if legacy.is_active:
            active_by_name[_person_key(legacy.name)].append(legacy)

    proposals: list[MatchProposal] = []
    for current in current_students:
        candidates = tuple(
            active_by_name.get(_person_key(current.last_name, current.first_name), ())
        )
        current_group = _normalized_text(current.group_name)
        group_matches = tuple(
            candidate
            for candidate in candidates
            if current_group
            and current_group in {_normalized_text(group) for group in candidate.groups}
        )
        parent_matches = tuple(
            candidate
            for candidate in candidates
            if _legacy_parent_matches(current.contact_names, candidate)
        )

        selected: LegacyStudent | None = None
        method = ""
        if len(candidates) == 1:
            selected = candidates[0]
            method = "Уникальное совпадение ФИО"
        elif len(group_matches) == 1:
            selected = group_matches[0]
            method = "ФИО и точное название группы"
        elif len(parent_matches) == 1:
            selected = parent_matches[0]
            method = "ФИО ученика и родителя"
        elif not candidates:
            method = "Действующий ученик с таким ФИО не найден в старой базе"
        else:
            method = "В старой базе несколько действующих учеников с таким ФИО"

        proposals.append(
            MatchProposal(
                current=current,
                candidates=candidates,
                selected=selected,
                method=method,
            )
        )

    selected_counts = Counter(
        proposal.selected.user_id for proposal in proposals if proposal.selected is not None
    )
    ready: list[MatchProposal] = []
    review: list[MatchProposal] = []
    missing: list[MatchProposal] = []
    for proposal in proposals:
        if proposal.selected is None:
            (review if proposal.candidates else missing).append(proposal)
        elif selected_counts[proposal.selected.user_id] == 1:
            ready.append(proposal)
        else:
            review.append(
                MatchProposal(
                    current=proposal.current,
                    candidates=proposal.candidates,
                    selected=proposal.selected,
                    method=(
                        "Один старый аккаунт совпал с несколькими строками текущей CRM; "
                        "нужно выбрать актуальную карточку"
                    ),
                )
            )

    return MatchResult(
        ready=tuple(ready),
        review=tuple(review),
        missing=tuple(missing),
        legacy_students=legacy_students,
    )


def _current_name(current: CrmStudentRow) -> str:
    return " ".join(
        value.strip()
        for value in (current.last_name, current.first_name)
        if value and value.strip()
    )


def _candidate_summary(candidates: tuple[LegacyStudent, ...]) -> str:
    return "\n".join(
        f"ID {candidate.user_id}: {candidate.name}; {candidate.balance or 0} AC; "
        f"{', '.join(candidate.groups) or 'без группы'}; "
        f"родитель: {', '.join(candidate.parents) or 'не указан'}"
        for candidate in candidates
    )


def _proposal_row(proposal: MatchProposal, *, decision: str) -> dict[str, object]:
    legacy = proposal.selected
    return {
        "Решение": decision,
        "Баланс к переносу": legacy.balance if legacy and legacy.balance is not None else 0,
        "Старый user_id": legacy.user_id if legacy else None,
        "ФИО в старой системе": legacy.name if legacy else None,
        "Старая группа": "\n".join(legacy.groups) if legacy else None,
        "Родитель в старой системе": "\n".join(legacy.parents) if legacy else None,
        "Дата рождения": legacy.birthday if legacy else None,
        "Старый email": legacy.email if legacy else None,
        "ID сделки amoCRM": proposal.current.deal_id,
        "UUID amoCRM": proposal.current.uuid,
        "ФИО в текущей CRM": _current_name(proposal.current),
        "Текущая группа": proposal.current.group_name,
        "Курс": proposal.current.course_name,
        "Преподаватель": proposal.current.teacher_name,
        "Contact ID": proposal.current.contact_ids,
        "Контакт": proposal.current.contact_names,
        "Основание": proposal.method,
        "Кандидаты старой базы": _candidate_summary(proposal.candidates),
        "Комментарий": None,
    }


def _style_title(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor=PURPLE)
    cell.font = Font(color=WHITE, bold=True, size=18)
    cell.alignment = Alignment(vertical="center")


def _write_data_sheet(
    workbook: Workbook,
    *,
    title: str,
    rows: list[dict[str, object]],
    tab_color: str,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_properties.tabColor = tab_color
    if not rows:
        sheet["A1"] = "Нет строк"
        return

    headers = list(rows[0])
    header_fill = PatternFill("solid", fgColor=PURPLE)
    header_font = Font(color=WHITE, bold=True)
    border = Border(bottom=Side(style="thin", color=BORDER_COLOR))
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_index, values in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=values.get(header))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if row_index % 2:
                cell.fill = PatternFill("solid", fgColor="FBF9FE")
            if header == "Баланс к переносу":
                cell.number_format = '0 "AC"'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    sheet.row_dimensions[1].height = 34
    sheet.sheet_view.showGridLines = False

    widths = {
        "Решение": 18,
        "Баланс к переносу": 18,
        "Старый user_id": 15,
        "ФИО в старой системе": 30,
        "Старая группа": 34,
        "Родитель в старой системе": 32,
        "Дата рождения": 16,
        "Старый email": 34,
        "ID сделки amoCRM": 20,
        "UUID amoCRM": 38,
        "ФИО в текущей CRM": 30,
        "Текущая группа": 36,
        "Курс": 32,
        "Преподаватель": 27,
        "Contact ID": 20,
        "Контакт": 32,
        "Основание": 48,
        "Кандидаты старой базы": 72,
        "Комментарий": 36,
    }
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 24)

    if "Решение" in headers:
        decision_column = get_column_letter(headers.index("Решение") + 1)
        validation = DataValidation(
            type="list",
            formula1='"ИМПОРТИРОВАТЬ,ПРОВЕРИТЬ,ПРОПУСТИТЬ"',
            allow_blank=False,
        )
        sheet.add_data_validation(validation)
        validation.add(f"{decision_column}2:{decision_column}{len(rows) + 1}")
        sheet.conditional_formatting.add(
            f"{decision_column}2:{decision_column}{len(rows) + 1}",
            FormulaRule(
                formula=[f'${decision_column}2="ИМПОРТИРОВАТЬ"'],
                fill=PatternFill("solid", fgColor=GREEN),
                font=Font(color=GREEN_DARK, bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            f"{decision_column}2:{decision_column}{len(rows) + 1}",
            FormulaRule(
                formula=[f'${decision_column}2="ПРОВЕРИТЬ"'],
                fill=PatternFill("solid", fgColor="FFF3BF"),
                font=Font(color="8A5A00", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            f"{decision_column}2:{decision_column}{len(rows) + 1}",
            FormulaRule(
                formula=[f'${decision_column}2="ПРОПУСТИТЬ"'],
                fill=PatternFill("solid", fgColor=RED),
                font=Font(color=RED_DARK, bold=True),
            ),
        )


def _write_summary(
    workbook: Workbook,
    *,
    result: MatchResult,
    current_count: int,
    sql_path: Path,
    crm_path: Path,
) -> None:
    sheet = workbook.active
    sheet.title = "Сводка"
    sheet.sheet_properties.tabColor = PURPLE
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Перенос астрокоинов из старой системы"
    _style_title(sheet["A1"])
    sheet.row_dimensions[1].height = 38

    sheet.merge_cells("A3:D3")
    sheet["A3"] = (
        "Это предварительное сопоставление. В старой базе нет CRM/LMS ID ученика, "
        "поэтому строки связаны по ФИО, группе и ФИО родителя. Перед загрузкой проверьте "
        "листы «К импорту» и «Требует проверки»."
    )
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A3"].fill = PatternFill("solid", fgColor="FFF3BF")
    sheet.row_dimensions[3].height = 58

    active_legacy = [student for student in result.legacy_students if student.is_active]
    positive_legacy = [student for student in active_legacy if (student.balance or 0) > 0]
    ready_positive = [
        proposal
        for proposal in result.ready
        if (proposal.selected.balance if proposal.selected else 0) > 0
    ]
    metrics = (
        ("Строк в текущей CRM", current_count, "Все ученики из выбранной текущей выгрузки"),
        ("Действующих учеников в старой базе", len(active_legacy), "Старые users с role = 4"),
        ("Положительных старых балансов", len(positive_legacy), "Баланс больше нуля"),
        (
            "Общая сумма старых балансов",
            sum(student.balance or 0 for student in active_legacy),
            "AC",
        ),
        ("Предложено к импорту", len(result.ready), "Однозначное соответствие один к одному"),
        (
            "Из них с положительным балансом",
            len(ready_positive),
            "Нулевой баланс тоже сохранен в таблице",
        ),
        (
            "Сумма предложенного переноса",
            sum(proposal.selected.balance or 0 for proposal in result.ready if proposal.selected),
            "AC",
        ),
        ("Требует ручной проверки", len(result.review), "Дубли или несколько старых кандидатов"),
        (
            "Не найдено в старой базе",
            len(result.missing),
            "Совпадение действующего ученика не найдено",
        ),
    )

    start_row = 6
    for column, value in enumerate(("Показатель", "Значение", "Пояснение"), start=1):
        cell = sheet.cell(start_row, column, value)
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.font = Font(color=DARK, bold=True)
    for row_index, metric in enumerate(metrics, start=start_row + 1):
        for column, value in enumerate(metric, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_index % 2:
                cell.fill = PatternFill("solid", fgColor=LAVENDER)

    details_row = start_row + len(metrics) + 3
    details = (
        ("Создано", datetime.now(UTC).astimezone().strftime("%d.%m.%Y %H:%M %Z")),
        ("SQL-дамп", str(sql_path.resolve())),
        ("SHA-256 SQL", hashlib.sha256(sql_path.read_bytes()).hexdigest()),
        ("CRM XLSX", str(crm_path.resolve())),
        ("SHA-256 CRM", hashlib.sha256(crm_path.read_bytes()).hexdigest()),
    )
    for row_offset, (label, value) in enumerate(details):
        sheet.cell(details_row + row_offset, 1, label).font = Font(bold=True, color=DARK)
        sheet.cell(details_row + row_offset, 2, value)
        sheet.merge_cells(
            start_row=details_row + row_offset,
            start_column=2,
            end_row=details_row + row_offset,
            end_column=4,
        )

    guide_row = details_row + len(details) + 2
    guides = (
        ("К импорту", "Однозначные пары. Решение заранее выставлено «ИМПОРТИРОВАТЬ»."),
        ("Требует проверки", "Проверьте кандидатов и вручную измените решение."),
        ("Нет в старой базе", "Текущие ученики, для которых старый активный аккаунт не найден."),
        ("Остатки старой базы", "Полный перечень старых учеников и их текущих балансов."),
    )
    for row_offset, (label, description) in enumerate(guides):
        row = guide_row + row_offset
        sheet.cell(row, 1, label).font = Font(bold=True, color=PURPLE)
        sheet.cell(row, 2, description)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 62
    sheet.column_dimensions["D"].width = 20


def build_workbook(
    *,
    result: MatchResult,
    current_count: int,
    sql_path: Path,
    crm_path: Path,
) -> Workbook:
    workbook = Workbook()
    _write_summary(
        workbook,
        result=result,
        current_count=current_count,
        sql_path=sql_path,
        crm_path=crm_path,
    )

    ready_rows = [
        _proposal_row(proposal, decision="ИМПОРТИРОВАТЬ") for proposal in result.ready
    ]
    review_rows = [_proposal_row(proposal, decision="ПРОВЕРИТЬ") for proposal in result.review]
    missing_rows = [_proposal_row(proposal, decision="ПРОПУСТИТЬ") for proposal in result.missing]

    ready_ids = {
        proposal.selected.user_id for proposal in result.ready if proposal.selected is not None
    }
    review_candidate_ids = {
        candidate.user_id for proposal in result.review for candidate in proposal.candidates
    }
    legacy_rows = []
    for student in result.legacy_students:
        if student.user_id in ready_ids:
            status = "Предложен к импорту"
        elif student.user_id in review_candidate_ids:
            status = "Есть в спорных совпадениях"
        elif not student.is_active:
            status = "Удален в старой системе"
        else:
            status = "Нет однозначного совпадения в текущей CRM"
        legacy_rows.append(
            {
                "Статус": status,
                "Баланс": student.balance if student.balance is not None else 0,
                "Старый user_id": student.user_id,
                "ФИО": student.name,
                "Группы": "\n".join(student.groups),
                "Родители": "\n".join(student.parents),
                "Дата рождения": student.birthday,
                "Email": student.email,
                "Удален": student.deleted_at,
                "Комментарий": None,
            }
        )

    _write_data_sheet(workbook, title="К импорту", rows=ready_rows, tab_color=GREEN_DARK)
    _write_data_sheet(
        workbook,
        title="Требует проверки",
        rows=review_rows,
        tab_color=YELLOW,
    )
    _write_data_sheet(
        workbook,
        title="Нет в старой базе",
        rows=missing_rows,
        tab_color=RED_DARK,
    )
    _write_data_sheet(
        workbook,
        title="Остатки старой базы",
        rows=legacy_rows,
        tab_color="777582",
    )
    return workbook


def export_preview(args: argparse.Namespace) -> dict[str, object]:
    legacy_students = load_legacy_students(args.sql_path)
    current_students = parse_crm_students(args.crm_path, sheet_name=args.crm_sheet)
    result = match_students(legacy_students, current_students)
    workbook = build_workbook(
        result=result,
        current_count=len(current_students),
        sql_path=args.sql_path,
        crm_path=args.crm_path,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    workbook.close()

    return {
        "output": str(args.output.resolve()),
        "current_students": len(current_students),
        "legacy_students": len(legacy_students),
        "legacy_active_students": sum(student.is_active for student in legacy_students),
        "ready": len(result.ready),
        "ready_balance": sum(
            proposal.selected.balance or 0
            for proposal in result.ready
            if proposal.selected is not None
        ),
        "review": len(result.review),
        "missing": len(result.missing),
    }


def main() -> int:
    result = export_preview(parse_args())
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
