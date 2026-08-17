from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

CRM_TEMPLATE_SHEET_NAME = "Шаблон"
CRM_LEGACY_SHEET_NAME = "Сделки"
CRM_TEMPLATE_COLUMNS: tuple[tuple[str, str, bool, str, str], ...] = (
    (
        "lms_student_id",
        "ID ученика",
        True,
        "Уникальный ID ученика из LMS или CRM",
        "1841",
    ),
    (
        "deal_id",
        "ID сделки amoCRM",
        True,
        "Числовой ID сделки ученика; используется для автоматической смены статуса",
        "1357",
    ),
    (
        "last_name",
        "Фамилия ребенка",
        True,
        "Фамилия полностью",
        "Писарева",
    ),
    (
        "first_name",
        "Имя ребенка",
        True,
        "Имя полностью",
        "Есения",
    ),
    (
        "birth_date",
        "Дата рождения",
        False,
        "Дата рождения ученика в формате ДД.ММ.ГГГГ",
        "15.04.2015",
    ),
    (
        "group_name",
        "Группа",
        True,
        "Единое точное название группы",
        "1 3D общ вт 17:30 (год 25-26)",
    ),
    (
        "course_name",
        "Курс",
        True,
        "Название учебного курса",
        "Minecraft",
    ),
    (
        "venue_name",
        "Площадка",
        True,
        "Адрес или название филиала",
        "Гагарина 64",
    ),
    (
        "teacher_name",
        "Преподаватель",
        True,
        "ФИО преподавателя как в его профиле",
        "Дворянинова Вероника",
    ),
    (
        "contact_ids",
        "ID контакта",
        False,
        "Contact ID родителя из amoCRM, не MAX ID",
        "30420713",
    ),
    (
        "contact_names",
        "Имя контакта",
        False,
        "ФИО родителя",
        "Никина Елена Александровна",
    ),
)
CRM_TEMPLATE_FIELDS = frozenset(column[0] for column in CRM_TEMPLATE_COLUMNS)
CRM_TEMPLATE_REQUIRED_FIELDS = frozenset(
    column[0] for column in CRM_TEMPLATE_COLUMNS if column[2]
)
CRM_TEMPLATE_HEADER_BY_FIELD = {
    column[0]: column[1] for column in CRM_TEMPLATE_COLUMNS
}


class CrmImportError(RuntimeError):
    pass


@dataclass(frozen=True)
class CrmStudentRow:
    row_number: int
    deal_id: str | None
    uuid: str | None
    lms_student_id: str | None
    first_name: str | None
    last_name: str | None
    group_name: str | None
    course_name: str | None
    venue_name: str | None
    teacher_name: str | None
    city: str | None
    status_name: str | None
    contact_ids: str | None
    contact_names: str | None
    birth_date: date | None = None

    @property
    def needs_generated_access_code(self) -> bool:
        return not self.contact_ids


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "deal_id": (
        "id \u0441\u0434\u0435\u043b\u043a\u0438",
        "id \u0441\u0434\u0435\u043b\u043a\u0438 amocrm",
        "id \u0441\u0434\u0435\u043b\u043a\u0438 amo crm",
        "id",
        "\u043d\u043e\u043c\u0435\u0440 \u0441\u0434\u0435\u043b\u043a\u0438",
    ),
    "uuid": ("uuid",),
    "lms_student_id": (
        "id \u0443\u0447\u0435\u043d\u0438\u043a\u0430",
        "id \u0440\u0435\u0431\u0435\u043d\u043a\u0430",
        "id \u0440\u0435\u0431\u0451\u043d\u043a\u0430",
        "student id",
        "lms student id",
        "backoffice_id",
    ),
    "first_name": (
        "\u0438\u043c\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430 \u0438\u0437 lms",
        "\u0438\u043c\u044f \u0440\u0435\u0431\u0451\u043d\u043a\u0430 \u0438\u0437 lms",
        "\u0438\u043c\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430",
        "\u0438\u043c\u044f \u0440\u0435\u0431\u0451\u043d\u043a\u0430",
    ),
    "birth_date": (
        "дата рождения",
        "день рождения",
        "дата рождения ребенка",
        "дата рождения ребёнка",
        "birth date",
    ),
    "last_name": (
        "\u0444\u0430\u043c\u0438\u043b\u0438\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430 \u0438\u0437 lms",  # noqa: E501
        "\u0444\u0430\u043c\u0438\u043b\u0438\u044f \u0440\u0435\u0431\u0451\u043d\u043a\u0430 \u0438\u0437 lms",  # noqa: E501
        "\u0444\u0430\u043c\u0438\u043b\u0438\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430",
        "\u0444\u0430\u043c\u0438\u043b\u0438\u044f \u0440\u0435\u0431\u0451\u043d\u043a\u0430",
    ),
    "group_name": (
        "\u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u0433\u0440\u0443\u043f\u043f\u044b lms",
        "\u0433\u0440\u0443\u043f\u043f\u0430 lms",
        "\u0433\u0440\u0443\u043f\u043f\u0430",
    ),
    "course_name": (
        "\u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043a\u0443\u0440\u0441\u0430 \u0438\u0437 lms",  # noqa: E501
        "\u043a\u0443\u0440\u0441",
        "\u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043a\u0443\u0440\u0441\u0430",
    ),
    "venue_name": (
        "\u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u043b\u043e\u0449\u0430\u0434\u043a\u0438",  # noqa: E501
        "\u043f\u043b\u043e\u0449\u0430\u0434\u043a\u0430",
        "\u043f\u043b\u043e\u0449\u0430\u0434\u043a\u0430 lms",
    ),
    "teacher_name": (
        "\u043f\u0440\u0435\u043f\u043e\u0434\u0430\u0432\u0430\u0442\u0435\u043b\u044c",
        "\u043f\u0435\u0434\u0430\u0433\u043e\u0433",
    ),
    "city": ("\u0433\u043e\u0440\u043e\u0434",),
    "status_name": (
        "\u0438\u043c\u044f \u0441\u0442\u0430\u0442\u0443\u0441\u0430",
        "\u0441\u0442\u0430\u0442\u0443\u0441",
        "\u0441\u0442\u0430\u0434\u0438\u044f",
    ),
    "contact_ids": (
        "id \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u043e\u0432",
        "id \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u0430",
        "contact id",
        "contact ids",
    ),
    "contact_names": (
        "\u0438\u043c\u0435\u043d\u0430 \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u043e\u0432",
        "\u043a\u043e\u043d\u0442\u0430\u043a\u0442\u044b",
        "\u0438\u043c\u044f \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u0430",
    ),
}


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return " ".join(text.split())


def normalize_header(value: Any) -> str:
    text = normalize_text(value) or ""
    return text.lower().replace("\u0451", "\u0435")


def parse_birth_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = normalize_text(value)
        if not text:
            return None
        parsed = None
        for date_format in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                parsed = datetime.strptime(text, date_format).date()
                break
            except ValueError:
                continue
        if parsed is None:
            raise ValueError("use DD.MM.YYYY")
    if parsed > date.today():
        raise ValueError("birth date is in the future")
    return parsed


def find_header_mapping(
    values: tuple[Any, ...],
    *,
    allowed_fields: frozenset[str] | None = None,
) -> dict[int, str]:
    mapping: dict[int, str] = {}
    known_headers = {
        normalize_header(alias): field
        for field, aliases in HEADER_ALIASES.items()
        if allowed_fields is None or field in allowed_fields
        for alias in aliases
    }

    for index, value in enumerate(values):
        header = normalize_header(value)
        if header in known_headers:
            mapping[index] = known_headers[header]

    return mapping


def build_crm_import_template() -> bytes:
    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "Инструкция"
    instruction.sheet_view.showGridLines = False
    instruction.merge_cells("A1:D1")
    instruction["A1"] = "Шаблон импорта Algo MAX"
    instruction["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    instruction["A1"].fill = PatternFill("solid", fgColor="6F30D0")
    instruction["A1"].alignment = Alignment(vertical="center")
    instruction.row_dimensions[1].height = 34

    instruction.merge_cells("A3:D3")
    instruction["A3"] = (
        "Заполняйте только лист «Шаблон». Одна строка = один ученик. "
        "Не меняйте названия колонок."
    )
    instruction["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    instruction.row_dimensions[3].height = 34

    instruction.merge_cells("A4:D4")
    instruction["A4"] = (
        "Дополнительные листы и колонки игнорируются. Один ID контакта можно повторить "
        "у нескольких детей одного родителя."
    )
    instruction["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    instruction.row_dimensions[4].height = 34

    instruction_headers = (
        "Колонка",
        "Обязательно",
        "Что указывать",
        "Пример",
    )
    for column_index, value in enumerate(instruction_headers, start=1):
        cell = instruction.cell(row=6, column=column_index, value=value)
        cell.font = Font(bold=True, color="24152F")
        cell.fill = PatternFill("solid", fgColor="FFD43B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_index, (_, header, required, description, example) in enumerate(
        CRM_TEMPLATE_COLUMNS,
        start=7,
    ):
        values = (
            header,
            "Да" if required else "Нет",
            description,
            example,
        )
        for column_index, value in enumerate(values, start=1):
            cell = instruction.cell(row=row_index, column=column_index, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_index % 2:
                cell.fill = PatternFill("solid", fgColor="F5F0FF")

    instruction.freeze_panes = "A7"
    instruction.auto_filter.ref = f"A6:D{6 + len(CRM_TEMPLATE_COLUMNS)}"
    for column_letter, width in {"A": 24, "B": 16, "C": 58, "D": 38}.items():
        instruction.column_dimensions[column_letter].width = width

    template = workbook.create_sheet(CRM_TEMPLATE_SHEET_NAME)
    template.sheet_view.showGridLines = False
    template.append([column[1] for column in CRM_TEMPLATE_COLUMNS])
    for cell in template[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6F30D0")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    template.row_dimensions[1].height = 34
    template.freeze_panes = "A2"
    last_column = template.cell(
        row=1,
        column=len(CRM_TEMPLATE_COLUMNS),
    ).column_letter
    template.auto_filter.ref = f"A1:{last_column}1"
    template_widths = (18, 22, 24, 22, 18, 38, 22, 28, 34, 20, 34)
    for column_index, width in enumerate(template_widths, start=1):
        template.column_dimensions[
            template.cell(row=1, column=column_index).column_letter
        ].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_crm_students(
    path: str | Path,
    *,
    sheet_name: str = CRM_TEMPLATE_SHEET_NAME,
) -> list[CrmStudentRow]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        raise CrmImportError(f"CRM file not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return parse_crm_workbook(workbook, sheet_name=sheet_name)
    finally:
        workbook.close()


def parse_crm_students_content(
    content: bytes,
    *,
    sheet_name: str = CRM_TEMPLATE_SHEET_NAME,
) -> list[CrmStudentRow]:
    if not content:
        raise CrmImportError("CRM file is empty")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CrmImportError("CRM file is not a readable XLSX workbook") from exc
    try:
        return parse_crm_workbook(workbook, sheet_name=sheet_name)
    finally:
        workbook.close()


def parse_crm_workbook(
    workbook: Any,
    *,
    sheet_name: str,
) -> list[CrmStudentRow]:
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        sheet = next(
            (
                workbook[candidate_name]
                for candidate_name in (CRM_TEMPLATE_SHEET_NAME, CRM_LEGACY_SHEET_NAME)
                if candidate_name in workbook.sheetnames
            ),
            None,
        )
        if sheet is None:
            best_score = 0
            for candidate in workbook.worksheets:
                score = max(
                    (
                        len(find_header_mapping(tuple(values)))
                        for _, values in zip(
                            range(30),
                            candidate.iter_rows(values_only=True),
                            strict=False,
                        )
                    ),
                    default=0,
                )
                if score > best_score:
                    best_score = score
                    sheet = candidate
            if sheet is None or best_score < 2:
                raise CrmImportError(f"Sheet not found: {sheet_name}")

    template_format = sheet.title == CRM_TEMPLATE_SHEET_NAME
    allowed_fields = CRM_TEMPLATE_FIELDS if template_format else None

    header_mapping: dict[int, str] | None = None
    result: list[CrmStudentRow] = []
    template_student_rows: dict[str, int] = {}

    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        if header_mapping is None:
            candidate = find_header_mapping(values, allowed_fields=allowed_fields)
            if len(candidate) >= 3:
                header_mapping = candidate
                if template_format:
                    mapped_fields = frozenset(candidate.values())
                    missing_fields = CRM_TEMPLATE_REQUIRED_FIELDS - mapped_fields
                    if missing_fields:
                        missing_headers = [
                            CRM_TEMPLATE_HEADER_BY_FIELD[field]
                            for field in CRM_TEMPLATE_HEADER_BY_FIELD
                            if field in missing_fields
                        ]
                        raise CrmImportError(
                            "В листе «Шаблон» нет обязательных колонок: "
                            + ", ".join(missing_headers)
                        )
            continue

        raw_row = {
            field: values[index]
            for index, field in header_mapping.items()
        }
        row = {field: normalize_text(value) for field, value in raw_row.items()}
        first_name = row.get("first_name")
        last_name = row.get("last_name")
        group_name = row.get("group_name")
        deal_id = row.get("deal_id")
        contact_ids = row.get("contact_ids")

        if not any([first_name, last_name, group_name, deal_id, contact_ids]):
            continue

        try:
            birth_date = parse_birth_date(raw_row.get("birth_date"))
        except ValueError as exc:
            raise CrmImportError(
                f"Строка {row_number}: дата рождения должна быть в формате ДД.ММ.ГГГГ"
            ) from exc

        if template_format:
            missing_fields = [
                field
                for field in CRM_TEMPLATE_HEADER_BY_FIELD
                if field in CRM_TEMPLATE_REQUIRED_FIELDS and not row.get(field)
            ]
            if missing_fields:
                missing_headers = [
                    CRM_TEMPLATE_HEADER_BY_FIELD[field] for field in missing_fields
                ]
                raise CrmImportError(
                    f"Строка {row_number}: заполните обязательные поля: "
                    + ", ".join(missing_headers)
                )

            normalized_student_id = "".join(
                (row["lms_student_id"] or "").upper().split()
            )
            previous_row = template_student_rows.get(normalized_student_id)
            if previous_row is not None:
                raise CrmImportError(
                    f"Строки {previous_row} и {row_number}: повторяется ID ученика "
                    f"«{row['lms_student_id']}»"
                )
            template_student_rows[normalized_student_id] = row_number

        result.append(
            CrmStudentRow(
                row_number=row_number,
                deal_id=deal_id,
                uuid=row.get("uuid"),
                lms_student_id=row.get("lms_student_id"),
                first_name=first_name,
                last_name=last_name,
                group_name=group_name,
                course_name=row.get("course_name"),
                venue_name=row.get("venue_name"),
                teacher_name=row.get("teacher_name"),
                city=row.get("city"),
                status_name=row.get("status_name"),
                contact_ids=contact_ids,
                contact_names=row.get("contact_names"),
                birth_date=birth_date,
            )
        )

    if header_mapping is None:
        raise CrmImportError("CRM header row was not found")

    return result
