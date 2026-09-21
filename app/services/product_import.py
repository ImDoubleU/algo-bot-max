from __future__ import annotations

import asyncio
import csv
import io
import re
from dataclasses import dataclass, field, replace
from decimal import Decimal, InvalidOperation
from typing import Any
from urllib.parse import urlsplit
from uuid import UUID, uuid4
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.enums import ProductStatus, StockMovementType, WarehouseType
from app.models.store import Product, ProductCategory, Warehouse, WarehouseInventory
from app.models.tenant import Tenant
from app.services.crm_sync import slugify
from app.services.product_media import (
    ProductMediaError,
    SavedProductImage,
    create_remote_product_image_client,
    remove_product_image,
    save_remote_product_image,
)
from app.services.warehouse import build_stock_movement
from app.services.warehouse_access import (
    ensure_warehouse_link,
    find_partner_warehouse_by_name,
    normalize_warehouse_name,
)

HEADER_ALIASES = {
    "sku": "sku",
    "артикул": "sku",
    "код": "sku",
    "product_sku": "sku",
    "name": "name",
    "название": "name",
    "товар": "name",
    "product_name": "name",
    "category": "category_name",
    "категория": "category_name",
    "category_name": "category_name",
    "category_slug": "category_slug",
    "price": "price_astrocoins",
    "цена": "price_astrocoins",
    "price_astrocoins": "price_astrocoins",
    "астрокоины": "price_astrocoins",
    "цена ac": "price_astrocoins",
    "quantity": "quantity",
    "остаток": "quantity",
    "количество": "quantity",
    "stock": "quantity",
    "warehouse": "warehouse_name",
    "склад": "warehouse_name",
    "warehouse_name": "warehouse_name",
    "warehouse_slug": "warehouse_slug",
    "description": "description",
    "описание": "description",
    "photo_url": "photo_url",
    "image_url": "photo_url",
    "фото": "photo_url",
    "фото url": "photo_url",
    "ссылка на фото": "photo_url",
    "ссылка на изображение": "photo_url",
    "изображение": "photo_url",
    "status": "status",
    "статус": "status",
}

PRODUCT_TEMPLATE_SHEET_NAME = "Товары"
PRODUCT_TEMPLATE_COLUMNS = (
    "Артикул",
    "Название",
    "Категория",
    "Цена AC",
    "Склад",
    "Остаток",
    "Статус",
    "Описание",
    "Ссылка на фото",
)
MAX_CONCURRENT_PRODUCT_IMAGE_DOWNLOADS = 4


class ProductImportError(RuntimeError):
    pass


@dataclass
class ProductImportResult:
    tenant_slug: str
    created_products: int = 0
    updated_products: int = 0
    created_categories: int = 0
    created_warehouses: int = 0
    updated_inventory: int = 0
    skipped_rows: int = 0
    errors: list[str] = field(default_factory=list)
    new_active_products: list[Product] = field(default_factory=list, repr=False)
    obsolete_photo_urls: list[str] = field(default_factory=list, repr=False)
    low_stock_items: list[tuple[Product, WarehouseInventory]] = field(
        default_factory=list,
        repr=False,
    )


@dataclass(frozen=True)
class ProductImportRow:
    row_number: int
    sku: str | None
    name: str
    category_name: str
    category_slug: str
    price_astrocoins: int
    quantity: int
    warehouse_name: str
    warehouse_slug: str
    description: str | None = None
    photo_url: str | None = None
    photo_thumbnail_url: str | None = None
    photo_master_url: str | None = None
    status: ProductStatus = ProductStatus.ACTIVE


def parse_product_rows(filename: str, content: bytes) -> list[ProductImportRow]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in {"csv", "xlsx"}:
        raise ProductImportError("Поддерживаются только файлы CSV и XLSX")
    try:
        raw_rows = _read_xlsx(content) if suffix == "xlsx" else _read_csv(content)
    except (BadZipFile, InvalidFileException, OSError, ValueError, csv.Error) as exc:
        raise ProductImportError(f"Не удалось прочитать файл {suffix.upper()}") from exc
    return _normalize_rows(raw_rows)


def _read_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [dict(row) for row in reader]


def _read_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = (
            workbook[PRODUCT_TEMPLATE_SHEET_NAME]
            if PRODUCT_TEMPLATE_SHEET_NAME in workbook.sheetnames
            else workbook.active
        )
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(value or "").strip() for value in rows[0]]
        result: list[dict[str, Any]] = []
        for values in rows[1:]:
            result.append(
                {
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                }
            )
        return result
    finally:
        workbook.close()


def _normalize_header(value: str) -> str:
    return HEADER_ALIASES.get(value.strip().lower(), value.strip().lower())


def _cell(row: dict[str, Any], key: str) -> Any:
    normalized = {_normalize_header(str(header)): value for header, value in row.items()}
    return normalized.get(key)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_product_category_name(value: Any) -> str:
    normalized = re.sub(r"\s+", " ", _text(value)).strip()
    normalized = re.sub(r"\s*/\s*", " / ", normalized)
    return normalized or "Без категории"


def product_category_key(value: Any) -> str:
    return normalize_product_category_name(value).casefold()


def _int(value: Any, *, default: int | None = None) -> int:
    text = _text(value).replace(" ", "").replace(",", ".")
    if not text:
        if default is None:
            raise ValueError("empty number")
        return default
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"неверное число: {value}") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"ожидается целое число: {value}")
    return int(number)


def _status(value: Any) -> ProductStatus:
    text = _text(value).lower()
    if text in {"", "active", "активен", "активный", "да", "1"}:
        return ProductStatus.ACTIVE
    if text in {"hidden", "скрыт", "скрытый", "нет", "0"}:
        return ProductStatus.HIDDEN
    if text in {"archived", "архив", "архивный"}:
        return ProductStatus.ARCHIVED
    raise ValueError(f"unknown status: {value}")


def _photo_url(value: Any) -> str | None:
    photo_url = _text(value)
    if not photo_url:
        return None
    if len(photo_url) > 500:
        raise ValueError("ссылка на фото длиннее 500 символов")
    try:
        parsed = urlsplit(photo_url)
    except ValueError as exc:
        raise ValueError("некорректная ссылка на фото") from exc
    if parsed.scheme.casefold() not in {"http", "https"} or not parsed.netloc:
        raise ValueError("ссылка на фото должна начинаться с http:// или https://")
    return photo_url


def _normalize_rows(raw_rows: list[dict[str, Any]]) -> list[ProductImportRow]:
    rows: list[ProductImportRow] = []
    seen_inventory_rows: dict[tuple[str, str], int] = {}
    for index, raw_row in enumerate(raw_rows, start=2):
        sku = _text(_cell(raw_row, "sku")).upper()
        name = _text(_cell(raw_row, "name"))
        if not sku and not name:
            continue
        if not name:
            raise ProductImportError(f"Строка {index}: не указано название товара")

        category_name = normalize_product_category_name(_cell(raw_row, "category_name"))
        category_slug = slugify(_text(_cell(raw_row, "category_slug")) or category_name)
        warehouse_name = _text(_cell(raw_row, "warehouse_name")) or "Общий склад"
        warehouse_slug = _text(_cell(raw_row, "warehouse_slug")) or slugify(warehouse_name)
        product_key = sku or f"{category_slug}:{name.casefold()}"
        inventory_key = (product_key, warehouse_slug)
        if inventory_key in seen_inventory_rows:
            first_row = seen_inventory_rows[inventory_key]
            product_label = sku or f"«{name}»"
            raise ProductImportError(
                f"Строки {first_row} и {index}: повторяется товар {product_label} на складе "
                f"«{warehouse_name}»"
            )
        seen_inventory_rows[inventory_key] = index

        try:
            price = _int(_cell(raw_row, "price_astrocoins"))
            quantity = _int(_cell(raw_row, "quantity"), default=0)
            status = _status(_cell(raw_row, "status"))
            photo_url = _photo_url(_cell(raw_row, "photo_url"))
        except ValueError as exc:
            raise ProductImportError(f"Строка {index}: {exc}") from exc

        if price < 0:
            raise ProductImportError(f"Строка {index}: цена не может быть отрицательной")
        if quantity < 0:
            raise ProductImportError(f"Строка {index}: остаток не может быть отрицательным")

        rows.append(
            ProductImportRow(
                row_number=index,
                sku=sku,
                name=name,
                category_name=category_name,
                category_slug=category_slug,
                price_astrocoins=price,
                quantity=quantity,
                warehouse_name=warehouse_name,
                warehouse_slug=warehouse_slug,
                description=_text(_cell(raw_row, "description")) or None,
                photo_url=photo_url,
                status=status,
            )
        )
    return rows


def generate_product_sku() -> str:
    return f"PRD-{uuid4().hex[:12].upper()}"


async def localize_product_import_photos(
    rows: list[ProductImportRow],
    *,
    media_root: str,
    media_base_url: str,
) -> tuple[list[ProductImportRow], list[SavedProductImage]]:
    saved_images: list[SavedProductImage] = []
    localized_images: dict[tuple[str, str], SavedProductImage] = {}
    base_url = media_base_url.rstrip("/")
    jobs: dict[tuple[str, str], ProductImportRow] = {}
    for row in rows:
        if not row.photo_url:
            continue
        product_key = row.sku or f"{row.category_slug}:{row.name.casefold()}"
        jobs.setdefault((product_key, row.photo_url), row)

    semaphore = asyncio.Semaphore(MAX_CONCURRENT_PRODUCT_IMAGE_DOWNLOADS)

    async def download(
        cache_key: tuple[str, str],
        row: ProductImportRow,
        *,
        client,
    ) -> tuple[tuple[str, str], ProductImportRow, SavedProductImage | None, Exception | None]:
        try:
            async with semaphore:
                image = await save_remote_product_image(
                    row.photo_url or "",
                    media_root=media_root,
                    client=client,
                )
            return cache_key, row, image, None
        except Exception as exc:
            return cache_key, row, None, exc

    if jobs:
        async with create_remote_product_image_client() as client:
            outcomes = await asyncio.gather(
                *(
                    download(cache_key, row, client=client)
                    for cache_key, row in jobs.items()
                )
            )

        failures: list[tuple[ProductImportRow, Exception]] = []
        for cache_key, row, saved_image, error in outcomes:
            if error is not None:
                failures.append((row, error))
                continue
            if saved_image is None:
                failures.append((row, ProductMediaError("Файл изображения не сохранен")))
                continue
            saved_images.append(saved_image)
            localized_images[cache_key] = saved_image

        if failures:
            await asyncio.gather(*(remove_product_image(image) for image in saved_images))
            failed_row, error = failures[0]
            if isinstance(error, ProductMediaError):
                raise ProductImportError(
                    f"Строка {failed_row.row_number}: не удалось загрузить фото: {error}"
                ) from error
            raise error

    localized_rows: list[ProductImportRow] = []
    for row in rows:
        if not row.photo_url:
            localized_rows.append(row)
            continue
        product_key = row.sku or f"{row.category_slug}:{row.name.casefold()}"
        saved_image = localized_images[(product_key, row.photo_url)]
        detail_url = f"{base_url}{saved_image.url_path}"
        localized_rows.append(
            replace(
                row,
                photo_url=detail_url,
                photo_thumbnail_url=(
                    f"{base_url}{saved_image.thumbnail_url_path}"
                    if saved_image.thumbnail_url_path
                    else detail_url
                ),
                photo_master_url=(
                    f"{base_url}{saved_image.master_url_path}"
                    if saved_image.master_url_path
                    else detail_url
                ),
            )
        )
    return localized_rows, saved_images


def build_product_import_template() -> bytes:
    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "Инструкция"
    instruction.append(["Шаблон массовой загрузки товаров"])
    instruction.append([])
    instruction.append(["1. Заполняйте только лист «Товары»."])
    instruction.append(
        [
            "2. Обязательные поля: название и цена. Если склад не указан, используется "
            "«Общий склад», остаток будет равен 0."
        ]
    )
    instruction.append(
        [
            "3. Артикул необязателен, но рекомендуется: по нему система надежно обновляет "
            "уже загруженный товар."
        ]
    )
    instruction.append(
        [
            "4. Для нескольких складов повторите артикул и данные товара в нескольких "
            "строках, меняя склад и остаток."
        ]
    )
    instruction.append(
        [
            "5. В поле «Ссылка на фото» укажите прямой адрес изображения либо публичную "
            "ссылку Яндекс Диска или Google Drive. Для Google Drive включите доступ "
            "«Все, у кого есть ссылка»."
        ]
    )
    instruction.append(["6. Статус: Активен, Скрыт или Архив. Пустое значение означает «Активен»."])
    instruction.append(
        [
            "7. Заполняйте лист «Товары». Листы «Инструкция» и «Пример» при импорте "
            "не обрабатываются."
        ]
    )
    instruction.column_dimensions["A"].width = 105
    instruction["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    instruction["A1"].fill = PatternFill("solid", fgColor="6F35D5")
    instruction["A1"].alignment = Alignment(vertical="center")
    instruction.row_dimensions[1].height = 28

    sheet = workbook.create_sheet(PRODUCT_TEMPLATE_SHEET_NAME)
    sheet.append(list(PRODUCT_TEMPLATE_COLUMNS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I1"
    widths = (20, 32, 24, 14, 28, 14, 16, 48, 56)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6F35D5")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    comments = {
        1: (
            "Необязательно, но рекомендуется. Один и тот же товар на разных складах "
            "должен иметь одинаковый артикул."
        ),
        2: "Обязательное поле.",
        3: "Если оставить пустым, будет создана категория «Без категории».",
        4: "Обязательное целое число, не меньше 0.",
        5: "Если оставить пустым, будет использован «Общий склад».",
        6: "Целое число, не меньше 0. По умолчанию 0.",
        7: "Активен, Скрыт или Архив. По умолчанию Активен.",
        8: "Необязательное описание товара.",
        9: (
            "Прямая ссылка на изображение или публичная ссылка Яндекс Диска/Google Drive "
            "http:// или https:// длиной до 500 символов."
        ),
    }
    for column, comment in comments.items():
        sheet.cell(row=1, column=column).comment = Comment(comment, "Algo MAX")

    status_validation = DataValidation(
        type="list",
        formula1='"Активен,Скрыт,Архив"',
        allow_blank=True,
    )
    status_validation.error = "Выберите: Активен, Скрыт или Архив"
    status_validation.errorTitle = "Некорректный статус"
    sheet.add_data_validation(status_validation)
    status_validation.add("G2:G5000")

    example = workbook.create_sheet("Пример")
    example.append(list(PRODUCT_TEMPLATE_COLUMNS))
    example.append(
        [
            "ROBOT-KIT-01",
            "Набор для робототехники",
            "Учебные наборы",
            750,
            "Главный склад",
            12,
            "Активен",
            "Набор для учебных проектов",
            "https://images.example.org/products/robot-kit.jpg",
        ]
    )
    example.append(
        [
            "YANDEX-PHOTO-01",
            "Товар с фото из Яндекс Диска",
            "Примеры ссылок",
            100,
            "Главный склад",
            1,
            "Активен",
            "Используйте стабильную ссылку, полученную через «Поделиться»",
            "https://disk.yandex.ru/i/PUBLIC_LINK_ID",
        ]
    )
    example.append(
        [
            "GOOGLE-PHOTO-01",
            "Товар с фото из Google Drive",
            "Примеры ссылок",
            100,
            "Главный склад",
            1,
            "Активен",
            "Откройте доступ «Все, у кого есть ссылка»",
            "https://drive.google.com/file/d/FILE_ID/view?usp=sharing",
        ]
    )
    example.append(
        [
            "ROBOT-KIT-01",
            "Набор для робототехники",
            "Учебные наборы",
            750,
            "Склад площадки",
            4,
            "Активен",
            "Набор для учебных проектов",
            "https://images.example.org/products/robot-kit.jpg",
        ]
    )
    example.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        example.column_dimensions[chr(64 + index)].width = width
        cell = example.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6F35D5")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


async def import_products_for_tenant(
    db: AsyncSession,
    *,
    tenant: Tenant,
    rows: list[ProductImportRow],
    actor_account_id: UUID | None = None,
) -> ProductImportResult:
    result = ProductImportResult(tenant_slug=tenant.slug)
    categories: dict[tuple[UUID, str], ProductCategory] = {}
    warehouses: dict[str, Warehouse] = {}

    for row in rows:
        warehouse_key = normalize_warehouse_name(row.warehouse_name)
        warehouse = warehouses.get(warehouse_key)
        if warehouse is None:
            warehouse, created_warehouse = await _get_or_create_warehouse(
                db,
                tenant=tenant,
                row=row,
            )
            warehouses[warehouse_key] = warehouse
            if created_warehouse:
                result.created_warehouses += 1

        product = await _find_import_product(
            db,
            tenant=tenant,
            warehouse=warehouse,
            row=row,
        )
        product_tenant_id = UUID(str(product.tenant_id if product is not None else tenant.id))
        category_key = (product_tenant_id, row.category_slug)
        category = categories.get(category_key)
        if category is None:
            category = await _get_or_create_category(
                db,
                tenant_id=product_tenant_id,
                row=row,
            )
            categories[category_key] = category

        category_created = getattr(category, "_import_created", False)
        if category_created:
            result.created_categories += 1
            category._import_created = False

        product_created = product is None
        if product_created:
            product = Product(
                tenant_id=tenant.id,
                category_id=category.id,
                sku=row.sku or generate_product_sku(),
                name=row.name,
                description=row.description,
                photo_url=row.photo_url,
                photo_thumbnail_url=row.photo_thumbnail_url,
                photo_master_url=row.photo_master_url,
                price_astrocoins=row.price_astrocoins,
                status=row.status,
            )
            db.add(product)
            await db.flush()
            result.created_products += 1
            if product.status == ProductStatus.ACTIVE:
                result.new_active_products.append(product)
        else:
            next_photo_urls = {
                row.photo_url,
                row.photo_thumbnail_url,
                row.photo_master_url,
            }
            for previous_url in (
                product.photo_url,
                product.photo_thumbnail_url,
                product.photo_master_url,
            ):
                if previous_url and previous_url not in next_photo_urls:
                    result.obsolete_photo_urls.append(previous_url)
            product.category_id = category.id
            product.name = row.name
            product.description = row.description
            product.photo_url = row.photo_url
            product.photo_thumbnail_url = row.photo_thumbnail_url
            product.photo_master_url = row.photo_master_url
            product.photo_crop_x = None
            product.photo_crop_y = None
            product.photo_crop_width = None
            product.photo_crop_height = None
            product.price_astrocoins = row.price_astrocoins
            product.status = row.status
            result.updated_products += 1

        inventory = await db.scalar(
            select(WarehouseInventory)
            .where(
                WarehouseInventory.warehouse_id == warehouse.id,
                WarehouseInventory.product_id == product.id,
            )
            .with_for_update()
            .options(
                selectinload(WarehouseInventory.product),
                selectinload(WarehouseInventory.warehouse),
            )
        )
        if inventory is None:
            inventory = WarehouseInventory(
                tenant_id=product.tenant_id,
                warehouse_id=warehouse.id,
                product_id=product.id,
            )
            inventory.product = product
            inventory.warehouse = warehouse
            db.add(inventory)

        reserved_quantity = int(inventory.reserved_quantity or 0)
        if row.quantity < reserved_quantity:
            raise ProductImportError(
                f"Строка {row.row_number}: остаток {row.quantity} меньше уже "
                f"зарезервированного количества {reserved_quantity}"
            )
        previous_quantity = int(inventory.available_quantity or 0)
        inventory.is_active = True
        inventory.available_quantity = row.quantity
        quantity_delta = row.quantity - previous_quantity
        if quantity_delta:
            db.add(
                build_stock_movement(
                    inventory=inventory,
                    movement_type=StockMovementType.ADJUSTMENT,
                    quantity=abs(quantity_delta),
                    actor_account_id=actor_account_id,
                    from_warehouse_id=warehouse.id if quantity_delta < 0 else None,
                    to_warehouse_id=warehouse.id if quantity_delta > 0 else None,
                    comment=(
                        f"Импорт товаров: остаток изменен с {previous_quantity} "
                        f"на {row.quantity}"
                    ),
                )
            )
        free_quantity = row.quantity - reserved_quantity
        if free_quantity <= 5 and not inventory.low_stock_notified:
            inventory.low_stock_notified = True
            result.low_stock_items.append((product, inventory))
        elif free_quantity > 5:
            inventory.low_stock_notified = False
        result.updated_inventory += 1

    return result


async def _get_or_create_category(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    row: ProductImportRow,
) -> ProductCategory:
    category, created = await get_or_create_product_category(
        db,
        tenant_id=tenant_id,
        category_name=row.category_name,
        category_slug=row.category_slug,
    )
    category._import_created = created
    return category


async def get_or_create_product_category(
    db: AsyncSession,
    *,
    tenant_id: UUID,
    category_name: str,
    category_slug: str | None = None,
) -> tuple[ProductCategory, bool]:
    canonical_name = normalize_product_category_name(category_name)
    canonical_slug = slugify(_text(category_slug) or canonical_name)
    categories = list(
        await db.scalars(
            select(ProductCategory)
            .where(ProductCategory.tenant_id == tenant_id)
            .order_by(ProductCategory.sort_order, ProductCategory.id)
        )
    )
    category = next(
        (item for item in categories if item.slug == canonical_slug),
        None,
    )
    if category is None:
        canonical_key = product_category_key(canonical_name)
        category = next(
            (item for item in categories if product_category_key(item.name) == canonical_key),
            None,
        )
    if category is not None:
        category.name = canonical_name
        return category, False

    category = ProductCategory(
        tenant_id=tenant_id,
        slug=canonical_slug,
        name=canonical_name,
        sort_order=100,
    )
    db.add(category)
    await db.flush()
    return category, True


async def _find_import_product(
    db: AsyncSession,
    *,
    tenant: Tenant,
    warehouse: Warehouse,
    row: ProductImportRow,
) -> Product | None:
    own_products = list(
        await db.scalars(
            select(Product)
            .where(Product.tenant_id == tenant.id)
            .options(selectinload(Product.category))
            .order_by(Product.created_at, Product.id)
        )
    )

    def matches(product: Product) -> bool:
        if row.sku:
            return product.sku == row.sku
        return (
            normalize_warehouse_name(product.name) == normalize_warehouse_name(row.name)
            and product.category is not None
            and product_category_key(product.category.name)
            == product_category_key(row.category_name)
        )

    own_product = next((product for product in own_products if matches(product)), None)
    if own_product is not None:
        return own_product

    shared_products = (
        (
            await db.scalars(
                select(Product)
                .join(WarehouseInventory, WarehouseInventory.product_id == Product.id)
                .where(WarehouseInventory.warehouse_id == warehouse.id)
                .options(selectinload(Product.category))
                .order_by(Product.created_at, Product.id)
            )
        )
        .unique()
        .all()
    )
    return next((product for product in shared_products if matches(product)), None)


async def _get_or_create_warehouse(
    db: AsyncSession,
    *,
    tenant: Tenant,
    row: ProductImportRow,
) -> tuple[Warehouse, bool]:
    warehouse = await find_partner_warehouse_by_name(
        db,
        tenant=tenant,
        name=row.warehouse_name,
    )
    if warehouse is not None:
        await ensure_warehouse_link(
            db,
            tenant_id=UUID(str(tenant.id)),
            warehouse_id=UUID(str(warehouse.id)),
        )
        return warehouse, False

    warehouse = Warehouse(
        tenant_id=tenant.id,
        slug=row.warehouse_slug,
        name=row.warehouse_name,
        warehouse_type=WarehouseType.COMMON,
    )
    db.add(warehouse)
    await db.flush()
    await ensure_warehouse_link(
        db,
        tenant_id=UUID(str(tenant.id)),
        warehouse_id=UUID(str(warehouse.id)),
    )
    return warehouse, True
