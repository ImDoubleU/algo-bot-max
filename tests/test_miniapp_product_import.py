import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app.db.base  # noqa: F401
from app.models.account import MaxAccount, StaffRoleAssignment
from app.models.base import Base
from app.models.enums import AssignmentStatus, StaffRole, WarehouseType
from app.models.store import (
    Product,
    ProductCategory,
    Warehouse,
    WarehouseInventory,
    WarehouseTenantLink,
)
from app.models.tenant import City, Partner, Tenant
from app.services.miniapp import import_miniapp_products
from app.services.product_import import (
    ProductImportError,
    build_product_import_template,
    parse_product_rows,
)
from app.services.product_media import SavedProductImage


@pytest.fixture
async def db_session():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with session_factory() as session:
        yield session
    await engine.dispose()


async def seed_admin(db_session) -> None:
    city = City(slug="nizhniy-novgorod", name="Нижний Новгород")
    partner = Partner(slug="partner-a", name="Партнер A")
    db_session.add_all([city, partner])
    await db_session.flush()

    tenant = Tenant(
        city_id=city.id,
        partner_id=partner.id,
        slug="nizhniy-novgorod-partner-a",
        name="Нижний Новгород / Партнер A",
    )
    account = MaxAccount(max_user_id=53364725, username="admin_user")
    db_session.add_all([tenant, account])
    await db_session.flush()

    db_session.add(
        StaffRoleAssignment(
            tenant_id=tenant.id,
            account_id=account.id,
            role=StaffRole.ADMIN,
            status=AssignmentStatus.ACTIVE,
        )
    )
    await db_session.commit()


async def test_admin_imports_products_from_miniapp_csv(db_session, tmp_path) -> None:
    await seed_admin(db_session)
    content = (
        "sku,name,category,price_astrocoins,quantity,warehouse\n"
        "PEN-LOGO,Ручка металл с лого,Канцелярия,120,18,Союзный 45\n"
    ).encode()

    result = await import_miniapp_products(
        db_session,
        max_user_id=53364725,
        tenant_slug="nizhniy-novgorod-partner-a",
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )

    assert result.created_products == 1
    assert result.created_categories == 1
    assert result.created_warehouses == 1
    assert result.updated_inventory == 1

    product = await db_session.scalar(select(Product).where(Product.sku == "PEN-LOGO"))
    category = await db_session.scalar(select(ProductCategory))
    warehouse = await db_session.scalar(select(Warehouse))
    inventory = await db_session.scalar(select(WarehouseInventory))

    assert product is not None
    assert product.name == "Ручка металл с лого"
    assert product.price_astrocoins == 120
    assert category is not None
    assert category.name == "Канцелярия"
    assert warehouse is not None
    assert warehouse.name == "Союзный 45"
    assert inventory is not None
    assert inventory.available_quantity == 18


async def test_product_import_merges_category_spacing_variants(db_session, tmp_path) -> None:
    await seed_admin(db_session)
    content = (
        "sku,name,category,price_astrocoins,quantity,warehouse\n"
        "MUG-1,Кружка,Кружки/бутылки,120,5,Главный склад\n"
        "BOTTLE-1,Бутылка,Кружки / бутылки,180,4,Главный склад\n"
    ).encode()

    parsed_rows = parse_product_rows("products.csv", content)
    assert {row.category_name for row in parsed_rows} == {"Кружки / бутылки"}
    assert len({row.category_slug for row in parsed_rows}) == 1

    result = await import_miniapp_products(
        db_session,
        max_user_id=53364725,
        tenant_slug="nizhniy-novgorod-partner-a",
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )

    categories = list(await db_session.scalars(select(ProductCategory)))
    assert result.created_categories == 1
    assert len(categories) == 1
    assert categories[0].name == "Кружки / бутылки"


async def test_product_import_reuses_existing_warehouse_with_different_slug(
    db_session,
    tmp_path,
) -> None:
    await seed_admin(db_session)
    tenant = await db_session.scalar(
        select(Tenant).where(Tenant.slug == "nizhniy-novgorod-partner-a")
    )
    assert tenant is not None
    existing_warehouse = Warehouse(
        tenant_id=tenant.id,
        slug="бор-центр",
        name="Бор Центр",
        warehouse_type=WarehouseType.COMMON,
        address="ул. Ленина, 157",
    )
    db_session.add(existing_warehouse)
    await db_session.commit()

    content = (
        "sku,name,category,price_astrocoins,quantity,warehouse\n"
        "NOTEBOOK,Блокнот,Канцелярия,120,18,БОР\u00a0  ЦЕНТР\n"
    ).encode()
    result = await import_miniapp_products(
        db_session,
        max_user_id=53364725,
        tenant_slug="nizhniy-novgorod-partner-a",
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )

    warehouses = list(await db_session.scalars(select(Warehouse)))
    inventory = await db_session.scalar(select(WarehouseInventory))
    assert result.created_warehouses == 0
    assert len(warehouses) == 1
    assert warehouses[0].id == existing_warehouse.id
    assert warehouses[0].name == "Бор Центр"
    assert warehouses[0].address == "ул. Ленина, 157"
    assert inventory is not None
    assert inventory.warehouse_id == existing_warehouse.id
    assert inventory.available_quantity == 18


async def test_product_import_links_same_named_warehouse_from_partner_city(
    db_session,
    tmp_path,
) -> None:
    await seed_admin(db_session)
    source_tenant = await db_session.scalar(
        select(Tenant).where(Tenant.slug == "nizhniy-novgorod-partner-a")
    )
    account = await db_session.scalar(select(MaxAccount).where(MaxAccount.max_user_id == 53364725))
    assert source_tenant is not None
    assert account is not None

    existing_warehouse = Warehouse(
        tenant_id=source_tenant.id,
        slug="nn-vaneeva-133",
        name="НН Ванеева 133",
        warehouse_type=WarehouseType.COMMON,
        address="Нижний Новгород, ул. Ванеева, 133",
    )
    source_category = ProductCategory(
        tenant_id=source_tenant.id,
        slug="uchebnye-nabory",
        name="Учебные наборы",
    )
    source_product = Product(
        tenant_id=source_tenant.id,
        category=source_category,
        sku="ROBOT",
        name="Робот",
        price_astrocoins=700,
    )
    bor_city = City(slug="bor", name="Бор")
    db_session.add_all([existing_warehouse, source_product, bor_city])
    await db_session.flush()
    db_session.add(
        WarehouseInventory(
            tenant_id=source_tenant.id,
            warehouse_id=existing_warehouse.id,
            product_id=source_product.id,
            available_quantity=5,
        )
    )
    bor_tenant = Tenant(
        city_id=bor_city.id,
        partner_id=source_tenant.partner_id,
        slug="bor-partner-a",
        name="Бор / Партнер A",
    )
    db_session.add(bor_tenant)
    await db_session.flush()
    db_session.add(
        StaffRoleAssignment(
            tenant_id=bor_tenant.id,
            account_id=account.id,
            role=StaffRole.ADMIN,
            status=AssignmentStatus.ACTIVE,
        )
    )
    await db_session.commit()

    content = (
        "sku,name,category,price_astrocoins,quantity,warehouse\n"
        "ROBOT,Робот,Учебные наборы,750,12,НН\u00a0  ВАНЕЕВА 133\n"
    ).encode()
    result = await import_miniapp_products(
        db_session,
        max_user_id=account.max_user_id,
        tenant_slug=bor_tenant.slug,
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )

    warehouses = list(await db_session.scalars(select(Warehouse)))
    products = list(await db_session.scalars(select(Product)))
    inventories = list(await db_session.scalars(select(WarehouseInventory)))
    link = await db_session.scalar(
        select(WarehouseTenantLink).where(
            WarehouseTenantLink.tenant_id == bor_tenant.id,
            WarehouseTenantLink.warehouse_id == existing_warehouse.id,
        )
    )
    assert result.created_warehouses == 0
    assert result.created_products == 0
    assert result.updated_products == 1
    assert len(warehouses) == 1
    assert len(products) == 1
    assert products[0].id == source_product.id
    assert products[0].tenant_id == source_tenant.id
    assert products[0].price_astrocoins == 750
    assert len(inventories) == 1
    assert inventories[0].warehouse_id == existing_warehouse.id
    assert inventories[0].available_quantity == 12
    assert link is not None


async def test_admin_imports_public_product_photo_url(
    db_session,
    tmp_path,
    monkeypatch,
) -> None:
    await seed_admin(db_session)
    photo_url = "https://cdn.example.org/catalog/pen.png?size=large"
    saved_path = tmp_path / "imported.png"

    async def fake_save_remote_product_image(
        value: str,
        *,
        media_root: str,
    ) -> SavedProductImage:
        assert value == photo_url
        assert media_root == str(tmp_path)
        saved_path.write_bytes(b"\x89PNG\r\n\x1a\n")
        return SavedProductImage(
            path=saved_path,
            url_path="/media/products/imported.png",
        )

    monkeypatch.setattr(
        "app.services.product_import.save_remote_product_image",
        fake_save_remote_product_image,
    )
    content = (
        "Артикул,Название,Цена AC,Остаток,Склад,Ссылка на фото\n"
        f"PEN-PHOTO,Ручка с фото,150,7,Главный склад,{photo_url}\n"
    ).encode()

    await import_miniapp_products(
        db_session,
        max_user_id=53364725,
        tenant_slug="nizhniy-novgorod-partner-a",
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )

    product = await db_session.scalar(select(Product).where(Product.sku == "PEN-PHOTO"))
    assert product is not None
    assert product.photo_url == "https://algo.test/media/products/imported.png"
    assert saved_path.exists()


async def test_admin_imports_product_without_sku_from_template(db_session, tmp_path) -> None:
    await seed_admin(db_session)
    rows = parse_product_rows("products.xlsx", build_product_import_template())
    assert rows == []

    content = (
        "Название,Категория,Цена AC,Склад,Остаток\nНабор наклеек,Сувениры,75,Главный склад,12\n"
    ).encode()
    result = await import_miniapp_products(
        db_session,
        max_user_id=53364725,
        tenant_slug="nizhniy-novgorod-partner-a",
        filename="products.csv",
        content=content,
        media_root=str(tmp_path),
        media_base_url="https://algo.test",
    )
    product = await db_session.scalar(select(Product).where(Product.name == "Набор наклеек"))
    assert result.created_products == 1
    assert product is not None
    assert product.sku.startswith("PRD-")


def test_product_import_template_contains_photo_link_column_and_example() -> None:
    content = build_product_import_template()
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Инструкция", "Товары", "Пример"]
        headers = [cell.value for cell in next(workbook["Товары"].iter_rows(max_row=1))]
        assert headers == [
            "Артикул",
            "Название",
            "Категория",
            "Цена AC",
            "Склад",
            "Остаток",
            "Статус",
            "Описание",
            "Ссылка на фото",
        ]
        assert workbook["Пример"]["I2"].value.startswith("https://")
        instructions = "\n".join(
            str(row[0].value or "") for row in workbook["Инструкция"].iter_rows()
        )
        assert "Яндекс Диска" in instructions
        assert "Google Drive" in instructions
    finally:
        workbook.close()


def test_product_import_accepts_semicolon_csv_and_rejects_duplicate_inventory() -> None:
    rows = parse_product_rows(
        "products.csv",
        ("Артикул;Название;Цена;Остаток;Склад\nPEN-1;Ручка;120;5;Главный склад\n").encode(),
    )
    assert rows[0].sku == "PEN-1"
    assert rows[0].quantity == 5

    with pytest.raises(ProductImportError, match="повторяется товар PEN-1"):
        parse_product_rows(
            "products.csv",
            (
                "sku,name,price,quantity,warehouse\n"
                "PEN-1,Ручка,120,5,Главный склад\n"
                "PEN-1,Ручка,120,7,Главный склад\n"
            ).encode(),
        )


@pytest.mark.parametrize("photo_url", ["ftp://files.example.org/photo.jpg", "not-a-url"])
def test_product_import_rejects_non_http_photo_url(photo_url: str) -> None:
    with pytest.raises(ProductImportError, match="ссылка на фото должна начинаться"):
        parse_product_rows(
            "products.csv",
            (
                "sku,name,price,quantity,warehouse,photo_url\n"
                f"PEN-1,Ручка,120,5,Главный склад,{photo_url}\n"
            ).encode(),
        )


@pytest.mark.parametrize(
    "filename, content",
    [("products.txt", b"sku,name"), ("products.xlsx", b"broken")],
)
def test_product_import_rejects_unsupported_or_broken_files(
    filename: str,
    content: bytes,
) -> None:
    with pytest.raises(ProductImportError):
        parse_product_rows(filename, content)
