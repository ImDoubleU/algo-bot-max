from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.crm_import import (
    CRM_TEMPLATE_COLUMNS,
    CRM_TEMPLATE_SHEET_NAME,
    CrmImportError,
    build_crm_import_template,
    parse_crm_students,
    parse_crm_students_content,
)

SHEET_DEALS = "\u0421\u0434\u0435\u043b\u043a\u0438"


def test_parse_crm_students_reads_active_export_shape(tmp_path: Path) -> None:
    path = tmp_path / "active.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_DEALS
    sheet.append(
        [
            "ID \u0441\u0434\u0435\u043b\u043a\u0438",
            "UUID",
            "ID \u0443\u0447\u0435\u043d\u0438\u043a\u0430",
            "\u0418\u043c\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430 \u0438\u0437 LMS",
            "\u0424\u0430\u043c\u0438\u043b\u0438\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430 \u0438\u0437 LMS",  # noqa: E501
            "\u041d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u0433\u0440\u0443\u043f\u043f\u044b LMS",  # noqa: E501
            "\u041d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043a\u0443\u0440\u0441\u0430 \u0438\u0437 LMS",  # noqa: E501
            "\u041d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u043f\u043b\u043e\u0449\u0430\u0434\u043a\u0438",  # noqa: E501
            "\u041f\u0440\u0435\u043f\u043e\u0434\u0430\u0432\u0430\u0442\u0435\u043b\u044c",
            "\u0413\u043e\u0440\u043e\u0434",
            "ID \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u043e\u0432",
            "\u0418\u043c\u0435\u043d\u0430 \u043a\u043e\u043d\u0442\u0430\u043a\u0442\u043e\u0432",
            "\u0418\u043c\u044f \u0441\u0442\u0430\u0442\u0443\u0441\u0430",
        ]
    )
    sheet.append(
        [
            1357,
            "uuid-1",
            "ST-001",
            "\u0410\u043b\u0438\u0441\u0430",
            "\u0412\u0430\u0441\u0438\u043b\u044c\u0435\u0432\u0430",
            "\u0421\u043e\u044e\u0437\u043d\u044b\u0439 45, \u0432\u0441 10:00",
            "Python Start",
            "\u0421\u043e\u044e\u0437\u043d\u044b\u0439 45",
            "\u041e\u043b\u0435\u0439\u043d\u0438\u043a \u0414",
            "\u041d\u0438\u0436\u043d\u0438\u0439 \u041d\u043e\u0432\u0433\u043e\u0440\u043e\u0434",
            "681, 682",
            "\u041c\u0430\u043c\u0430 \u0410\u043b\u0438\u0441\u044b",
            "\u0410\u043a\u0442\u0438\u0432\u0435\u043d",
        ]
    )
    workbook.save(path)

    rows = parse_crm_students(path)

    assert len(rows) == 1
    assert rows[0].deal_id == "1357"
    assert rows[0].lms_student_id == "ST-001"
    assert rows[0].first_name == "\u0410\u043b\u0438\u0441\u0430"
    assert rows[0].venue_name == "\u0421\u043e\u044e\u0437\u043d\u044b\u0439 45"
    assert (
        rows[0].city
        == "\u041d\u0438\u0436\u043d\u0438\u0439 \u041d\u043e\u0432\u0433\u043e\u0440\u043e\u0434"
    )
    assert rows[0].contact_ids == "681, 682"
    assert rows[0].needs_generated_access_code is False


def test_parse_crm_students_marks_missing_contact_id(tmp_path: Path) -> None:
    path = tmp_path / "active.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_DEALS
    sheet.append(
        [
            "ID \u0441\u0434\u0435\u043b\u043a\u0438",
            "\u0418\u043c\u044f \u0440\u0435\u0431\u0435\u043d\u043a\u0430 \u0438\u0437 LMS",
            "\u041d\u0430\u0437\u0432\u0430\u043d\u0438\u0435 \u0433\u0440\u0443\u043f\u043f\u044b LMS",  # noqa: E501
            "\u0413\u043e\u0440\u043e\u0434",
        ]
    )
    sheet.append(
        [
            1358,
            "\u0418\u0432\u0430\u043d",
            "\u0413\u0430\u0433\u0430\u0440\u0438\u043d\u0430 64, \u0441\u0431 18:00",
            "\u0411\u043e\u0440",
        ]
    )
    workbook.save(path)

    rows = parse_crm_students(path)

    assert len(rows) == 1
    assert rows[0].contact_ids is None
    assert rows[0].needs_generated_access_code is True


def test_standard_template_ignores_other_sheets_and_columns() -> None:
    workbook = load_workbook(BytesIO(build_crm_import_template()))
    sheet = workbook[CRM_TEMPLATE_SHEET_NAME]
    values_by_field = {
        "lms_student_id": "ST-001",
        "deal_id": "1357",
        "last_name": "\u0412\u0430\u0441\u0438\u043b\u044c\u0435\u0432\u0430",
        "first_name": "\u0410\u043b\u0438\u0441\u0430",
        "birth_date": "15.04.2015",
        "group_name": "Python, \u0432\u0442 17:30",
        "course_name": "Python Start",
        "venue_name": "\u0413\u0430\u0433\u0430\u0440\u0438\u043d\u0430 64",
        "teacher_name": "\u0418\u0432\u0430\u043d\u043e\u0432\u0430 \u0410\u043d\u043d\u0430",
        "contact_ids": "30420713",
        "contact_names": "Васильева Елена",
    }
    sheet.append([values_by_field[field] for field, *_ in CRM_TEMPLATE_COLUMNS])
    extra_column = len(CRM_TEMPLATE_COLUMNS) + 1
    sheet.cell(row=1, column=extra_column, value="Лишняя колонка")
    sheet.cell(row=2, column=extra_column, value="must-be-ignored")
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    rows = parse_crm_students_content(output.getvalue())

    assert len(rows) == 1
    assert rows[0].lms_student_id == "ST-001"
    assert rows[0].first_name == "\u0410\u043b\u0438\u0441\u0430"
    assert rows[0].deal_id == "1357"
    assert rows[0].birth_date == date(2015, 4, 15)


def test_standard_template_rejects_duplicate_student_ids() -> None:
    workbook = load_workbook(BytesIO(build_crm_import_template()))
    sheet = workbook[CRM_TEMPLATE_SHEET_NAME]
    values = [example for _field, _header, _required, _description, example in CRM_TEMPLATE_COLUMNS]
    sheet.append(values)
    sheet.append(values)
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(CrmImportError, match="повторяется ID ученика"):
        parse_crm_students_content(output.getvalue())


def test_standard_template_rejects_missing_required_value() -> None:
    workbook = load_workbook(BytesIO(build_crm_import_template()))
    sheet = workbook[CRM_TEMPLATE_SHEET_NAME]
    values = [example for _field, _header, _required, _description, example in CRM_TEMPLATE_COLUMNS]
    last_name_index = next(
        index for index, (field, *_) in enumerate(CRM_TEMPLATE_COLUMNS) if field == "last_name"
    )
    values[last_name_index] = None
    sheet.append(values)
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(CrmImportError, match="Фамилия ребенка"):
        parse_crm_students_content(output.getvalue())
